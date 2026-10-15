#!/usr/bin/env python3
"""Раскладка клинкетов по рядам на полу контейнера.

До этого площадь под задвижки считалась формулой от диаметра. Формула
занижает: она не учитывает, что задвижка лежит на боку и занимает пол
строительной длиной на полную высоту с маховиком, а поперек контейнера
их входит ровно столько, сколько влезет целых штук.

Здесь раскладка честная, ряд за рядом: сколько задвижек каждого
диаметра встает поперек, сколько получается рядов, сколько метров пола
они съедают и влезают ли в один ярус.

Габариты справочные, по типовым судовым бронзовым клинкетам. Их надо
замерить при погрузке: справочник уже ошибся на массах почти вдвое.

Укладка принята такая: задвижка лежит на боку, шток горизонтально
поперек контейнера, маховик не несет нагрузки. Высота стопки при этом
равна диаметру фланца, а не полной высоте задвижки.

Результат: outputs/marine_equipment/Раскладка_клинкетов.xlsx
"""

import pathlib
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_container_spb import asrow  # noqa: E402
from build_container_27t import collect27  # noqa: E402
from build_shipment_spb import item_kind, tier  # noqa: E402
from build_two_containers import split  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
L1_FILL = PatternFill("solid", fgColor="EAF3EA")
L2_FILL = PatternFill("solid", fgColor="FFF2CC")
WARN = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

from gate_geometry import (  # noqa: E402
    BOARD, CRATE_H, DIM, FLOOR_L, FLOOR_W, HEIGHT, across, place, rows_of,
)


def head(ws, row, ncols, height=36):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def gates(items):
    """Клинкеты нижнего яруса по диаметрам."""
    g = defaultdict(int)
    for r in items:
        if tier(asrow(r)) == 1 and item_kind(asrow(r)) == "клинкет (задвижка)":
            g[int(r["Ду"] or 0)] += int(r["Штук"])
    return g


def sheet(wb, name, title, items, note):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE

    g = gates(items)
    rows = rows_of(g)
    layers = place(rows)

    row = 4
    ws.cell(row=row, column=1, value="Сколько чего").font = SUB
    row += 1
    for i, h in enumerate(["Диаметр", "Штук", "На бок: длина х высота, мм",
                           "Высота стопки, мм", "В ряд поперек, шт",
                           "Рядов", "Съедает длины, м"], start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 7, height=40)
    total_len = 0.0
    for du in sorted(g, reverse=True):
        row += 1
        n = g[du]
        if du not in DIM:
            for i, v in enumerate([f"Ду{du}", n, "габарита нет", "", "", "",
                                   ""], start=1):
                cell = ws.cell(row=row, column=i, value=v)
                cell.font = BASE
                cell.border = BORDER
                cell.fill = WARN
            continue
        sl, fl, h = DIM[du]
        across = int(FLOOR_W * 1000 // h)
        nrows = -(-n // across)
        used = nrows * sl / 1000
        total_len += used
        for i, v in enumerate([f"Ду{du}", n, f"{sl} x {h}", fl, across,
                               nrows, round(used, 2)], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i == 7:
                cell.number_format = "#,##0.00"
    row += 1
    for i, v in enumerate(["ИТОГО", sum(g.values()), "", "", "",
                           len(rows),
                           round(total_len, 2)], start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = BOLD
        cell.border = BORDER
        if i == 7:
            cell.number_format = "#,##0.00"
    row += 1
    cell = ws.cell(row=row, column=1,
                   value=f"Пол контейнера {FLOOR_L} м. Клинкеты просят "
                         f"{total_len:.2f} м, это "
                         f"{total_len / FLOOR_L * 100:.0f}% длины. "
                         + ("В один ярус не ложатся — "
                            f"нужно {layers} яруса через доску."
                            if layers > 1 else
                            "Ложатся в один ярус."))
    cell.font = BOLD if layers > 1 else BASE
    row += 2

    ws.cell(row=row, column=1, value="Ряд за рядом, от двери вглубь").font = SUB
    row += 1
    for i, h in enumerate(["Ярус", "Ряд", "Что кладем", "Штук в ряду",
                           "Глубина ряда, м", "Нарастающим от торца, м",
                           "Высота стопки яруса, м"], start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 7, height=32)
    k = defaultdict(int)
    for r in rows:
        row += 1
        k[r["layer"]] += 1
        values = [f"Ярус {r['layer']}", k[r["layer"]], f"Клинкет Ду{r['du']}",
                  r["qty"], round(r["depth"], 2), round(r["used"], 2),
                  round(r["stack"], 2)]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.fill = L1_FILL if r["layer"] == 1 else L2_FILL
            if i in (5, 6, 7):
                cell.number_format = "#,##0.00"

    row += 2
    ws.cell(row=row, column=1, value="Высота пирога").font = SUB
    row += 1
    stack1 = max([r["stack"] for r in rows if r["layer"] == 1] or [0])
    stack2 = max([r["stack"] for r in rows if r["layer"] == 2] or [0])
    h_gates = BOARD + stack1 + (BOARD + stack2 if stack2 else 0) + BOARD
    for label, value, note_ in [
        ("Настил по полу, м", BOARD, "необрезная доска 25 мм"),
        ("Ярус 1 клинкетов, м", stack1, "высота равна диаметру фланца"),
        ("Доска между ярусами, м", BOARD if stack2 else 0, ""),
        ("Ярус 2 клинкетов, м", stack2, ""),
        ("Доска поверх, м", BOARD, "под ящики"),
        ("ВСЕГО под клинкеты, м", round(h_gates, 3), ""),
        ("Остается до потолка, м", round(HEIGHT - h_gates, 3),
         "сюда встают ящики 0,91 м и сверху фильтры"),
        ("Ярусов ящиков поместится", int((HEIGHT - h_gates) / CRATE_H), ""),
    ]:
        ws.cell(row=row, column=1, value=label).font = BASE
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0.000"
        ws.cell(row=row, column=3, value=note_).font = NOTE
        row += 1

    widths(ws, [30, 12, 26, 14, 20, 24, 22])
    return total_len, layers


def main():
    wb = Workbook()
    wb.remove(wb.active)

    items27, left, kg, m3 = collect27()
    boxes, load, tkg, rest, moves = split()

    res = {}
    res["Усиленный 27 т"] = sheet(
        wb, "Усиленный 27 т",
        "Усиленный двадцатифутовый, 27 т: раскладка клинкетов",
        items27,
        "Задвижки лежат на боку, шток горизонтально поперек контейнера. "
        "Габариты справочные, замерить при погрузке.")
    for n in sorted(boxes):
        res[f"Контейнер {n}"] = sheet(
            wb, f"Контейнер {n}",
            f"Двадцатифутовый, контейнер {n} из двух: раскладка клинкетов",
            boxes[n],
            "Задвижки лежат на боку, шток горизонтально поперек "
            "контейнера. Габариты справочные, замерить при погрузке.")

    ws = wb.create_sheet("Справочник габаритов", 0)
    ws["A1"] = "Габариты клинкетов, по которым считали"
    ws["A1"].font = TITLE
    ws["A2"] = ("Типовые судовые бронзовые клинкеты. Это справочник, а не "
                "замеры: массы из того же справочника оказались завышены "
                "почти вдвое. Замеряем при погрузке и правим таблицу в "
                "build_gate_layout.py.")
    ws["A2"].font = NOTE
    row = 4
    for i, h in enumerate(["Диаметр", "Строительная длина, мм",
                           "Диаметр фланца, мм", "Полная высота, мм",
                           "В ряд поперек контейнера, шт",
                           "Занимает пола на штуку, кв. м"], start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 6, height=40)
    for du in sorted(DIM, reverse=True):
        row += 1
        sl, fl, h = DIM[du]
        across = int(FLOOR_W * 1000 // h)
        for i, v in enumerate([f"Ду{du}", sl, fl, h, across,
                               round(sl * h / 1e6, 3)], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i == 6:
                cell.number_format = "#,##0.000"
    row += 2
    for text in [
        "Задвижку кладем на бок: шток горизонтально, маховик нагрузки не "
        "несет. Тогда высота стопки равна диаметру фланца, а не полной "
        "высоте задвижки, и ярусов помещается больше.",
        "Поперек контейнера задвижки встают полной высотой: ширина "
        f"{FLOOR_W} м делится на эту высоту, дробные места не считаем.",
        "Валетом, через одну, головой в разные стороны, можно выиграть "
        "около десяти процентов длины. В расчете этого запаса нет, он в "
        "руках бригады.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1
    widths(ws, [16, 24, 22, 20, 28, 26])

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Раскладка_клинкетов.xlsx"
    wb.save(path)
    print(f"готово: {path}\n")
    for name, (length, layers) in res.items():
        print(f"{name:<18} клинкеты просят {length:>5.2f} м пола при "
              f"{FLOOR_L} — ярусов {layers}")


if __name__ == "__main__":
    main()
