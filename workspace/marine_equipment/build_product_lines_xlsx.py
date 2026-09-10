#!/usr/bin/env python3
"""Подборки по товарным линейкам: фильтры забортной воды, клинкеты.

Покупатель спрашивает не «арматуру», а «фильтр Ду150» или «клинкет
Ду300». В учете эти позиции разбросаны по локациям и названы
по-разному: клинкет и задвижка клинкетная — одно и то же. Здесь они
собраны вместе и разложены по диаметрам.

Результат: outputs/marine_equipment/Фильтры_и_клинкеты.xlsx
"""

import pathlib
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_marine_registry import SOURCES, parse  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
SUBHEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
NOPRICE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DISCLAIMER = ("Количества по учету на 16.06.2025 и не подтверждены "
              "инвентаризацией. Перед отправкой покупателю пересчитать.")

# основное изделие и комплектующие к нему
FILTERS = re.compile(r"фильтр.*(заборт|заборн)|фильтр фланц.*заборт", re.I)
FILTER_PARTS = re.compile(r"сетка на фильтр", re.I)

GATES = re.compile(r"^клинкет|^задвижка клинкетная", re.I)
GATE_PARTS = re.compile(r"блины к клинкету|электропривод к задвижке|"
                        r"^задвижки ", re.I)

COLUMNS = ["№", "Диаметр", "Наименование", "Чертеж", "Материал", "Наличие",
           "Цена за ед., руб", "Сумма, руб", "Состояние", "Комплектность",
           "Локация"]


def head(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def diameter_num(name):
    m = re.search(r"[дД][уУ]\s*(\d{1,3})", name)
    return int(m.group(1)) if m else 0


def diameter(name):
    n = diameter_num(name)
    return f"Ду{n}" if n else "не указан"


def material(name):
    low = name.lower()
    if "бронз" in low:
        return "бронза"
    if "титан" in low:
        return "титан"
    if "медь" in low or "медн" in low:
        return "медь"
    if "сталь" in low or "стальн" in low:
        return "сталь"
    return ""


def block(ws, row, rows, caption, counter):
    """Один блок таблицы: заголовок и позиции."""
    ws.cell(row=row, column=1, value=caption)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = SUBHEAD_FILL
        cell.font = Font(name=FONT, size=11, bold=True, color="1F3864")
        cell.border = BORDER
    first = row + 1

    for r in rows:
        row += 1
        counter[0] += 1
        values = [counter[0], diameter(r["Наименование"]), r["Наименование"],
                  r["Чертеж"], material(r["Наименование"]), r["Наличие"],
                  r["Цена за ед., руб"] or None, r["Сумма"] or None,
                  r["Состояние"], r["Комплектность"], r["Локация"]]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 3))
            if i in (7, 8):
                cell.number_format = "#,##0"
        if not r["Цена за ед., руб"]:
            ws.cell(row=row, column=7).fill = NOPRICE_FILL

    row += 1
    ws.cell(row=row, column=3, value=f"Итого: {caption.lower()}").font = Font(
        name=FONT, bold=True)
    for col in (6, 8):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=SUM({letter}{first}:{letter}{row - 1})")
        cell.font = Font(name=FONT, bold=True)
        cell.border = BORDER
        if col == 8:
            cell.number_format = "#,##0"
    return row + 2


def sheet_line(wb, name, title, main, parts, main_caption, parts_caption):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = NOTE

    start = 4
    for i, h in enumerate(COLUMNS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLUMNS))

    counter = [0]
    row = block(ws, start + 1, main, main_caption, counter)
    if parts:
        row = block(ws, row, parts, parts_caption, counter)

    # сводка по диаметрам: то, что спрашивает покупатель
    ws.cell(row=row, column=1, value="Сколько чего по диаметрам").font = TITLE
    row += 1
    for i, h in enumerate(["Диаметр", "Штук", "Сумма, руб"], start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 3, height=22)

    agg = defaultdict(lambda: [0, 0])
    for r in main:
        key = diameter_num(r["Наименование"])
        agg[key][0] += r["Наличие"]
        agg[key][1] += r["Сумма"]
    for key in sorted(agg):
        row += 1
        label = f"Ду{key}" if key else "не указан"
        for i, v in enumerate([label, agg[key][0], agg[key][1] or None],
                              start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i == 3:
                cell.number_format = "#,##0"

    ws.freeze_panes = f"C{start + 1}"
    widths(ws, [5, 12, 54, 20, 12, 10, 15, 15, 12, 14, 16])
    return ws


def main():
    rows = []
    for spec in SOURCES:
        rows.extend(parse(spec))
    for r in rows:
        r["Сумма"] = (r["Наличие"] or 0) * (r["Цена за ед., руб"] or 0)
    live = [r for r in rows if (r["Наличие"] or 0) > 0]

    filters = sorted([r for r in live if FILTERS.search(r["Наименование"])],
                     key=lambda r: diameter_num(r["Наименование"]))
    filter_parts = [r for r in live if FILTER_PARTS.search(r["Наименование"])]

    gates = sorted([r for r in live if GATES.search(r["Наименование"])],
                   key=lambda r: (diameter_num(r["Наименование"]),
                                  r["Наименование"]))
    gate_parts = [r for r in live if GATE_PARTS.search(r["Наименование"])]

    wb = Workbook()
    wb.remove(wb.active)
    sheet_line(wb, "Фильтры забортной воды",
               "Фильтры забортной воды: все, что есть на складе",
               filters, filter_parts,
               "Фильтры забортной воды", "Сетки и комплектующие")
    sheet_line(wb, "Клинкеты и задвижки",
               "Клинкеты и задвижки: все, что есть на складе "
               "(клинкет и задвижка клинкетная — одно и то же)",
               gates, gate_parts,
               "Клинкеты и задвижки клинкетные", "Комплектующие и приводы")

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Фильтры_и_клинкеты.xlsx"
    wb.save(path)

    print(f"готово: {path}")
    print(f"фильтры: {len(filters)} позиций, "
          f"{sum(r['Наличие'] for r in filters):.0f} шт, "
          f"{sum(r['Сумма'] for r in filters):,.0f} руб".replace(",", " "))
    print(f"  комплектующие: {len(filter_parts)} позиций")
    print(f"клинкеты: {len(gates)} позиций, "
          f"{sum(r['Наличие'] for r in gates):.0f} шт, "
          f"{sum(r['Сумма'] for r in gates):,.0f} руб".replace(",", " "))
    print(f"  комплектующие: {len(gate_parts)} позиций")


if __name__ == "__main__":
    main()
