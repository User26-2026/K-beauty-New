#!/usr/bin/env python3
"""Перечень к отгрузке контейнером: факт погрузки 26.10.2026.

Контейнер загружен и ушел в Петербург 26.10.2026. Фактический
перечень — ФАКТ_КОНТЕЙНЕР_26.10.26, он и есть источник для
приложения 4: это то, что действительно передано Покупателю.
Перечень менеджера от 21.10 остается рядом как ПЛАН, по нему
считается сверка «что собирались везти против того, что уехало».

Это рабочий перечень менеджера, а не наш свод из учета. В нем уже
проставлены чертежи, состояние, паспортные веса и число маховиков, а
позиции разложены по группам в том порядке, в каком их грузят.

Чем он лучше нашего прежнего свода:

- веса паспортные и по БРОНЗЕ, а не по латуни: Ду350 — 212,5 кг
  против 228, Ду300 — 192, Ду250 — 125, Ду200 — 76,5 против 88;
- чертеж есть почти у каждой позиции, в том числе у задвижек, где у
  нас его не было совсем;
- состояние проставлено по всем позициям, а не у 54 из 810;
- отдельной колонкой идет число маховиков: у задвижек Ду80 их 24 на
  36 штук, у Ду300 — 11 на 12. Маховик снимается в ящик с биркой, и
  без него задвижка неполная.

Итог факта: 222 позиции, 1 810 шт, 19 266,72 кг — пересчет сходится с
«ИТОГО» файла, это контроль корректности разбора. План был 218
позиций, 1 775 шт, 19 209,42 кг.

В факте есть и упаковочные листы, по одному на ящик или коробку:
их разбирает packing_lists.py.

Источники:
  data/marine_equipment/spb_list/Факт_контейнер_26.10.26.xlsx   — факт
  data/marine_equipment/spb_list/Перечень_контейнер_21.10.26.xlsx — план
"""

import pathlib
import re

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
SPB = ROOT / "data" / "marine_equipment" / "spb_list"
SRC = SPB / "Факт_контейнер_26.10.26.xlsx"
PLAN = SPB / "Перечень_контейнер_21.10.26.xlsx"
SHEET = "Перечень для отгрузки"

# колонки листа, 1-based
COL = dict(num=1, photo=2, name=3, du=4, ru=5, draw=6, cond=7, qty=8,
           qty_total=9, unit_kg=10, kg=11, handwheel=12, group_kg=13,
           pack=14)

NO_DRAW = re.compile(r"без\s*чертежа|^-+$|^\.+$", re.I)

# Позиции, снятые с отгрузки по решению собственника. Ключ — очищенное
# наименование, значение — почему сняли. В спецификацию и в расчеты
# такие строки не попадают, но остаются здесь видимыми: иначе при
# следующей выгрузке перечня они вернутся молча.
#
# На факте 26.10 это правило уже не срабатывает: стальной Ду300 в
# контейнер действительно не поехал, бригада его не грузила. Список
# оставлен для сверки с планом и на будущие отгрузки.
EXCLUDED = {
    "задвижка клинкетная стальная ду300":
        "снята с отгрузки 22.10.2026, решение собственника",
}


def key(name):
    """Наименование для сверки с исключениями: без лишних пробелов."""
    return re.sub(r"\s+", " ", str(name or "")).strip().lower()


def text(v):
    return str(v).strip() if v is not None else ""


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = text(v).replace(",", ".").replace(" ", "")
    return float(s) if re.match(r"^-?\d+(\.\d+)?$", s) else None


def material(name, group):
    """Материал позиции по названию и группе перечня."""
    s = f"{name} {group}".lower()
    if "титан" in s:
        return "титан"
    if "амг" in s or "грибк" in s:
        # АМг — алюминиево-магниевый сплав
        return "алюминий"
    if "нерж" in s:
        return "нержавейка"
    if re.search(r"сталь|стальн", s):
        return "сталь"
    # Вся остальная арматура бронзовая — подтверждено собственником
    # 25.10.2026. Раньше класс назывался «бронза/латунь» и в ломе
    # считался по латуни как осторожная оценка.
    return "бронза"


def read(src=SRC, skip_excluded=True):
    """Строки перечня. Подзаголовки групп становятся полем «Группа».

    По умолчанию читается факт погрузки. Чтобы получить план, передай
    PLAN: сверка «план против факта» считается именно так. В сверке
    план читается с skip_excluded=False — иначе снятые позиции исчезнут
    с обеих сторон и проверить, что они не поехали, будет нечем.
    """
    ws = load_workbook(src, data_only=True)[SHEET]
    rows = []
    group = ""
    for r in ws.iter_rows(min_row=2, values_only=True):
        def cell(key):
            i = COL[key] - 1
            return r[i] if i < len(r) else None

        name = text(cell("name"))
        n = cell("num")
        if not name:
            # подзаголовок группы стоит в первой колонке без номера
            head = text(r[0])
            if head and not num(head):
                group = head
            continue
        if name.lower().startswith("итого"):
            continue
        if n is None and not num(cell("qty")):
            # строка-подзаголовок внутри группы, набранная в колонке
            # наименования: «Клапан пожарный», «Титановая арматура»
            group = name
            continue
        qty = num(cell("qty")) or 0
        if qty <= 0:
            continue
        if skip_excluded and key(name) in EXCLUDED:
            continue
        unit = num(cell("unit_kg")) or 0
        draw = text(cell("draw"))
        rows.append({
            "Наименование": name,
            "Группа": group,
            "Ду": int(num(cell("du")) or 0),
            "Ру": text(cell("ru")),
            "Чертеж": "" if NO_DRAW.search(draw) else draw,
            "Состояние": text(cell("cond")),
            "Штук": int(qty),
            "вес_ед": unit,
            "вес": num(cell("kg")) or qty * unit,
            "Маховиков": num(cell("handwheel")),
            "Упаковка": text(cell("pack")),
            "Материал": material(name, group),
        })
    return rows


if __name__ == "__main__":
    rows = read()
    if EXCLUDED:
        print("снято с отгрузки:",
              "; ".join(f"{k} — {v}" for k, v in EXCLUDED.items()))
    print(f"{len(rows)} позиций, {sum(r['Штук'] for r in rows):,} шт, "
          f"{sum(r['вес'] for r in rows):,.2f} кг".replace(",", " "))
    from collections import defaultdict
    agg = defaultdict(lambda: [0, 0, 0.0])
    for r in rows:
        a = agg[r["Группа"]]
        a[0] += 1
        a[1] += r["Штук"]
        a[2] += r["вес"]
    print(f"\n{'группа':<40}{'поз':>5}{'шт':>7}{'кг':>10}")
    for g, (n, q, kg) in agg.items():
        print(f"{g[:40]:<40}{n:>5}{q:>7}{kg:>10.1f}")
    mat = defaultdict(lambda: [0, 0.0])
    for r in rows:
        m = mat[r["Материал"]]
        m[0] += 1
        m[1] += r["вес"]
    print("\nматериал:", {k: f"{v[0]} поз, {v[1]:.0f} кг"
                         for k, v in mat.items()})
