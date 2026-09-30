#!/usr/bin/env python3
"""Склад в Петербурге: приемка двух контейнеров и цена продажи.

Арматура едет не под конкретную сделку, а на склад, с которого будет
продаваться. Значит нужны две вещи, которых не было в плане отгрузки:
чем принимать груз и от чего считать цену.

Себестоимость на питерском складе — это цена товара плюс доставка,
разнесенная по весу. Расходы на контейнеры вводятся руками: ставки
на момент сборки файла еще нет.

Результат: outputs/marine_equipment/Склад_Питер.xlsx
"""

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_shipment_spb import item_kind, material  # noqa: E402
from build_two_containers import BOX_KG, NAMES, split  # noqa: E402
from build_weight_xlsx import du  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
CALC_FILL = PatternFill("solid", fgColor="EAF3EA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COSTS = "Расходы на доставку"
PRICE = "Цена на складе в Питере"

CONDITIONS = ["Новое", "Б/у", "Восстановленное"]

# строки расходов на контейнер
COST_LINES = [
    "Аренда или покупка контейнера",
    "Фрахт Владивосток — Петербург",
    "Автодоставка до терминала во Владивостоке",
    "Погрузка, кран, такелаж",
    "Тара: поддоны, доска, брус, ящики",
    "Упаковочные материалы",
    "Страхование груза",
    "Выгрузка и автодоставка на склад в Питере",
    "Прочее",
]


def head(ws, row, ncols, height=40):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_costs(wb, load):
    """Расходы по каждому контейнеру и цена килограмма доставки."""
    ws = wb.create_sheet(COSTS, 0)
    ws["A1"] = "Во что обойдется доставка двух контейнеров"
    ws["A1"].font = TITLE
    ws["A2"] = ("Желтое вводим сами: ставок на момент сборки файла нет. "
                "Все остальное считается. Цена килограмма отсюда идет в "
                "себестоимость каждой позиции.")
    ws["A2"].font = NOTE

    row = 4
    for i, h in enumerate(["Расход", NAMES[1], NAMES[2], "Итого"], start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 4, height=24)

    first = row + 1
    for name in COST_LINES:
        row += 1
        ws.cell(row=row, column=1, value=name).font = BASE
        ws.cell(row=row, column=1).border = BORDER
        for col in (2, 3):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.number_format = "#,##0"
        cell = ws.cell(row=row, column=4, value=f"=B{row}+C{row}")
        cell.fill = CALC_FILL
        cell.border = BORDER
        cell.number_format = "#,##0"

    row += 1
    total_row = row
    ws.cell(row=row, column=1, value="ИТОГО расходов").font = BOLD
    for col in (2, 3, 4):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=SUM({letter}{first}:{letter}{row - 1})")
        cell.font = BOLD
        cell.fill = CALC_FILL
        cell.border = BORDER
        cell.number_format = "#,##0"

    row += 2
    ws.cell(row=row, column=1, value="Что везем").font = SUB
    row += 1
    weight_row = row
    ws.cell(row=row, column=1, value="Вес груза, кг").font = BASE
    for col, n in ((2, 1), (3, 2)):
        cell = ws.cell(row=row, column=col, value=round(load[n][0]))
        cell.font = BASE
        cell.border = BORDER
        cell.number_format = "#,##0"
    cell = ws.cell(row=row, column=4, value=f"=B{row}+C{row}")
    cell.border = BORDER
    cell.number_format = "#,##0"

    row += 1
    rate_row = row
    ws.cell(row=row, column=1,
            value="Доставка за килограмм, руб").font = BOLD
    for col in (2, 3, 4):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=IFERROR({letter}{total_row}"
                             f"/{letter}{weight_row},0)")
        cell.font = BOLD
        cell.fill = CALC_FILL
        cell.border = BORDER
        cell.number_format = "#,##0.00"

    row += 2
    for text in [
        "Для сравнения: транспортная компания по распискам от 24.09.2026 "
        "обошлась в 81 руб за килограмм. Контейнер должен выйти в разы "
        "дешевле, иначе он не нужен.",
        f"Загрузка по весу: {NAMES[1]} {load[1][0] / 1000:.1f} т из "
        f"{BOX_KG // 1000}, {NAMES[2]} {load[2][0] / 1000:.1f} т из "
        f"{BOX_KG // 1000}.",
        "Хранение в Питере в эти расходы не входит. Если арматура "
        "пролежит на складе год-два, хранение может обойтись дороже "
        "самой доставки — ставку надо узнать заранее.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = NOTE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 28
        row += 1

    widths(ws, [46, 18, 18, 18])
    return rate_row


COLS = [
    ("№", 5),
    ("Контейнер", 12),
    ("Тип", 20),
    ("Диаметр", 10),
    ("Наименование", 46),
    ("Материал", 18),
    ("Состояние", 13),
    ("Отправлено, шт", 13),
    ("Принято в Питере, шт", 14),
    ("Расхождение", 12),
    ("Вес ед., кг", 11),
    ("Доставка на ед., руб", 14),
    ("Цена склада, руб", 14),
    ("Себестоимость в Питере, руб", 15),
    ("Доставка в цене", 12),
    ("Цена продажи в Питере, руб", 15),
    ("Наценка", 10),
    ("№ места", 9),
    ("Примечание", 26),
]

L = {name: get_column_letter(i) for i, (name, _) in enumerate(COLS, start=1)}


def sheet_price(wb, boxes, rate_row):
    ws = wb.create_sheet(PRICE)
    ws["A1"] = "Приемка и цена продажи на складе в Петербурге"
    ws["A1"].font = TITLE
    ws["A2"] = ("Принимаем груз по колонке «Принято», расхождение "
                "считается само. Себестоимость в Питере — цена склада "
                "плюс доставка по весу. От нее ставим цену продажи.")
    ws["A2"].font = NOTE

    start = 5
    for i, (name, _) in enumerate(COLS, start=1):
        ws.cell(row=start - 1, column=i, value=name)
    head(ws, start - 1, len(COLS))

    items = [(n, r) for n in sorted(boxes) for r in boxes[n]]
    for k, (n, r) in enumerate(items, start=1):
        row = start + k - 1
        d = du(r["Наименование"])
        rate = f"'{COSTS}'!${get_column_letter(1 + n)}${rate_row}"
        values = [
            k, NAMES[n], item_kind(r), f"Ду{d}" if d else "",
            r["Наименование"], material(r), None,
            r["Наличие"], None,
            f'=IF({L["Принято в Питере, шт"]}{row}="","",'
            f'{L["Принято в Питере, шт"]}{row}'
            f'-{L["Отправлено, шт"]}{row})',
            round(r["вес_ед"], 1),
            f'={L["Вес ед., кг"]}{row}*{rate}',
            r["Цена за ед., руб"] or None,
            f'=IF({L["Цена склада, руб"]}{row}="","",'
            f'{L["Цена склада, руб"]}{row}'
            f'+{L["Доставка на ед., руб"]}{row})',
            f'=IFERROR({L["Доставка на ед., руб"]}{row}'
            f'/{L["Себестоимость в Питере, руб"]}{row},"")',
            None,
            f'=IFERROR({L["Цена продажи в Питере, руб"]}{row}'
            f'/{L["Себестоимость в Питере, руб"]}{row}-1,"")',
            None, None,
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(i in (5, 19)))
        for name in ("Доставка на ед., руб", "Цена склада, руб",
                     "Себестоимость в Питере, руб",
                     "Цена продажи в Питере, руб"):
            ws[f"{L[name]}{row}"].number_format = "#,##0"
        ws[f'{L["Вес ед., кг"]}{row}'].number_format = "#,##0.0"
        for name in ("Доставка в цене", "Наценка"):
            ws[f"{L[name]}{row}"].number_format = "0%"
        for name in ("Состояние", "Принято в Питере, шт",
                     "Цена склада, руб", "Цена продажи в Питере, руб",
                     "№ места", "Примечание"):
            ws[f"{L[name]}{row}"].fill = INPUT_FILL
        for name in ("Расхождение", "Доставка на ед., руб",
                     "Себестоимость в Питере, руб", "Доставка в цене",
                     "Наценка"):
            ws[f"{L[name]}{row}"].fill = CALC_FILL

    last = start + len(items) - 1

    dv = DataValidation(type="list",
                        formula1='"' + ",".join(CONDITIONS) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{L["Состояние"]}{start}:{L["Состояние"]}{last}')

    total = last + 1
    ws.cell(row=total, column=5, value="ИТОГО").font = BOLD
    for name in ("Отправлено, шт", "Принято в Питере, шт", "Расхождение"):
        cell = ws[f"{L[name]}{total}"]
        cell.value = f"=SUM({L[name]}{start}:{L[name]}{last})"
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"

    ws.auto_filter.ref = f"A{start - 1}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = f"E{start}"
    widths(ws, [w for _, w in COLS])
    return start, last


def sheet_totals(wb, start, last, boxes):
    ws = wb.create_sheet("Итоги по складу")
    ws["A1"] = "Что лежит на складе в Питере и сколько это стоит"
    ws["A1"].font = TITLE
    ws["A2"] = ("Считается по листу цены. Пока приемка не заполнена, "
                "часть строк будет пустой.")
    ws["A2"].font = NOTE

    def rng(name):
        return f"'{PRICE}'!${L[name]}${start}:${L[name]}${last}"

    row = 4
    ws.cell(row=row, column=1, value="По контейнерам").font = SUB
    row += 1
    for i, h in enumerate(["Контейнер", "Позиций", "Отправлено, шт",
                           "Принято, шт", "Цена склада, руб",
                           "Доставка, руб", "Себестоимость в Питере, руб"],
                          start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 7, height=28)

    first = row + 1
    for n in sorted(boxes):
        row += 1
        crit = f'"{NAMES[n]}"'
        cells = [
            NAMES[n],
            f'=COUNTIF({rng("Контейнер")},{crit})',
            f'=SUMIF({rng("Контейнер")},{crit},{rng("Отправлено, шт")})',
            f'=SUMIF({rng("Контейнер")},{crit},'
            f'{rng("Принято в Питере, шт")})',
            f'=SUMPRODUCT(({rng("Контейнер")}={crit})*'
            f'IFERROR({rng("Отправлено, шт")}*{rng("Цена склада, руб")},0))',
            f'=SUMPRODUCT(({rng("Контейнер")}={crit})*'
            f'IFERROR({rng("Отправлено, шт")}*'
            f'{rng("Доставка на ед., руб")},0))',
            f"=E{row}+F{row}",
        ]
        for i, v in enumerate(cells, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i >= 5:
                cell.number_format = "#,##0"
    row += 1
    ws.cell(row=row, column=1, value="ИТОГО").font = BOLD
    for col in range(2, 8):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=SUM({letter}{first}:{letter}{row - 1})")
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"

    row += 2
    ws.cell(row=row, column=1, value="Приемка").font = SUB
    row += 1
    for name, formula in [
        ("Строк принято",
         f'=COUNT({rng("Принято в Питере, шт")})'),
        ("Строк осталось принять",
         f'=COUNTA({rng("Наименование")})'
         f'-COUNT({rng("Принято в Питере, шт")})'),
        ("Расхождение всего, шт", f'=SUM({rng("Расхождение")})'),
        ("Строк с расхождением",
         f'=COUNTIF({rng("Расхождение")},"<>0")'
         f'-COUNTBLANK({rng("Расхождение")})'),
    ]:
        ws.cell(row=row, column=1, value=name).font = BASE
        cell = ws.cell(row=row, column=2, value=formula)
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"
        row += 1

    for i, w in enumerate([34, 14, 16, 14, 18, 16, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


RULES = [
    ("Сначала расходы, потом цены",
     "Пока на листе расходов нули, доставка на единицу будет нулевой, "
     "и себестоимость в Питере совпадет с ценой склада. Ставить цену "
     "продажи от такой цифры нельзя."),
    ("Принимаем по местам, а не по памяти",
     "Номер места в этой таблице должен совпадать с биркой на ящике и "
     "строкой упаковочного листа. Расхождение фиксируем в день выгрузки, "
     "иначе повторим владивостокскую историю с недостоверными остатками."),
    ("Состояние переносим из отгрузки",
     "Новое и б/у на питерском складе продаются по разной цене. Если "
     "состояние не проставили при погрузке, на приемке это уже спорно."),
    ("Цена продажи считается от себестоимости в Питере",
     "Это цена склада плюс доставка. На тяжелой арматуре доставка "
     "заметна, на мелочи почти нет — поэтому наценку ставим по строке, "
     "а не одним процентом на весь склад."),
    ("Хранение считаем отдельно",
     "Доставка — разовый расход, хранение идет каждый месяц. Если "
     "арматура лежит год-два, хранение способно перекрыть доставку. "
     "Ставку за квадратный метр или за паллетоместо надо знать до того, "
     "как контейнеры выехали."),
    ("Позиции без цены не продаем вслепую",
     "В складских файлах часть позиций идет без цены. Пока цена не "
     "проставлена, себестоимость в Питере по ним не считается."),
]


def sheet_rules(wb):
    ws = wb.create_sheet("Как работать с файлом")
    ws["A1"] = "Порядок работы"
    ws["A1"].font = TITLE
    for i, h in enumerate(["№", "Правило", "Почему так"], start=1):
        ws.cell(row=3, column=i, value=h)
    head(ws, 3, 3, height=22)
    for n, (rule, why) in enumerate(RULES, start=1):
        row = 3 + n
        for i, v in enumerate([n, rule, why], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i > 1))
        ws.row_dimensions[row].height = 52
    for i, w in enumerate([5, 44, 84], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def main():
    boxes, load, rest = split()

    wb = Workbook()
    wb.remove(wb.active)
    rate_row = sheet_costs(wb, load)
    start, last = sheet_price(wb, boxes, rate_row)
    sheet_totals(wb, start, last, boxes)
    sheet_rules(wb)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Склад_Питер.xlsx"
    wb.save(path)

    total = sum(len(boxes[n]) for n in boxes)
    value = sum(r["Сумма"] for n in boxes for r in boxes[n])
    print(f"готово: {path}")
    print(f"позиций на приемку: {total}, "
          f"{sum(r['Наличие'] for n in boxes for r in boxes[n]):.0f} штук")
    print(f"стоимость по ценам склада: {value / 1e6:.1f} млн руб")
    for n in sorted(boxes):
        print(f"  {NAMES[n]}: {len(boxes[n])} поз, "
              f"{load[n][0] / 1000:.1f} т, "
              f"{sum(r['Сумма'] for r in boxes[n]) / 1e6:.1f} млн")


if __name__ == "__main__":
    main()
