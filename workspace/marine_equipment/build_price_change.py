#!/usr/bin/env python3
"""Что изменилось в цене после перехода на ставку за килограмм.

Собственник назначил ставку 1 862,50 руб за килограмм на бронзу и
латунь. До этого цена бралась из действующих приложений к договору,
где она стояла по позициям: где числом, где диапазоном, где «по
согласованию».

Ставка за килограмм — это другая логика ценообразования, и она дает
совсем другую раскладку. Крупное дорожает: килограмм один и тот же, а
клинкет Ду300 весит 192 кг. Мелкое дешевеет: клапан на полтора
килограмма по ставке стоит под три тысячи, а в договоре стоял пять.

Здесь считается ровно это: по каждой позиции старая цена против
новой, разница за единицу и на всю позицию.

Правила сравнения:

- старая цена берется из приложений к договору по чертежу, иначе по
  очищенному наименованию;
- где в договоре диапазон, для сравнения берется ВЕРХНЯЯ граница —
  так же, как мы считаем свою цену продажи, а исходная запись остается
  рядом отдельной колонкой;
- позиции, которых в договоре не было или где стояло «по
  согласованию», в сравнение не идут: сравнивать не с чем, они вынесены
  отдельной вкладкой;
- сталь, титан, нержавейка и алюминий считаются по-старому, ставка на
  них не распространяется.

Результат: outputs/marine_equipment/Изменение_цен_по_ставке.xlsx
"""

import pathlib
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import upd_shipments  # noqa: E402
from build_contract_annexes import (  # noqa: E402
    AGREE, BRONZE, annex4, clean, norm_draw, price_index,
)
from build_container_spb import asrow, plural  # noqa: E402
from build_logistics_compare import rows_all  # noqa: E402
from build_shipment_spb import material  # noqa: E402
from build_spb_list_xlsx import group_of  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
DOWN = PatternFill("solid", fgColor="FCE4D6")   # подешевело
UP = PatternFill("solid", fgColor="E2EFDA")     # подорожало
GREY = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RATE = upd_shipments.BRONZE_RATE


def head(ws, row, ncols, height=34):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, values, wrap=(), center=(), money=(), pct=(), fill=None,
        font=None):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font or BASE
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(i in wrap),
                                   horizontal="center" if i in center
                                   else None)
        if i in money:
            cell.number_format = "#,##0"
        if i in pct:
            cell.number_format = "+0%;-0%"
        if fill:
            cell.fill = fill
    return row + 1


def collect():
    """Позиции со старой и новой ценой."""
    by_draw, by_name = price_index()
    rows, _ = annex4()
    index = {}
    for r in rows_all():
        index[(r["Наименование"], r.get("Чертеж") or "")] = r
    out = []
    for r in rows:
        src = index.get((r["name"], r["draw"]))
        old_text = (by_draw.get(norm_draw(r["draw"]))
                    or by_name.get(clean(r["name"])) or "")
        old = upd_shipments.bounds(old_text) if AGREE not in old_text.lower() \
            else []
        new = upd_shipments.bounds(r["price"])
        out.append({
            "Наименование": r["name"],
            "Чертеж": r["draw"],
            "Ду": r["du"],
            "Штук": r["qty"],
            "вес_ед": r["unit_kg"],
            "Группа": group_of(r["name"]),
            "Материал": material(asrow(src)) if src else "",
            "по ставке": "ставка" in r["how"],
            "старый текст": old_text,
            "старая": max(old) * 1000 if old and max(old) < 1000 else
                      (max(old) if old else None),
            "новая": max(new) if new else None,
        })
    return out


COLS = [
    ("№", 5), ("Наименование", 48), ("Чертеж", 20), ("Ду", 6),
    ("Штук", 7), ("Вес ед., кг", 10), ("Как записано в договоре", 18),
    ("Старая цена за ед., руб", 13), ("Новая цена за ед., руб", 13),
    ("Разница за ед., руб", 12), ("Разница", 9),
    ("Разница на позицию, руб", 14),
]


def sheet_rows(wb, name, rows, title):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    total = sum(r["дельта_всего"] for r in rows)
    kg = sum(r["вес_ед"] * r["Штук"] for r in rows)
    ws["A2"] = (f"{len(rows)} "
                f"{plural(len(rows), 'позиция', 'позиции', 'позиций')}, "
                f"{sum(r['Штук'] for r in rows):,.0f} шт, {kg:,.0f} кг. "
                f"Итого {total:+,.0f} руб. Внутри группы — по размеру "
                f"разницы на позицию.".replace(",", " "))
    ws["A2"].font = NOTE

    start = 4
    for i, (h, _) in enumerate(COLS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLS))
    row = start + 1
    n = 0
    by_group = defaultdict(list)
    for r in rows:
        by_group[r["Группа"]].append(r)
    for g in sorted(by_group,
                    key=lambda x: sum(abs(i["дельта_всего"])
                                      for i in by_group[x]), reverse=True):
        items = by_group[g]
        row = put(ws, row, [
            f"{g} — {len(items)} "
            f"{plural(len(items), 'позиция', 'позиции', 'позиций')}, "
            f"{sum(i['дельта_всего'] for i in items):+,.0f} руб"
            .replace(",", " ")
        ], fill=GROUP_FILL, font=SUB)
        for r in items:
            n += 1
            row = put(ws, row, [
                n, r["Наименование"], r["Чертеж"] or "—", r["Ду"] or "",
                r["Штук"], round(r["вес_ед"], 1), r["старый текст"],
                round(r["старая"]), round(r["новая"]), round(r["дельта"]),
                r["доля"], round(r["дельта_всего"]),
            ], wrap=(2,), center=(1, 4, 5, 6, 11),
                money=(8, 9, 10, 12), pct=(11,),
                fill=DOWN if r["дельта"] < 0 else UP)
    ws.freeze_panes = f"A{start + 1}"
    widths(ws, [w for _, w in COLS])
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1


def sheet_nocmp(wb, rows):
    ws = wb.create_sheet("Не с чем сравнить")
    ws["A1"] = "Позиции по ставке, которых в договоре не было"
    ws["A1"].font = TITLE
    ws["A2"] = (f"{len(rows)} "
                f"{plural(len(rows), 'позиция', 'позиции', 'позиций')}, "
                f"{sum(r['Штук'] for r in rows):,.0f} шт, "
                f"{sum(r['вес_ед'] * r['Штук'] for r in rows):,.0f} кг на "
                f"{sum(r['новая'] * r['Штук'] for r in rows):,.0f} руб. "
                f"В договоре их не было совсем или стояло «по "
                f"согласованию» — сравнивать не с чем, это новая цена, "
                f"а не изменение.".replace(",", " "))
    ws["A2"].font = NOTE
    cols = [("№", 5), ("Наименование", 52), ("Чертеж", 20), ("Ду", 6),
            ("Штук", 7), ("Вес ед., кг", 10), ("Цена за ед., руб", 13),
            ("Сумма, руб", 13)]
    start = 4
    for i, (h, _) in enumerate(cols, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(cols))
    row = start + 1
    for n, r in enumerate(sorted(rows, key=lambda x: -x["новая"] * x["Штук"]),
                          start=1):
        row = put(ws, row, [n, r["Наименование"], r["Чертеж"] or "—",
                            r["Ду"] or "", r["Штук"], round(r["вес_ед"], 1),
                            round(r["новая"]),
                            round(r["новая"] * r["Штук"])],
                  wrap=(2,), center=(1, 4, 5, 6), money=(7, 8))
    ws.freeze_panes = f"A{start + 1}"
    widths(ws, [w for _, w in cols])


def sheet_summary(wb, rows, cmp, down, up, same, nocmp):
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Что изменилось в цене после перехода на ставку за килограмм"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Ставка {RATE:,.2f} руб за килограмм на бронзу и латунь. "
                f"Сравнение со старой ценой из приложений к договору."
                .replace(",", " ").replace(".00", ",00"))
    ws["A2"].font = NOTE

    minus = abs(sum(r["дельта_всего"] for r in down))
    plus = sum(r["дельта_всего"] for r in up)
    rows_sum = [
        ("Подешевело", len(down),
         sum(r["Штук"] for r in down),
         sum(r["вес_ед"] * r["Штук"] for r in down), -minus,
         "ставка дает цену ниже договорной"),
        ("Подорожало", len(up), sum(r["Штук"] for r in up),
         sum(r["вес_ед"] * r["Штук"] for r in up), plus,
         "ставка дает цену выше договорной"),
        ("Без изменения", len(same), sum(r["Штук"] for r in same),
         sum(r["вес_ед"] * r["Штук"] for r in same), 0, ""),
        ("ИТОГО сравнимых", len(cmp), sum(r["Штук"] for r in cmp),
         sum(r["вес_ед"] * r["Штук"] for r in cmp), plus - minus,
         "позиции, где в договоре была цена"),
        ("Не с чем сравнить", len(nocmp), sum(r["Штук"] for r in nocmp),
         sum(r["вес_ед"] * r["Штук"] for r in nocmp),
         sum(r["новая"] * r["Штук"] for r in nocmp),
         "в договоре цены не было, это новая цена"),
    ]
    cols = [("Что", 22), ("Позиций", 9), ("Штук", 9), ("Вес, кг", 11),
            ("Деньги, руб", 15), ("Пояснение", 44)]
    start = 4
    for i, (h, _) in enumerate(cols, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(cols))
    row = start + 1
    for a, n, q, kg, money, note in rows_sum:
        row = put(ws, row, [a, n, q, round(kg), round(money), note],
                  wrap=(6,), center=(2, 3, 4), money=(4, 5),
                  font=BOLD if a.startswith("ИТОГО") else BASE,
                  fill=DOWN if a == "Подешевело" else
                       UP if a == "Подорожало" else
                       GROUP_FILL if a.startswith("ИТОГО") else None)
    row += 1
    ws.cell(row=row, column=1, value="Почему так выходит").font = SUB
    row += 1
    for t in [
        "Ставка за килограмм не знает про размер изделия. Клинкет Ду300 "
        "весит 192 кг, и по ставке он стоит 357 600 вместо 200 000 по "
        "договору. Клапан на полтора килограмма по ставке стоит 2 794 "
        "вместо 5 000.",
        "Поэтому дорожает крупное, а дешевеет мелочь. По штукам мелочи "
        "больше, по деньгам решает крупное.",
        "Сравнение честное только там, где в договоре была цена. Где "
        "стояло «по согласованию», сравнивать не с чем — такие позиции "
        "вынесены отдельной вкладкой.",
        "Где в договоре диапазон, для сравнения взята ВЕРХНЯЯ граница. "
        "Исходная запись стоит рядом в колонке «Как записано в "
        "договоре», чтобы было видно, с чем сравнивали.",
        "Сталь, титан, нержавейка и алюминий по ставке не считаются, в "
        "это сравнение они не входят.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + t)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1
    widths(ws, [w for _, w in cols])


def main():
    rows = collect()
    cmp = [r for r in rows if r["старая"] and r["новая"] and r["по ставке"]]
    for r in cmp:
        r["дельта"] = r["новая"] - r["старая"]
        r["доля"] = r["дельта"] / r["старая"]
        r["дельта_всего"] = r["дельта"] * r["Штук"]
    down = sorted([r for r in cmp if r["дельта"] < 0],
                  key=lambda x: x["дельта_всего"])
    up = sorted([r for r in cmp if r["дельта"] > 0],
                key=lambda x: -x["дельта_всего"])
    same = [r for r in cmp if r["дельта"] == 0]
    nocmp = [r for r in rows if r["по ставке"] and not r["старая"]]

    wb = Workbook()
    sheet_summary(wb, rows, cmp, down, up, same, nocmp)
    sheet_rows(wb, "Подешевело", down,
               "Позиции, где ставка за килограмм дает цену НИЖЕ договорной")
    sheet_rows(wb, "Подорожало", up,
               "Позиции, где ставка за килограмм дает цену ВЫШЕ договорной")
    sheet_nocmp(wb, nocmp)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Изменение_цен_по_ставке.xlsx"
    wb.save(path)
    print(f"готово: {path}\n")
    print(f"ставка {RATE:,.2f} руб за кг".replace(",", " "))
    print(f"сравнимых позиций {len(cmp)}: подешевело {len(down)}, "
          f"подорожало {len(up)}, без изменения {len(same)}")
    print(f"не с чем сравнить {len(nocmp)} — в договоре цены не было")
    print(f"\nпо деньгам: минус {abs(sum(r['дельта_всего'] for r in down)):,.0f}"
          f", плюс {sum(r['дельта_всего'] for r in up):,.0f}, "
          f"итого {sum(r['дельта_всего'] for r in cmp):+,.0f} руб"
          .replace(",", " "))
    print("\nСАМОЕ СИЛЬНОЕ СНИЖЕНИЕ по деньгам:")
    for r in down[:10]:
        print(f"  {r['Наименование'][:40]:<40}{r['Штук']:>5} шт  "
              f"{r['старая']:>9,.0f} -> {r['новая']:>8,.0f}  "
              f"{r['доля']:>+6.0%}  {r['дельта_всего']:>+11,.0f}"
              .replace(",", " "))


if __name__ == "__main__":
    main()
