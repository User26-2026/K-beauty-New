#!/usr/bin/env python3
"""Позиции отгрузки 26.10.2026, по которым цена не закрыта.

Контейнер уехал, а в приложении № 4 часть позиций стоит «по
согласованию». По пункту 3.5 Договора такая позиция проданной не
считается: товар физически у Покупателя, а расчета по нему нет.
Пока эти строки не закрыты, спецификация неполная.

Файл собран, чтобы закрыть их одним проходом. По каждой позиции
рядом лежат все цифры, на которые можно опереться:

  - цена по ставке — что дала бы бронза по 1 862,50 руб за кг. Для
    нового это НИЖНЯЯ граница: новое дороже бывшего в употреблении;
  - цена из Договора — то, что стояло по этой позиции в действующих
    приложениях, если она там нашлась;
  - цена конкурента — камчатский прайс от 03.10.2026, по чертежу или
    по группе с диаметром;
  - лом — абсолютный минимум, ниже которого продавать незачем.

Колонка «СОГЛАСОВАНО» желтая и пустая: ее заполняет собственник, и
после этого цифры переносятся в price_fixes или прямо в приложение.

Результат: outputs/marine_equipment/Цены_к_согласованию_26.10.xlsx
"""

import pathlib
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import scrap_prices  # noqa: E402
import upd_shipments  # noqa: E402
from build_competitor_valuation import competitor_index, pick  # noqa: E402
from build_contract_annexes import (AGREE, annex4, clean,  # noqa: E402
                                    norm_draw, price_index, upper)
from build_container_spb import plural  # noqa: E402
from build_spb_list_xlsx import group_of  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [
    ("№", 5), ("Наименование", 46), ("Чертеж", 18), ("Ду", 6), ("Ру", 8),
    ("Сост.", 8), ("Шт", 6), ("Вес ед., кг", 9), ("Вес всего, кг", 10),
    ("Почему без цены", 26), ("По ставке за ед., руб", 12),
    ("Из Договора за ед., руб", 12), ("У конкурента за ед., руб", 12),
    ("Лом за ед., руб", 10), ("СОГЛАСОВАНО за ед., руб", 14),
    ("Сумма по согласованной, руб", 14),
]


def rub(n):
    return f"{n:,.0f}".replace(",", " ")


def head(ws, row, ncols, height=40):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def put(ws, row, values, wrap=(), center=(), money=(), fill=None, font=None):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font or BASE
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(i in wrap),
                                   horizontal="center" if i in center
                                   else None)
        if i in money:
            cell.number_format = "#,##0"
        if fill:
            cell.fill = fill
    return row + 1


def collect():
    """Строки приложения 4 без твердой цены, с подсказками по цене."""
    rows, _ = annex4()
    by_draw, by_name = price_index()
    cby_draw, cby_type, _ = competitor_index()
    out = []
    for r in rows:
        if r["price"] and r["price"] != AGREE:
            continue
        mat = r["material"]
        rate = (upd_shipments.BRONZE_RATE * r["unit_kg"]
                if r["unit_kg"] else None)
        text = (by_draw.get(norm_draw(r["draw"]))
                or by_name.get(clean(r["name"])) or "")
        contract = upper(text) if text and AGREE not in text.lower() else ""
        comp, _ = pick(cby_draw, cby_type, r["draw"],
                       group_of(r["name"]), r["du"], mat)
        scrap = scrap_prices.value(mat, r["unit_kg"])
        out.append(dict(r, rate=rate, contract=contract, comp=comp,
                        scrap=scrap))
    return out


def main():
    rows = collect()
    kg = sum(r["kg"] for r in rows)
    qty = sum(r["qty"] for r in rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "К согласованию"
    ws["A1"] = "Цены, которые надо закрыть по отгрузке 26.10.2026"
    ws["A1"].font = TITLE
    ws["A2"] = (f"{len(rows)} "
                f"{plural(len(rows), 'позиция', 'позиции', 'позиций')}, "
                f"{rub(qty)} шт, {rub(round(kg))} кг. Товар уехал в "
                f"контейнере, но по пункту 3.5 Договора до согласования "
                f"цены проданным не считается. Желтая колонка — под "
                f"заполнение, остальное справочно.")
    ws["A2"].font = NOTE

    start = 4
    for i, (h, _) in enumerate(COLS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLS))
    row = start + 1

    groups = defaultdict(list)
    for r in rows:
        groups[r["how"].split(";")[0].strip()].append(r)

    n = 0
    for why in sorted(groups, key=lambda g: -sum(x["kg"] for x in groups[g])):
        rs = sorted(groups[why], key=lambda x: -x["kg"])
        cell = ws.cell(row=row, column=1)
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=len(COLS))
        cell.value = (f"{why} — {len(rs)} "
                      f"{plural(len(rs), 'позиция', 'позиции', 'позиций')}, "
                      f"{rub(sum(x['qty'] for x in rs))} шт, "
                      f"{rub(round(sum(x['kg'] for x in rs)))} кг")
        cell.font = SUB
        cell.fill = GROUP_FILL
        for c in range(1, len(COLS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
        for r in rs:
            n += 1
            put(ws, row, [
                n, r["name"], r["draw"] or "—", r["du"] or "", r["ru"] or "",
                r["cond"], r["qty"], round(r["unit_kg"], 1), round(r["kg"]),
                why,
                round(r["rate"]) if r["rate"] else "",
                r["contract"] or "",
                round(r["comp"]) if r["comp"] else "",
                round(r["scrap"]) if r["scrap"] else "",
                "", "",
            ], wrap=(2, 10), center=(1, 4, 5, 6, 7, 8, 9),
                money=(11, 13, 14, 15, 16))
            ws.cell(row=row, column=15).fill = FILL_IN
            ws.cell(row=row, column=16).value = (
                f"=O{row}*G{row}")
            ws.cell(row=row, column=16).number_format = "#,##0"
            row += 1

    last = row - 1
    row = put(ws, row, [
        "ИТОГО", f"{len(rows)} позиций", "", "", "", "", qty, "",
        round(kg), "", "", "", "", "", "",
        f"=SUM(P{start + 1}:P{last})",
    ], center=(7, 9), money=(16,), font=BOLD, fill=GROUP_FILL)
    ws.cell(row=row - 1, column=16).number_format = "#,##0"

    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(COLS))}{last}"
    for i, (_, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.print_title_rows = f"{start}:{start}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1

    row += 1
    ws.cell(row=row, column=1, value="Как читать подсказки").font = SUB
    row += 1
    for t in [
        "«По ставке» — цена от веса по бронзе, 1 862,50 руб за "
        "килограмм. Для НОВОГО это нижняя граница: новое дороже "
        "бывшего в употреблении, потому оно и выведено из-под ставки.",
        "«Из Договора» — что стояло по этой позиции в действующих "
        "приложениях. Если пусто, в Договоре цены по ней не было.",
        "«У конкурента» — камчатский прайс от 03.10.2026, подобран по "
        "номеру чертежа, а где его нет — по группе с диаметром. Это "
        "ориентир рынка, а не наша цена.",
        "«Лом» — цена приемки за эту массу. Абсолютный минимум: ниже "
        "продавать незачем, проще сдать.",
        "Заполненную желтую колонку переносим в приложение № 4 и "
        "подписываем. До этого по пункту 3.5 товар проданным не "
        "считается, хотя физически он уже у Покупателя.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + t)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=len(COLS))
        ws.row_dimensions[row].height = 28
        row += 1

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Цены_к_согласованию_26.10.xlsx"
    wb.save(path)
    print(f"готово: {path}\n")
    print(f"{len(rows)} позиций, {rub(qty)} шт, {rub(round(kg))} кг\n")
    for why in sorted(groups, key=lambda g: -sum(x["kg"] for x in groups[g])):
        rs = groups[why]
        has = sum(1 for r in rs if r["rate"] or r["contract"] or r["comp"])
        print(f"  {why:<44}{len(rs):>3} поз, "
              f"{rub(round(sum(x['kg'] for x in rs))):>6} кг, "
              f"подсказка есть у {has}")
    rate_sum = sum((r["rate"] or 0) * r["qty"] for r in rows)
    print(f"\nесли закрыть все по ставке: {rub(round(rate_sum))} руб — "
          f"это НИЖНЯЯ оценка, новое дороже")


if __name__ == "__main__":
    main()
