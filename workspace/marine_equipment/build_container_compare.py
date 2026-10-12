#!/usr/bin/env python3
"""Сравнение двадцатифутового и сорокафутового контейнера.

Два сценария под реальные предложения перевозчика: двадцатифутовый
усиленный на 27 тонн за 315 тысяч и сорокафутовый на 26 тонн за 435
тысяч. По каждому считается, что войдет, что останется и во что
обойдется тонна.

Главное, что показывает расчет: двадцатифутовый до своих 27 тонн не
догрузится. Внутри у него 33 куб. м, и объем кончается раньше веса.
У сорокафутового 67 куб. м, там наоборот упор в вес.

Веса не справочные: 85% тоннажа стоит на замерах менеджера из
приложений 2 и 3.

Результат: outputs/marine_equipment/Сравнение_контейнеров.xlsx
"""

import pathlib
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_container_22t import (  # noqa: E402
    BOX_KG, BOX_PER_LAYER, asrow, collect, sheet_items,
)
from build_shipment_spb import tier  # noqa: E402
from build_spb_list_xlsx import GROUPS, group_of  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
WIN_FILL = PatternFill("solid", fgColor="EAF3EA")
LOSE_FILL = PatternFill("solid", fgColor="FDF2E3")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# предложения перевозчика на руках
OPTIONS = [
    dict(name="20-футовый усиленный", kg=27_000, m3=33.0, price=315_000,
         short="20 футов"),
    dict(name="40-футовый", kg=26_000, m3=67.2, price=435_000,
         short="40 футов"),
]
# третий вариант для справки: тот же размер, но обычный контейнер
PLAIN20 = dict(name="20-футовый обычный", kg=22_000, m3=33.0, price=236_000)


def head(ws, row, ncols, height=36):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def run(opt):
    a, b, rest, kg, m3 = collect(opt["kg"], opt["m3"])
    items = a + b
    # m3 из collect — занятый объем, opt["m3"] — объем контейнера
    return dict(opt, items=items, rest=rest, kg=kg,
                m3_used=m3, kg_limit=opt["kg"], m3_limit=opt["m3"],
                positions=len(items), units=sum(r["Штук"] for r in items),
                rest_kg=sum(r["вес"] for r in rest),
                rest_units=sum(r["Штук"] for r in rest),
                per_ton=opt["price"] / (kg / 1000) if kg else 0,
                limited=("объем" if m3 / opt["m3"] > kg / opt["kg"]
                         else "вес"))


def sheet_compare(wb, runs, total_kg):
    ws = wb.create_sheet("Сравнение", 0)
    ws["A1"] = "Двадцатифутовый или сорокафутовый"
    ws["A1"].font = TITLE
    ws["A2"] = ("Считано по весам менеджера: 85% тоннажа стоит на его "
                "замерах, остальное пересчитано от соседнего диаметра. "
                "Объем — по габаритам позиций плюс упаковка.")
    ws["A2"].font = NOTE

    rows = [
        ("Цена за контейнер, руб", lambda r: r["price"], "#,##0"),
        ("Паспортная загрузка, т", lambda r: r["kg_limit"] / 1000, "0.0"),
        ("Внутренний объем, куб. м", lambda r: r["m3_limit"], "0.0"),
        ("", None, None),
        ("Фактически войдет, т", lambda r: r["kg"] / 1000, "0.00"),
        ("Использовано грузоподъемности",
         lambda r: r["kg"] / r["kg_limit"], "0%"),
        ("Занято объема", lambda r: r["m3_used"] / r["m3_limit"], "0%"),
        ("Во что упирается", lambda r: r["limited"], None),
        ("", None, None),
        ("Позиций", lambda r: r["positions"], "#,##0"),
        ("Штук", lambda r: r["units"], "#,##0"),
        ("В ящиках, т", lambda r: r["boxed_kg"] / 1000, "0.00"),
        ("Ящиков собрать", lambda r: r["boxes"], "#,##0"),
        ("Мест под ящики в контейнере",
         lambda r: BOX_PER_LAYER * 2 * (2 if r["m3_limit"] > 40 else 1),
         "#,##0"),
        ("", None, None),
        ("Остается во Владивостоке, т",
         lambda r: r["rest_kg"] / 1000, "0.00"),
        ("Остается позиций", lambda r: len(r["rest"]), "#,##0"),
        ("", None, None),
        ("Цена за тонну, руб", lambda r: r["per_ton"], "#,##0"),
        ("Остаток вывезти ТК по 81 руб/кг",
         lambda r: r["rest_kg"] * 81, "#,##0"),
        ("Итого за всю арматуру, руб",
         lambda r: r["price"] + r["rest_kg"] * 81, "#,##0"),
    ]

    start = 4
    ws.cell(row=start, column=1, value="Показатель")
    for i, r in enumerate(runs, start=2):
        ws.cell(row=start, column=i, value=r["name"])
    head(ws, start, 1 + len(runs), height=30)

    row = start
    for label, fn, fmt in rows:
        row += 1
        ws.cell(row=row, column=1, value=label).font = (BOLD if label
                                                        else BASE)
        if fn is None:
            continue
        values = [fn(r) for r in runs]
        best = None
        if fmt == "#,##0" and "Цена за тонну" in label:
            best = values.index(min(values))
        if "Фактически войдет" in label or "Использовано" in label:
            best = values.index(max(values))
        if "Остается во Владивостоке" in label or "Итого за всю" in label:
            best = values.index(min(values))
        for i, v in enumerate(values, start=2):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if fmt:
                cell.number_format = fmt
            if best is not None:
                cell.fill = WIN_FILL if i - 2 == best else LOSE_FILL
                if i - 2 == best:
                    cell.font = BOLD

    row += 2
    ws.cell(row=row, column=1, value="Для справки: обычный 20-футовый"
            ).font = SUB
    row += 1
    for label, value, fmt in [
        ("Цена за контейнер, руб", PLAIN20["price"], "#,##0"),
        ("Паспортная загрузка, т", PLAIN20["kg"] / 1000, "0.0"),
        ("Фактически войдет, т", min(runs[0]["kg"], PLAIN20["kg"]) / 1000,
         "0.00"),
        ("Цена за тонну, руб",
         PLAIN20["price"] / (min(runs[0]["kg"], PLAIN20["kg"]) / 1000),
         "#,##0"),
    ]:
        ws.cell(row=row, column=1, value=label).font = BASE
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = fmt
        row += 1
    cell = ws.cell(row=row, column=1,
                   value="Размер тот же, 33 куб. м. Усиленный отличается "
                         "только грузоподъемностью, а она здесь не "
                         "работает: объем кончается раньше веса.")
    cell.font = NOTE
    cell.alignment = Alignment(wrap_text=True)

    row += 3
    ws.cell(row=row, column=1, value="Что войдет по группам").font = SUB
    row += 1
    ws.cell(row=row, column=1, value="Группа")
    for i, r in enumerate(runs, start=2):
        ws.cell(row=row, column=i, value=f"{r['short']}, штук")
    ws.cell(row=row, column=2 + len(runs), value="Всего на складе, штук")
    head(ws, row, 2 + len(runs), height=30)

    totals = defaultdict(float)
    per_run = [defaultdict(float) for _ in runs]
    for i, r in enumerate(runs):
        for item in r["items"]:
            per_run[i][group_of(item["Наименование"])] += item["Штук"]
    # весь склад считаем один раз: пул у обоих сценариев одинаковый
    for item in runs[0]["items"] + runs[0]["rest"]:
        totals[group_of(item["Наименование"])] += item["Штук"]

    order = [g for g in GROUPS if g in totals]
    order += [g for g in sorted(totals) if g not in GROUPS]
    for g in order:
        row += 1
        ws.cell(row=row, column=1, value=g).font = BASE
        for i, pr in enumerate(per_run, start=2):
            cell = ws.cell(row=row, column=i, value=pr[g] or None)
            cell.font = BASE
            cell.border = BORDER
            cell.number_format = "#,##0"
            if pr[g] < totals[g]:
                cell.fill = WARN_FILL
        cell = ws.cell(row=row, column=2 + len(runs), value=totals[g])
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"
    row += 1
    ws.cell(row=row, column=1, value="ИТОГО").font = BOLD
    for i in range(2, 3 + len(runs)):
        letter = get_column_letter(i)
        cell = ws.cell(row=row, column=i,
                       value=f"=SUM({letter}{row - len(order)}:"
                             f"{letter}{row - 1})")
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"
    note = ws.cell(row=row + 1, column=1,
                   value="Оранжевым помечено там, где группа уезжает "
                         "не полностью.")
    note.font = NOTE

    row += 3
    ws.cell(row=row, column=1, value="Выводы").font = SUB
    row += 1
    a20, a40 = runs
    for text in [
        f"Вся арматура — {total_kg / 1000:.1f} т. Двадцатифутовый увезет "
        f"{a20['kg'] / 1000:.1f} т, сорокафутовый {a40['kg'] / 1000:.1f} т.",
        f"Двадцатифутовый до своих {a20['kg_limit'] / 1000:.0f} тонн не "
        f"догрузится: объем 33 куб. м кончается на "
        f"{a20['kg'] / 1000:.1f} тоннах, это "
        f"{a20['kg'] / a20['kg_limit'] * 100:.0f}% грузоподъемности. "
        "Платить за усиленный контейнер не за что.",
        f"Сорокафутовый упирается в вес: {a40['kg'] / 1000:.1f} т при "
        f"{a40['m3_used'] / a40['m3_limit'] * 100:.0f}% занятого объема. "
        "Там еще половина кузова свободна.",
        f"Разница в цене {(a40['price'] - a20['price']) / 1000:.0f} тысяч "
        f"за {(a40['kg'] - a20['kg']) / 1000:.1f} тонны сверху — это "
        f"{(a40['price'] - a20['price']) / ((a40['kg'] - a20['kg']) / 1000):,.0f}"
        " руб за дополнительную тонну.".replace(",", " "),
        "Транспортная компания берет 81 000 руб за тонну. Второй "
        "двадцатифутовый на остаток обойдется дешевле.",
        "Объем посчитан, а не измерен. Если упаковка выйдет рыхлее на "
        "15%, двадцатифутовый возьмет не 21,7 т, а 19. У сорокафутового "
        "такого риска нет — там запас по объему двукратный.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1

    widths(ws, [38, 22, 22, 22])
    return ws


def main():
    runs = []
    total_kg = 0
    for opt in OPTIONS:
        r = run(opt)
        # в ящики идет только то, что лежит верхними ярусами
        boxed = sum(i["вес"] for i in r["items"]
                    if tier(asrow(i)) in (3, 4))
        r["boxes"] = int(-(-boxed // BOX_KG))
        r["boxed_kg"] = boxed
        runs.append(r)
        total_kg = max(total_kg, r["kg"] + r["rest_kg"])

    wb = Workbook()
    wb.remove(wb.active)
    sheet_compare(wb, runs, total_kg)
    for r in runs:
        sheet_items(wb, f"{r['short']} — едет",
                    f"{r['name']}: что входит", r["items"],
                    f"{r['positions']} позиций, {r['units']:.0f} штук, "
                    f"{r['kg'] / 1000:.2f} т, {r['m3_used']:.1f} куб. м")
        sheet_items(wb, f"{r['short']} — остается",
                    f"{r['name']}: что остается во Владивостоке", r["rest"],
                    f"{len(r['rest'])} позиций, {r['rest_units']:.0f} штук, "
                    f"{r['rest_kg'] / 1000:.2f} т")

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Сравнение_контейнеров.xlsx"
    wb.save(path)

    print(f"готово: {path}\n")
    print(f"вся арматура: {total_kg / 1000:.2f} т\n")
    for r in runs:
        print(f"{r['name']}: {r['price']:,} руб".replace(",", " "))
        print(f"   войдет {r['kg'] / 1000:>5.2f} т "
              f"({r['kg'] / r['kg_limit'] * 100:>3.0f}% веса), "
              f"{r['m3_used']:>4.1f} куб.м "
              f"({r['m3_used'] / r['m3_limit'] * 100:>3.0f}% объема), "
              f"упор в {r['limited']}")
        print(f"   {r['positions']} позиций, {r['units']:.0f} штук, "
              f"{r['boxes']} ящиков")
        print(f"   останется {r['rest_kg'] / 1000:.2f} т "
              f"({len(r['rest'])} позиций), "
              f"{r['per_ton']:,.0f} руб за тонну\n".replace(",", " "))


if __name__ == "__main__":
    main()
