#!/usr/bin/env python3
"""Бланк отгрузки: что отправили, новое или б/у, почем и сколько
доставки легло на штуку.

Пустая форма на 300 строк. Вписываем наименование, единицу измерения,
состояние, количество, цену за штуку и доставку за штуку — остальное
считается: цена с доставкой, сумма, доставка на всю позицию и итог по
строке.

Две первые строки заполнены как пример, их стираем перед работой.

Результат: outputs/marine_equipment/Отгрузка_образец.xlsx
"""

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
EXAMPLE = Font(name=FONT, size=10, italic=True, color="808080")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
CALC_FILL = PatternFill("solid", fgColor="EAF3EA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHEET = "Отгрузка"
ROWS = 300          # столько пустых строк готовим
START = 8           # первая строка данных
HEADER = START - 1  # строка заголовков таблицы

UNITS = ["шт", "компл", "кг", "м", "пара", "набор"]
CONDITIONS = ["Новое", "Б/у"]

# колонка: (заголовок, ширина, вводим руками)
COLUMNS = [
    ("№", 5, False),
    ("Наименование товара", 52, True),
    ("Ед. изм.", 10, True),
    ("Новое или б/у", 14, True),
    ("Количество", 11, True),
    ("Цена за шт, руб", 14, True),
    ("Стоимость доставки за шт, руб", 16, True),
    ("Цена с доставкой за шт, руб", 16, False),
    ("Сумма, руб", 14, False),
    ("Доставка всего, руб", 14, False),
    ("Итого по строке, руб", 15, False),
    ("№ места", 10, True),
    ("Примечание", 30, True),
]

N, NAME, UNIT, COND, QTY, PRICE, SHIP, FULL, SUM_, SHIPSUM, TOTAL, PLACE, \
    REM = [get_column_letter(i) for i in range(1, len(COLUMNS) + 1)]

EXAMPLES = [
    ("Задвижка клинкетная Ду250", "шт", "Б/у", 2, 150000, 14300, "12",
     "фланцы закрыты фанерой"),
    ("Клапан проходной штуцерный Ду20 Ру40", "шт", "Новое", 40, 2000, 65,
     "3", "ящик с перегородками"),
]

HOWTO = [
    ("Состояние ставим при погрузке",
     "Новое или б/у видно только тогда, когда позицию берут в руки и "
     "кладут в ящик. Потом по памяти это не восстановить."),
    ("Доставка за штуку считается от веса",
     "За контейнер платят за тонны. Разделите все расходы на отгрузку "
     "на общий вес груза — получится цена килограмма, умножьте на вес "
     "штуки. Блок для этого счета стоит вверху листа отгрузки."),
    ("Поровну на штуку разносить нельзя",
     "Задвижка Ду250 весит 220 кг, штуцерный клапан — 1 кг. Если "
     "разделить доставку поровну, мелочь подорожает в разы, а тяжелая "
     "арматура уедет дешевле себестоимости."),
    ("Цена с доставкой — это не цена продажи",
     "Это то, во что позиция обошлась на складе у покупателя. Цену "
     "ставим от нее, иначе тяжелое уедет в минус."),
    ("Новое и б/у не смешиваем в одном месте",
     "Покупатель принимает груз по упаковочному листу. Если в ящике "
     "вперемешку, спор о состоянии придет на весь ящик."),
    ("Номер места — как на бирке",
     "Номер в этой таблице должен совпадать с биркой на ящике и со "
     "строкой упаковочного листа."),
]


def head(ws, row, ncols, height=42):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def sheet_form(wb):
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = "Отгрузка: что отправили и почем это обошлось"
    ws["A1"].font = TITLE
    ws["A2"] = ("Желтое вписываем сами, зеленое считается. "
                "Две первые строки — пример, сотрите их перед работой.")
    ws["A2"].font = NOTE

    # шапка документа
    fields = [("Покупатель", "Дата отгрузки"),
              ("Куда везем", "Номер контейнера")]
    for i, (left, right) in enumerate(fields):
        row = 4 + i
        ws.cell(row=row, column=1, value=left).font = BOLD
        cell = ws.cell(row=row, column=2)
        cell.fill = INPUT_FILL
        cell.border = BORDER
        ws.cell(row=row, column=6, value=right).font = BOLD
        cell = ws.cell(row=row, column=7)
        cell.fill = INPUT_FILL
        cell.border = BORDER

    # счет доставки на килограмм
    ws.cell(row=4, column=9, value="Все расходы на доставку, руб").font = BOLD
    cell = ws.cell(row=4, column=11)
    cell.fill = INPUT_FILL
    cell.border = BORDER
    cell.number_format = "#,##0"
    ws.cell(row=5, column=9, value="Общий вес груза, кг").font = BOLD
    cell = ws.cell(row=5, column=11)
    cell.fill = INPUT_FILL
    cell.border = BORDER
    cell.number_format = "#,##0"
    ws.cell(row=6, column=9,
            value="Доставка за килограмм, руб").font = BOLD
    cell = ws.cell(row=6, column=11, value="=IFERROR(K4/K5,0)")
    cell.fill = CALC_FILL
    cell.border = BORDER
    cell.font = BOLD
    cell.number_format = "#,##0.00"
    cell = ws.cell(row=6, column=12,
                   value="умножьте на вес одной штуки — получится "
                         "доставка за штуку")
    cell.font = NOTE

    for i, (name, _, _) in enumerate(COLUMNS, start=1):
        ws.cell(row=HEADER, column=i, value=name)
    head(ws, HEADER, len(COLUMNS))

    last = START + ROWS - 1
    for n in range(1, ROWS + 1):
        row = START + n - 1
        example = EXAMPLES[n - 1] if n <= len(EXAMPLES) else None
        values = [
            n,
            example[0] if example else None,
            example[1] if example else None,
            example[2] if example else None,
            example[3] if example else None,
            example[4] if example else None,
            example[5] if example else None,
            f'=IF({PRICE}{row}="","",{PRICE}{row}+{SHIP}{row})',
            f'=IF({PRICE}{row}="","",{QTY}{row}*{PRICE}{row})',
            f'=IF({SHIP}{row}="","",{QTY}{row}*{SHIP}{row})',
            f'=IF({PRICE}{row}="","",{SUM_}{row}+{SHIPSUM}{row})',
            example[6] if example else None,
            example[7] if example else None,
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = EXAMPLE if example else BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(i in (2, 13)))
        for letter in (PRICE, SHIP, FULL, SUM_, SHIPSUM, TOTAL):
            ws[f"{letter}{row}"].number_format = "#,##0"
        for i, (_, _, manual) in enumerate(COLUMNS, start=1):
            if i == 1:
                continue
            letter = get_column_letter(i)
            ws[f"{letter}{row}"].fill = INPUT_FILL if manual else CALC_FILL

    dv_unit = DataValidation(type="list", formula1='"' + ",".join(UNITS) + '"',
                             allow_blank=True)
    dv_cond = DataValidation(type="list",
                             formula1='"' + ",".join(CONDITIONS) + '"',
                             allow_blank=True)
    ws.add_data_validation(dv_unit)
    ws.add_data_validation(dv_cond)
    dv_unit.add(f"{UNIT}{START}:{UNIT}{last}")
    dv_cond.add(f"{COND}{START}:{COND}{last}")

    total = last + 1
    ws.cell(row=total, column=2, value="ИТОГО").font = BOLD
    for letter in (QTY, SUM_, SHIPSUM, TOTAL):
        cell = ws[f"{letter}{total}"]
        cell.value = f"=SUM({letter}{START}:{letter}{last})"
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"
    cell = ws.cell(row=total, column=len(COLUMNS),
                   value="строк заполнено: "
                         f"=COUNTA({NAME}{START}:{NAME}{last})")
    cell.value = f'=\"строк заполнено: \"&COUNTA({NAME}{START}:{NAME}{last})'
    cell.font = NOTE

    ws.auto_filter.ref = f"A{HEADER}:{get_column_letter(len(COLUMNS))}{last}"
    ws.freeze_panes = f"C{START}"
    for i, (_, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return last


def sheet_totals(wb, last):
    ws = wb.create_sheet("Итоги")
    ws["A1"] = "Новое и б/у: сколько чего отгружено"
    ws["A1"].font = TITLE
    ws["A2"] = ("Считается по колонке «Новое или б/у». Строка «не "
                "указано» показывает, что еще не заполнили.")
    ws["A2"].font = NOTE

    def rng(letter):
        return f"'{SHEET}'!${letter}${START}:${letter}${last}"

    row = 4
    for i, h in enumerate(["Состояние", "Строк", "Количество", "Сумма, руб",
                           "Доставка, руб", "Итого с доставкой, руб"],
                          start=1):
        ws.cell(row=row, column=i, value=h)
    head(ws, row, 6, height=30)

    for name in CONDITIONS + ["не указано"]:
        row += 1
        if name == "не указано":
            count = (f'=COUNTIFS({rng(NAME)},"<>",{rng(COND)},"")')
            crit = '""'
        else:
            count = f'=COUNTIF({rng(COND)},"{name}")'
            crit = f'"{name}"'
        values = [
            name,
            count,
            f"=SUMIF({rng(COND)},{crit},{rng(QTY)})",
            f"=SUMIF({rng(COND)},{crit},{rng(SUM_)})",
            f"=SUMIF({rng(COND)},{crit},{rng(SHIPSUM)})",
            f"=D{row}+E{row}",
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i >= 3:
                cell.number_format = "#,##0"

    row += 1
    ws.cell(row=row, column=1, value="ИТОГО").font = BOLD
    for col in range(2, 7):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=SUM({letter}{row - 3}:{letter}{row - 1})")
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"

    row += 2
    ws.cell(row=row, column=1, value="Доля доставки в цене").font = TITLE
    row += 1
    for name, formula in [
        ("Товар без доставки, руб", f"=SUM({rng(SUM_)})"),
        ("Доставка, руб", f"=SUM({rng(SHIPSUM)})"),
        ("Доставка в процентах от товара",
         f"=IFERROR(SUM({rng(SHIPSUM)})/SUM({rng(SUM_)}),0)"),
        ("Всего с доставкой, руб",
         f"=SUM({rng(SUM_)})+SUM({rng(SHIPSUM)})"),
    ]:
        ws.cell(row=row, column=1, value=name).font = BASE
        cell = ws.cell(row=row, column=2, value=formula)
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = ("0.0%" if "процентах" in name else "#,##0")
        row += 1

    for i, w in enumerate([34, 14, 14, 18, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def sheet_howto(wb):
    ws = wb.create_sheet("Как заполнять")
    ws["A1"] = "Как вести бланк"
    ws["A1"].font = TITLE
    for i, h in enumerate(["№", "Правило", "Почему так"], start=1):
        ws.cell(row=3, column=i, value=h)
    head(ws, 3, 3, height=22)
    for n, (rule, why) in enumerate(HOWTO, start=1):
        row = 3 + n
        for i, v in enumerate([n, rule, why], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i > 1))
        ws.row_dimensions[row].height = 50
    for i, w in enumerate([5, 40, 86], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def main():
    wb = Workbook()
    last = sheet_form(wb)
    sheet_totals(wb, last)
    sheet_howto(wb)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Отгрузка_образец.xlsx"
    wb.save(path)
    print(f"готово: {path}")
    print(f"строк под заполнение: {ROWS}, "
          f"колонок: {len(COLUMNS)}, пример в первых {len(EXAMPLES)} строках")


if __name__ == "__main__":
    main()
