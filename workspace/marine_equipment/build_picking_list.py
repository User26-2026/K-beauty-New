#!/usr/bin/env python3
"""Перечень арматуры для отгрузки: рабочий бланк для склада.

Менеджер отдает этот файл бригаде, бригада идет по складу и вписывает
фактическое количество и недостающие чертежи. После этого у нас первый
раз появляется достоверная картина: сейчас остатки сверялись 16.06.2025
и уже дважды не сошлись на отгрузках.

Порядок — как грузим: сначала вся арматура от самой крупной к мелкой,
внутри группы по убыванию диаметра, и только в конце вентиляционные
головки. Головки по диаметру крупные, но это не арматура: они легкие,
объемные и в контейнер идут последними. Бригада идет сверху вниз и в
том же порядке кладет в контейнер.

Ничего лишнего: наименование, диаметр, давление, чертеж, количество и
вес. Цены, локации и расчеты в этот файл не идут — он для рук, а не для
переговоров.

Результат: outputs/marine_equipment/Перечень_для_отгрузки.xlsx
"""

import pathlib
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_container_spb import plural  # noqa: E402
from build_logistics_compare import rows_all  # noqa: E402
from build_spb_list_xlsx import group_of  # noqa: E402
from build_weight_xlsx import du as du_of  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")   # что заполняет склад
NODRAW = PatternFill("solid", fgColor="FFFF00")    # чертежа нет
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="1F3864")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Группы, которые идут после всей арматуры, в порядке перечисления.
TAIL = ["Головки вентиляционные"]

COLS = [
    ("№", 5),
    ("Наименование", 54),
    ("Ду", 8),
    ("Ру", 8),
    ("Чертеж", 22),
    ("Кол-во по учету, шт", 12),
    ("Вес 1 шт, кг", 11),
    ("Вес всего, кг", 12),
    ("ФАКТ, шт", 12),
    ("ЧЕРТЕЖ дописать", 24),
]


def ru_of(r):
    """Давление: из списка менеджера или из названия."""
    if r.get("Ру"):
        return r["Ру"]
    m = re.search(r"ру\s*(\d{1,3}(?:[.,]\d)?)", r["Наименование"], re.I)
    return m.group(1) if m else ""


def head(ws, row, ncols, height=36):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def load():
    rows = []
    for r in rows_all():
        rows.append({
            "Наименование": r["Наименование"],
            "Ду": int(r.get("Ду") or du_of(r["Наименование"]) or 0),
            "Ру": ru_of(r),
            "Чертеж": r.get("Чертеж") or "",
            "Штук": r["Штук"],
            "вес_ед": r["вес_ед"],
            "вес": r["вес"],
            "Группа": group_of(r["Наименование"]),
        })
    return rows


def group_order(by_group):
    """Сначала арматура по убыванию диаметра, в конце вентиляционные головки.

    Головки вентиляционные по диаметру дотягивают до Ду200 и в общем
    списке вставали четвертыми, между фильтрами и фланцевой угловой.
    Арматура это не они: головка легкая и объемная, грузится последней.
    """
    def key(g):
        rows = by_group[g]
        return (TAIL.index(g) if g in TAIL else -1,
                -max((r["Ду"] for r in rows), default=0),
                -sum(r["вес"] for r in rows))
    return sorted(by_group, key=key)


def main():
    rows = load()
    by_group = defaultdict(list)
    for r in rows:
        by_group[r["Группа"]].append(r)

    wb = Workbook()
    ws = wb.active
    ws.title = "Перечень для отгрузки"
    ws["A1"] = "Перечень арматуры для отгрузки в Петербург"
    ws["A1"].font = TITLE
    ws["A2"] = (f"{len(rows)} "
                f"{plural(len(rows), 'позиция', 'позиции', 'позиций')}, "
                f"{sum(r['Штук'] for r in rows):,.0f} шт, "
                f"{sum(r['вес'] for r in rows) / 1000:.2f} т. Порядок как "
                "грузим: сверху самая крупная арматура, дальше по "
                "убыванию диаметра, в самом конце вентиляционные "
                "головки.".replace(",", " "))
    ws["A2"].font = NOTE
    ws["A3"] = ("Желтые колонки заполняет склад: сколько нашлось по факту "
                "и номер чертежа там, где его нет. Желтая ячейка в колонке "
                "«Чертеж» — чертеж отсутствует, его и надо вписать.")
    ws["A3"].font = NOTE

    start = 5
    for i, (h, _) in enumerate(COLS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLS))

    row = start
    n = 0
    nodraw = 0
    for g in group_order(by_group):
        items = sorted(by_group[g], key=lambda r: (-r["Ду"],
                                                   r["Наименование"]))
        row += 1
        qty = sum(r["Штук"] for r in items)
        kg = sum(r["вес"] for r in items)
        ws.cell(row=row, column=1,
                value=f"{g} — {len(items)} "
                      f"{plural(len(items), 'позиция', 'позиции', 'позиций')}"
                      f", {qty:,.0f} шт, {kg:,.0f} кг".replace(",", " "))
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = GROUP_FILL
            cell.font = SUB
            cell.border = BORDER
        ws.row_dimensions[row].height = 22

        for r in items:
            n += 1
            row += 1
            values = [n, r["Наименование"], r["Ду"] or "", r["Ру"] or "",
                      r["Чертеж"], r["Штук"], round(r["вес_ед"], 1),
                      round(r["вес"]), None, None]
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=i, value=v)
                cell.font = BASE
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top",
                                           wrap_text=(i == 2),
                                           horizontal=("center"
                                                       if i in (1, 3, 4, 6)
                                                       else None))
                if i in (6, 8):
                    cell.number_format = "#,##0"
                if i == 7:
                    cell.number_format = "#,##0.0"
            ws.cell(row=row, column=6).font = BOLD
            for c in (9, 10):
                ws.cell(row=row, column=c).fill = FILL_IN
            if not r["Чертеж"]:
                ws.cell(row=row, column=5).fill = NODRAW
                nodraw += 1
            ws.row_dimensions[row].height = 26

    last = row
    row += 1
    ws.cell(row=row, column=2, value="ИТОГО").font = BOLD
    for col in (6, 8):
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col,
                       value=f"=SUM({letter}{start + 1}:{letter}{last})")
        cell.font = BOLD
        cell.border = BORDER
        cell.number_format = "#,##0"
    cell = ws.cell(row=row, column=9,
                   value=f"=SUM(I{start + 1}:I{last})")
    cell.font = BOLD
    cell.border = BORDER
    cell.fill = FILL_IN
    cell.number_format = "#,##0"

    row += 2
    ws.cell(row=row, column=1, value="Как заполнять").font = SUB
    row += 1
    for text in [
        "Идти сверху вниз: сначала клинкеты и крупное, потом мелочь, "
        "в конце вентиляционные головки. В этом же порядке потом "
        "грузим контейнер.",
        "В колонку ФАКТ вписать, сколько реально нашлось. Если не "
        "нашлось совсем — ставить ноль, а не оставлять пусто: пустая "
        "ячейка значит «не смотрели».",
        f"Желтым подсвечены {nodraw} позиций без чертежа. Чертеж выбит "
        "на корпусе или на бирке, его надо переписать в последнюю "
        "колонку. Без чертежа позицию трудно продать и невозможно "
        "сверить с прайсом.",
        "Если нашлось больше, чем в учете, это нормально и важно: "
        "учет сверялся 16.06.2025 и уже дважды не сошелся. Пишем факт.",
        "Если позиция в разобранном виде, без штока, без клиньев или "
        "некомплект — пометить прямо в строке рядом с количеством.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1

    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(COLS))}{last}"
    ws.freeze_panes = f"C{start + 1}"
    ws.print_title_rows = f"{start}:{start}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    widths(ws, [w for _, w in COLS])

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Перечень_для_отгрузки.xlsx"
    wb.save(path)

    print(f"готово: {path}\n")
    print(f"{len(rows)} позиций, {sum(r['Штук'] for r in rows):,.0f} шт, "
          f"{sum(r['вес'] for r in rows) / 1000:.2f} т, "
          f"без чертежа {nodraw}".replace(",", " "))
    print(f"\n{'группа':<36}{'поз':>5}{'шт':>7}{'кг':>9}{'макс Ду':>9}")
    for g in group_order(by_group):
        items = by_group[g]
        print(f"{g:<36}{len(items):>5}{sum(r['Штук'] for r in items):>7.0f}"
              f"{sum(r['вес'] for r in items):>9.0f}"
              f"{max((r['Ду'] for r in items), default=0):>9}")


if __name__ == "__main__":
    main()
