#!/usr/bin/env python3
"""Журнал фактических отгрузок: что уехало, почем и что это значит.

Здесь только факт с экспедиторских расписок: вес, места, деньги за
перевозку, объявленная ценность. Отсюда берем настоящую доставку на
единицу — вместо оценки — и настоящий вес, чтобы поправить справочник.

Заодно показываем расхождения с учетом: по обеим отгруженным позициям
остатки не сошлись.

Результат: outputs/marine_equipment/Отгрузки_факт.xlsx
"""

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_marine_registry import SOURCES, parse  # noqa: E402
from build_weight_xlsx import unit_weight as ref_weight  # noqa: E402
from shipments_done import (  # noqa: E402
    SHIPMENTS, cost, per_kg, per_unit, unit_weight, units,
)

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
CALC_FILL = PatternFill("solid", fgColor="EAF3EA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def head(ws, row, ncols, height=36):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, values, fmt=None, font=None):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font or BASE
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(i in (2, 3)))
        if fmt and i in fmt:
            cell.number_format = fmt[i]
    return row + 1


def stock_of(name, location):
    """Сколько числится в учете по этой позиции."""
    for spec in SOURCES:
        for r in parse(spec):
            if r["Наименование"] == name and r["Локация"] == location:
                return r["Наличие"] or 0, r["Цена за ед., руб"] or 0
    return None, 0


def sheet_items(wb):
    ws = wb.create_sheet("Отгружено")
    ws["A1"] = "Что уже уехало"
    ws["A1"].font = TITLE
    ws["A2"] = ("Цифры с экспедиторских расписок: вес брутто вместе с "
                "тарой, деньги — вся сумма за отгрузку.")
    ws["A2"].font = NOTE

    cols = ["Дата", "Куда", "Покупатель", "Наименование", "Ед.",
            "Состояние", "Кол-во", "Мест", "Вес всего, кг", "Вес ед., кг",
            "Доставка всего, руб", "Доставка на ед., руб", "Руб за кг",
            "Объявленная ценность, руб", "Документ"]
    for i, h in enumerate(cols, start=1):
        ws.cell(row=4, column=i, value=h)
    head(ws, 4, len(cols))

    fmt = {9: "#,##0", 10: "#,##0.0", 11: "#,##0", 12: "#,##0",
           13: "#,##0.0", 14: "#,##0"}
    row = 5
    for s in SHIPMENTS:
        for i in s["items"]:
            row = put(ws, row, [
                s["date"], s["city"], s["buyer"], i["name"], i["unit"],
                i["condition"] or "не указано", i["qty"], s["places"],
                s["kg"], unit_weight(s), cost(s), per_unit(s), per_kg(s),
                s["declared"], s["doc"],
            ], fmt)
            for col in (10, 12, 13):
                ws.cell(row=row - 1, column=col).fill = CALC_FILL

    row = put(ws, row, ["ИТОГО", None, None, None, None, None,
                        sum(units(s) for s in SHIPMENTS), None,
                        sum(s["kg"] for s in SHIPMENTS), None,
                        sum(cost(s) for s in SHIPMENTS), None, None,
                        sum(s["declared"] for s in SHIPMENTS), None],
              fmt, BOLD)

    widths(ws, [12, 18, 30, 34, 7, 12, 9, 7, 13, 12, 16, 16, 11, 18, 16])
    ws.freeze_panes = "D5"
    return ws


def sheet_costs(wb):
    ws = wb.create_sheet("Расходы по распискам")
    ws["A1"] = "Из чего сложилась доставка"
    ws["A1"].font = TITLE
    ws["A2"] = ("Перевозка — это только часть счета. Организация "
                "перевозки, упаковка и страхование дают больше.")
    ws["A2"].font = NOTE

    cols = ["Документ", "Куда", "Услуга", "Сумма, руб", "Доля"]
    for i, h in enumerate(cols, start=1):
        ws.cell(row=4, column=i, value=h)
    head(ws, 4, len(cols), height=24)

    row = 5
    for s in SHIPMENTS:
        total = cost(s)
        for name, value in s["services"]:
            row = put(ws, row, [s["doc"], s["city"], name, value,
                                value / total if total else 0],
                      {4: "#,##0", 5: "0.0%"})
        row = put(ws, row, [s["doc"], s["city"], "Итого по расписке",
                            total, 1], {4: "#,##0", 5: "0.0%"}, BOLD)
        row += 1

    widths(ws, [18, 20, 40, 16, 10])
    return ws


def sheet_weights(wb):
    ws = wb.create_sheet("Вес факт и справочник")
    ws["A1"] = "Насколько справочный вес расходится с фактическим"
    ws["A1"].font = TITLE
    ws["A2"] = ("Весов в учете нет, поэтому загрузка контейнера считалась "
                "по справочным массам. Вот первые две проверки по "
                "реальным отгрузкам.")
    ws["A2"].font = NOTE

    cols = ["Наименование", "Штук", "Вес ед. по расписке, кг",
            "Вес ед. по справочнику, кг", "Расхождение", "Вывод"]
    for i, h in enumerate(cols, start=1):
        ws.cell(row=4, column=i, value=h)
    head(ws, 4, len(cols))

    row = 5
    for s in SHIPMENTS:
        for i in s["items"]:
            fact = unit_weight(s)
            ref = ref_weight(i["registry"])
            diff = (fact - ref) / ref if ref else 0
            verdict = ("справочник завышает вес" if diff < -0.15
                       else "справочник занижает вес" if diff > 0.15
                       else "сходится")
            row = put(ws, row, [i["name"], i["qty"], fact, ref, diff,
                                verdict],
                      {3: "#,##0.0", 4: "#,##0.0", 5: "0%"})
            if abs(diff) > 0.15:
                ws.cell(row=row - 1, column=5).fill = WARN_FILL

    row += 1
    for text in [
        "Вес по расписке — брутто, вместе с обрешеткой и поддоном, то "
        "есть сам товар легче.",
        "Двух замеров мало, чтобы менять справочник целиком. Но если "
        "задвижки действительно вдвое легче расчета, в контейнер войдет "
        "заметно больше, чем показывает план загрузки.",
        "Что делать: при следующих отгрузках записывать вес каждого "
        "места и подставлять его в колонку фактического веса файла "
        "Вес_арматуры.xlsx.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = NOTE
        cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 26
        row += 1

    widths(ws, [36, 9, 20, 22, 14, 34])
    return ws


def sheet_gaps(wb):
    ws = wb.create_sheet("Расхождения с учетом")
    ws["A1"] = "Учет против факта отгрузки"
    ws["A1"].font = TITLE
    ws["A2"] = ("Обе отгруженные позиции с учетом не сошлись. Это та же "
                "причина, по которой до инвентаризации нельзя называть "
                "покупателю количества.")
    ws["A2"].font = NOTE

    cols = ["Наименование", "Локация", "Числится в учете",
            "Отгружено по факту", "Расхождение", "Что это значит"]
    for i, h in enumerate(cols, start=1):
        ws.cell(row=4, column=i, value=h)
    head(ws, 4, len(cols))

    row = 5
    for s in SHIPMENTS:
        for i in s["items"]:
            stock, _ = stock_of(i["registry"], i["location"])
            stock = 0 if stock is None else stock
            diff = stock - i["qty"]
            if diff < 0:
                meaning = ("отгрузили больше, чем числилось: товар был, "
                           "в учете его не было")
            elif diff > 0:
                meaning = "остаток после отгрузки"
            else:
                meaning = "сошлось в ноль"
            row = put(ws, row, [i["registry"], i["location"], stock,
                                i["qty"], diff, meaning],
                      {3: "#,##0", 4: "#,##0", 5: "#,##0"})
            if diff < 0:
                ws.cell(row=row - 1, column=5).fill = WARN_FILL

    row += 1
    for text in [
        "Клинкеты Ду150 числились в нуле, а отгружено 45 новых штук.",
        "Фильтров забортной воды Ду300 числился один, уехало два.",
        "Пока инвентаризация не сделана, любые количества в КП и "
        "объявлениях — предположение, а не факт.",
    ]:
        cell = ws.cell(row=row, column=1, value="— " + text)
        cell.font = NOTE
        cell.alignment = Alignment(wrap_text=True)
        row += 1

    widths(ws, [36, 14, 16, 18, 14, 46])
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_items(wb)
    sheet_costs(wb)
    sheet_weights(wb)
    sheet_gaps(wb)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Отгрузки_факт.xlsx"
    wb.save(path)

    print(f"готово: {path}")
    for s in SHIPMENTS:
        print(f"  {s['date']} {s['city']:<16} {units(s):>3} шт, "
              f"{s['kg']:>5} кг, {cost(s):>8,.0f} руб  "
              f"= {per_unit(s):>9,.0f} руб/шт, {per_kg(s):>5.1f} руб/кг"
              .replace(",", " "))
    print(f"  всего доставки: "
          f"{sum(cost(s) for s in SHIPMENTS):,.0f} руб".replace(",", " "))


if __name__ == "__main__":
    main()
