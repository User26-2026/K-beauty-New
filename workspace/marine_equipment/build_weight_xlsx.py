#!/usr/bin/env python3
"""Оценка веса арматуры.

Весов в учете нет ни по одной позиции, поэтому вес считается по
справочным массам судовой арматуры: масса зависит от типа и диаметра.
Это оценка с погрешностью около трети в обе стороны, а не факт.
Чтобы получить настоящую цифру, надо взвесить 5-10 характерных позиций
разных диаметров и подставить их в колонку фактического веса.

Вес нужен для трех вещей: загрузки контейнера, расчета доставки
покупателю и цены лома, если часть пойдет в металл.

Результат: outputs/marine_equipment/Вес_арматуры.xlsx
"""

import pathlib
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_marine_registry import SOURCES, parse  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Справочная масса одной единицы, кг. Первый ключ — тип арматуры,
# внутри — диаметр условного прохода.
WEIGHTS = {
    "клапан штуцерный": {3: 0.3, 6: 0.4, 10: 0.6, 15: 1.0, 20: 1.5, 25: 2.5,
                         32: 3.5, 50: 7, 65: 10},
    "клапан фланцевый": {20: 4, 25: 5, 32: 7, 40: 10, 50: 14, 65: 20, 70: 23,
                         80: 28, 100: 40, 125: 55, 150: 75, 200: 130},
    "клинкет/задвижка": {50: 12, 60: 16, 65: 18, 70: 22, 80: 28, 100: 45,
                         125: 60, 150: 85, 200: 140, 250: 220, 300: 300,
                         350: 400},
    "захлопка": {50: 12, 100: 35},
    "кингстон": {50: 20, 80: 35, 100: 45},
    "коробка клапанная": {50: 25, 80: 45, 100: 60},
    "кран": {10: 0.8, 15: 1.2, 20: 1.8, 32: 4},
}
# если диаметр не указан, берем среднее по типу
DEFAULTS = {"клапан штуцерный": 2, "клапан фланцевый": 20,
            "клинкет/задвижка": 40, "захлопка": 15, "кингстон": 35,
            "коробка клапанная": 40, "кран": 2, "прочая арматура": 5}

COLUMNS = ["№", "Тип", "Диаметр", "Наименование", "Штук",
           "Справочный вес ед., кг", "Вес позиции, кг",
           "Фактический вес ед., кг", "Вес по факту, кг", "Локация"]


def du(name):
    m = re.search(r"[дД][уУ]\s*(\d{1,3})", name)
    return int(m.group(1)) if m else None


def kind(name):
    low = name.lower()
    if re.search(r"клинкет|задвижк", low):
        return "клинкет/задвижка"
    if re.search(r"штуцерн|штущерн|муфтов|цапков|под дюрит|педальн|манометр",
                 low):
        return "клапан штуцерный"
    if "коробка" in low:
        return "коробка клапанная"
    if "захлопк" in low:
        return "захлопка"
    if "кингстон" in low:
        return "кингстон"
    if low.startswith("кран"):
        return "кран"
    if "фланц" in low:
        return "клапан фланцевый"
    return "прочая арматура"


def unit_weight(name):
    """Справочная масса единицы: по типу и ближайшему диаметру."""
    k = kind(name)
    table = WEIGHTS.get(k)
    d = du(name)
    if table and d:
        if d in table:
            return table[d]
        nearest = min(table, key=lambda x: abs(x - d))
        return table[nearest]
    return DEFAULTS.get(k, 5)


def head(ws, row, ncols, height=32):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_items(wb, rows):
    ws = wb.create_sheet("Расчет по позициям")
    ws["A1"] = "Вес арматуры: расчет по позициям"
    ws["A1"].font = TITLE
    ws["A2"] = ("Справочный вес проставлен по типу и диаметру. Желтые колонки "
                "заполняем после взвешивания, тогда вес пересчитается сам.")
    ws["A2"].font = NOTE

    start = 4
    for i, h in enumerate(COLUMNS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLUMNS))

    for n, r in enumerate(rows, start=1):
        row = start + n
        w = unit_weight(r["Наименование"])
        d = du(r["Наименование"])
        values = [n, kind(r["Наименование"]), f"Ду{d}" if d else "не указан",
                  r["Наименование"], r["Наличие"], w,
                  f"=E{row}*F{row}", None,
                  f"=IF(H{row}=\"\",\"\",E{row}*H{row})", r["Локация"]]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 4))
            if i in (6, 7, 8, 9):
                cell.number_format = "#,##0.0"
        for col in (8,):
            ws.cell(row=row, column=col).fill = INPUT_FILL

    last = start + len(rows)
    total = last + 1
    ws.cell(row=total, column=4, value="ИТОГО").font = Font(name=FONT, bold=True)
    for col in (5, 7, 9):
        letter = get_column_letter(col)
        cell = ws.cell(row=total, column=col,
                       value=f"=SUM({letter}{start + 1}:{letter}{last})")
        cell.font = Font(name=FONT, bold=True)
        cell.border = BORDER
        cell.number_format = "#,##0" if col == 5 else "#,##0.0"
    ws.cell(row=total + 1, column=4, value="Итого тонн (справочно)").font = BASE
    cell = ws.cell(row=total + 1, column=7, value=f"=G{total}/1000")
    cell.number_format = "#,##0.0"
    cell.font = Font(name=FONT, bold=True)

    ws.auto_filter.ref = f"A{start}:J{last}"
    ws.freeze_panes = f"D{start + 1}"
    widths(ws, [5, 20, 12, 52, 8, 16, 15, 16, 15, 16])
    return ws


def sheet_summary(wb, rows):
    ws = wb.create_sheet("Сводка по весу", 0)
    ws["A1"] = "Сколько весит арматура"
    ws["A1"].font = TITLE
    ws["A2"] = ("Весов в учете нет. Цифры получены по справочным массам "
                "судовой арматуры, погрешность около трети в обе стороны. "
                "Чтобы знать точно, надо взвесить 5-10 позиций разных "
                "диаметров.")
    ws["A2"].font = NOTE

    total_kg = sum(r["Наличие"] * unit_weight(r["Наименование"]) for r in rows)
    facts = [
        ("Позиций арматуры", len(rows), ""),
        ("Штук всего", sum(r["Наличие"] for r in rows), ""),
        ("Вес, тонн (оценка)", round(total_kg / 1000, 1),
         f"разброс примерно {total_kg / 1000 * 0.7:.0f}-"
         f"{total_kg / 1000 * 1.3:.0f} тонн"),
    ]
    for i, (label, value, note) in enumerate(facts):
        row = 4 + i
        ws.cell(row=row, column=1, value=label).font = Font(name=FONT, size=10,
                                                            bold=True)
        ws.cell(row=row, column=2, value=value).font = Font(name=FONT, size=10,
                                                            bold=True)
        ws.cell(row=row, column=3, value=note).font = NOTE

    # по типам
    start = 8
    ws.cell(row=start - 1, column=1, value="По типам арматуры").font = TITLE
    for i, h in enumerate(["Тип", "Позиций", "Штук", "Вес, кг", "Вес, тонн"],
                          start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, 5, height=24)

    agg = defaultdict(lambda: [0, 0, 0.0])
    for r in rows:
        a = agg[kind(r["Наименование"])]
        a[0] += 1
        a[1] += r["Наличие"]
        a[2] += r["Наличие"] * unit_weight(r["Наименование"])
    row = start
    for k, (p, u, w) in sorted(agg.items(), key=lambda kv: -kv[1][2]):
        row += 1
        for i, v in enumerate([k, p, u, round(w), round(w / 1000, 2)], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i == 4:
                cell.number_format = "#,##0"
            if i == 5:
                cell.number_format = "#,##0.00"

    # по диаметрам
    row += 2
    ws.cell(row=row, column=1, value="По диаметрам").font = TITLE
    row += 1
    head_row = row
    for i, h in enumerate(["Диаметр", "Позиций", "Штук", "Вес, кг",
                           "Вес, тонн"], start=1):
        ws.cell(row=head_row, column=i, value=h)
    head(ws, head_row, 5, height=24)

    agg2 = defaultdict(lambda: [0, 0, 0.0])
    for r in rows:
        a = agg2[du(r["Наименование"]) or 0]
        a[0] += 1
        a[1] += r["Наличие"]
        a[2] += r["Наличие"] * unit_weight(r["Наименование"])
    for key in sorted(agg2):
        row += 1
        p, u, w = agg2[key]
        label = f"Ду{key}" if key else "не указан"
        for i, v in enumerate([label, p, u, round(w), round(w / 1000, 2)],
                              start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            if i == 4:
                cell.number_format = "#,##0"
            if i == 5:
                cell.number_format = "#,##0.00"

    widths(ws, [26, 12, 12, 14, 14, 40])
    return ws


def main():
    rows = []
    for spec in SOURCES:
        rows.extend(parse(spec))
    live = [r for r in rows
            if (r["Наличие"] or 0) > 0 and r["Товарная группа"] == "Арматура"]
    live.sort(key=lambda r: (kind(r["Наименование"]),
                             du(r["Наименование"]) or 0,
                             r["Наименование"]))

    wb = Workbook()
    wb.remove(wb.active)
    sheet_summary(wb, live)
    sheet_items(wb, live)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Вес_арматуры.xlsx"
    wb.save(path)

    total_kg = sum(r["Наличие"] * unit_weight(r["Наименование"]) for r in live)
    print(f"готово: {path}")
    print(f"арматуры: {len(live)} позиций, "
          f"{sum(r['Наличие'] for r in live):.0f} штук")
    print(f"вес по справочным массам: {total_kg:,.0f} кг "
          f"= {total_kg / 1000:.1f} тонн".replace(",", " "))
    print(f"разброс с учетом погрешности: {total_kg / 1000 * 0.7:.0f}-"
          f"{total_kg / 1000 * 1.3:.0f} тонн\n")

    agg = defaultdict(lambda: [0, 0.0])
    for r in live:
        a = agg[kind(r["Наименование"])]
        a[0] += r["Наличие"]
        a[1] += r["Наличие"] * unit_weight(r["Наименование"])
    for k, (u, w) in sorted(agg.items(), key=lambda kv: -kv[1][1]):
        print(f"  {k:<20} {u:>5.0f} шт  {w:>9,.0f} кг  "
              f"{w / 1000:>5.1f} т".replace(",", " "))


if __name__ == "__main__":
    main()
