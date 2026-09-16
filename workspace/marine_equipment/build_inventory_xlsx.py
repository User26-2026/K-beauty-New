#!/usr/bin/env python3
"""Инвентаризационная ведомость по складу судового оборудования.

Учетные файлы последний раз сверялись 16.06.2025. Все, что продано
после, в них не отражено, поэтому остатки недостоверны. Пока не
проведена инвентаризация, нельзя ни публиковать объявления, ни
рассылать КП с количествами, ни называть стоимость склада.

Ведомость печатается и заполняется на складе от руки либо ведется
в ноутбуке: колонка «Факт» заполняется по пересчету, расхождение
и его сумма считаются формулами.

Результат: outputs/marine_equipment/Инвентаризация.xlsx
"""

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_marine_registry import SOURCES, parse  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
LOC_FILL = PatternFill("solid", fgColor="E8EEF7")
TOP_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    "№", "Локация", "Место в файле", "Наименование", "Чертеж",
    "Числится", "Факт", "Расхождение", "Цена за ед., руб",
    "Сумма расхождения", "Кто считал", "Дата", "Комментарий",
]

RULES = [
    ("Считаем по факту, а не по бумаге",
     "В колонку «Факт» пишем то, что реально лежит и посчитано руками. "
     "Если позиции нет — ставим 0, а не оставляем пусто. Пустая ячейка "
     "означает «еще не считали»."),
    ("Начинаем с дорогого",
     "Лист «Считать первыми» — позиции, которые дают 80% стоимости склада. "
     "Их пересчитываем в первую очередь: по ним уже готовятся объявления "
     "и коммерческие предложения."),
    ("Одна позиция — один человек",
     "Кто считал и когда — обязательные колонки. Если потом всплывет "
     "расхождение, будет понятно, к кому идти с вопросом."),
    ("Комплектность отмечаем отдельно",
     "Если позиция на месте, но разукомплектована или в браке — пишем "
     "это в комментарий. Такая позиция стоит дешевле и в КП идет иначе."),
    ("Нашли то, чего нет в ведомости",
     "Дописываем новой строкой внизу: локация, наименование, количество. "
     "Учет неполный, часть товара может быть нигде не записана."),
    ("Не считаем на глаз",
     "Светильники и мелкая арматура лежат коробками и ящиками. "
     "Пересчитываем коробки и содержимое одной коробки, дальше умножаем, "
     "и это отмечаем в комментарии."),
    ("После инвентаризации фиксируем каждую отгрузку",
     "Причина расхождений в том, что продажи не заносились в файлы "
     "15 месяцев. Дальше каждая отгрузка отмечается в тот же день, "
     "иначе через год повторится то же самое."),
]


def head(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fill_sheet(ws, rows, title, subtitle, highlight=False):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = NOTE

    start = 4
    for i, h in enumerate(COLUMNS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLUMNS))

    for n, r in enumerate(rows, start=1):
        row = start + n
        values = [
            n, r["Локация"], r["Группа в файле"] or r["Оборудование"] or "",
            r["Наименование"], r["Чертеж"], r["Наличие"], None,
            f"=IF(G{row}=\"\",\"\",G{row}-F{row})",
            r["Цена за ед., руб"] or None,
            f"=IF(G{row}=\"\",\"\",H{row}*I{row})",
            "", "", "",
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i in (4, 13)))
            if i in (9, 10):
                cell.number_format = "#,##0"
        ws.cell(row=row, column=2).fill = LOC_FILL
        for col in (7, 11, 12, 13):
            ws.cell(row=row, column=col).fill = INPUT_FILL
        if highlight:
            ws.cell(row=row, column=4).fill = TOP_FILL

    last = start + len(rows)
    s = last + 2
    labels = [
        ("Позиций в ведомости", f"=COUNTA(D{start + 1}:D{last})"),
        ("Пересчитано", f"=COUNT(G{start + 1}:G{last})"),
        ("Осталось пересчитать", f"=D{s}-D{s + 1}"),
        ("Позиций с расхождением", f"=COUNTIF(H{start + 1}:H{last},\"<>0\")"
                                   f"-COUNTBLANK(H{start + 1}:H{last})"),
        ("Сумма расхождений, руб", f"=SUM(J{start + 1}:J{last})"),
    ]
    for i, (label, formula) in enumerate(labels):
        ws.cell(row=s + i, column=3, value=label).font = BASE
        cell = ws.cell(row=s + i, column=4, value=formula)
        cell.font = Font(name=FONT, bold=(i >= 3),
                         color="C00000" if i >= 3 else "000000")
        if i == 4:
            cell.number_format = "#,##0"

    ws.auto_filter.ref = f"A{start}:M{last}"
    ws.freeze_panes = f"D{start + 1}"
    widths(ws, [5, 16, 22, 52, 20, 10, 9, 12, 14, 16, 14, 12, 30])


def sheet_rules(wb):
    ws = wb.create_sheet("Как считать", 0)
    ws["A1"] = "Инвентаризация склада: правила"
    ws["A1"].font = TITLE
    ws["A2"] = ("Учетные файлы последний раз сверялись 16.06.2025, прошло "
                "15 месяцев. Все проданное за это время в них не отражено.")
    ws["A2"].font = NOTE

    for i, h in enumerate(["№", "Правило", "Почему"], start=1):
        ws.cell(row=4, column=i, value=h)
    head(ws, 4, 3)
    for n, (rule, why) in enumerate(RULES, start=1):
        row = 4 + n
        for i, v in enumerate([n, rule, why], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i in (2, 3)))
        ws.row_dimensions[row].height = 46
    widths(ws, [5, 40, 95])
    return ws


def main():
    rows = []
    for spec in SOURCES:
        rows.extend(parse(spec))
    for r in rows:
        r["Сумма"] = (r["Наличие"] or 0) * (r["Цена за ед., руб"] or 0)

    live = [r for r in rows if (r["Наличие"] or 0) > 0]
    # порядок обхода склада: по локации, внутри — по разделу учета
    live.sort(key=lambda r: (r["Локация"], r["Группа в файле"] or "",
                             r["Наименование"]))

    # позиции, дающие 80% стоимости, считаем первыми
    by_value = sorted([r for r in live if r["Сумма"] > 0],
                      key=lambda r: -r["Сумма"])
    total = sum(r["Сумма"] for r in by_value)
    top, acc = [], 0
    for r in by_value:
        top.append(r)
        acc += r["Сумма"]
        if acc >= total * 0.8:
            break

    wb = Workbook()
    wb.remove(wb.active)
    sheet_rules(wb)
    fill_sheet(wb.create_sheet("Считать первыми"), top,
               "Считать первыми: 80% стоимости склада",
               "По этим позициям готовятся объявления и КП. "
               "Пока они не пересчитаны, публиковать нельзя.",
               highlight=True)
    fill_sheet(wb.create_sheet("Ведомость полная"), live,
               "Инвентаризационная ведомость: все позиции",
               "Заполняем желтые колонки. Расхождение и сумма считаются сами. "
               "Позиции нет — ставим 0, а не пустую ячейку.")

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Инвентаризация.xlsx"
    wb.save(path)
    print(f"готово: {path}")
    print(f"всего позиций в ведомости: {len(live)}")
    print(f"из них считать первыми: {len(top)} "
          f"на {sum(r['Сумма'] for r in top):,.0f} руб".replace(",", " "))


if __name__ == "__main__":
    main()
