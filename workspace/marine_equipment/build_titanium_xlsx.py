#!/usr/bin/env python3
"""Отдельный перечень титановой арматуры.

Титан вынесен отдельно, потому что стоит дороже остальной арматуры
и как изделие, и как лом. Цен на складе по нему нет ни на одну позицию,
поэтому часть колонок заполняется руками.

Результат: outputs/marine_equipment/Титановые_клапаны.xlsx
"""

import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_marine_registry import SOURCES, parse  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "outputs" / "marine_equipment"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITANIUM = re.compile(r"титан", re.I)

COLUMNS = ["№", "Наименование", "Чертеж", "Диаметр", "Числится",
           "Факт после пересчета", "Цена за ед., руб", "Сумма, руб",
           "Вес ед., кг", "Локация", "Комментарий"]

def head(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def diameter(name):
    m = re.search(r"[дД][уУ]\s*(\d{1,3})", name)
    return f"Ду{m.group(1)}" if m else ""


def main():
    rows = []
    for spec in SOURCES:
        rows.extend(parse(spec))
    items = [r for r in rows
             if (r["Наличие"] or 0) > 0
             and TITANIUM.search(r["Наименование"] + " "
                                 + (r["Группа в файле"] or ""))]
    items.sort(key=lambda r: -(r["Наличие"] or 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Титановая арматура"
    ws["A1"] = "Титановая арматура на складе"
    ws["A2"] = ("Желтые колонки заполняем сами. Цен по титану в учете нет "
                "ни на одну позицию.")
    ws["A1"].font = TITLE
    ws["A2"].font = NOTE

    start = 4
    for i, h in enumerate(COLUMNS, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(COLUMNS))

    for n, r in enumerate(items, start=1):
        row = start + n
        values = [n, r["Наименование"], r["Чертеж"],
                  diameter(r["Наименование"]), r["Наличие"], None, None,
                  f"=IF(OR(F{row}=\"\",G{row}=\"\"),\"\",F{row}*G{row})",
                  None, r["Локация"], ""]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 2))
            if i in (7, 8):
                cell.number_format = "#,##0"
        for col in (6, 7, 9, 11):
            ws.cell(row=row, column=col).fill = INPUT_FILL

    last = start + len(items)
    total = last + 1
    ws.cell(row=total, column=2, value="ИТОГО").font = Font(name=FONT, bold=True)
    for col in (5, 6, 8):
        letter = get_column_letter(col)
        cell = ws.cell(row=total, column=col,
                       value=f"=SUM({letter}{start + 1}:{letter}{last})")
        cell.font = Font(name=FONT, bold=True)
        cell.border = BORDER
        if col == 8:
            cell.number_format = "#,##0"
    ws.cell(row=total + 1, column=2, value="Общий вес, кг").font = BASE
    ws.cell(row=total + 1, column=5,
            value=f"=SUMPRODUCT(F{start + 1}:F{last},I{start + 1}:I{last})"
            ).font = Font(name=FONT, bold=True)

    ws.freeze_panes = f"B{start + 1}"
    widths(ws, [5, 52, 20, 11, 11, 20, 16, 16, 13, 16, 34])

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Титановые_клапаны.xlsx"
    wb.save(path)
    print(f"готово: {path}")
    for r in items:
        print(f"  {r['Наличие']:>3.0f} шт  {r['Наименование']}  "
              f"черт. {r['Чертеж']}  ({r['Локация']})")
    print(f"  всего {len(items)} позиции, "
          f"{sum(r['Наличие'] for r in items):.0f} штук")


if __name__ == "__main__":
    main()
