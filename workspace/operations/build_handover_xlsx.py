#!/usr/bin/env python3
"""График передачи дел — рабочий файл для увольняющейся сотрудницы.

Отдельный файл от общего плана: общий план содержит оценку склада,
стратегию продажи и внутренние решения, его сотруднице не отправляем.
Здесь только то, что она должна сдать, и в какой день.

Результат: outputs/operations/Передача_дел_график.xlsx
"""

import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "operations" / "Передача_дел_график.xlsx"

FONT = "Arial"
START = dt.date(2026, 9, 10)
WEEKDAYS = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F3864")
BASE = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
DAY_FILL = PatternFill("solid", fgColor="E8EEF7")
KEY_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# день (номер рабочего дня), задача, кому передать, важность
SCHEDULE = [
    (1, "Логины и пароли: личные кабинеты ДЭК и Водоканала", "Москва", "критично"),
    (1, "Логины и пароли: рабочая почта, Фарпост, площадки объявлений", "Москва", "критично"),
    (1, "Лицевые счета по всем четырем объектам", "Москва", "критично"),
    (1, "ОСТАТКИ: начать список того, что отгружено с 16.06.2025", "Москва", "критично"),

    (2, "ОСТАТКИ: список отгрузок за 15 месяцев — что, сколько, кому, когда", "Москва", "критично"),
    (2, "Открытые заказы по косметике: что оплачено и еще не отгружено", "Москва", "критично"),
    (2, "Заказы в пути: что едет и когда придет", "Москва", "критично"),

    (3, "ОСТАТКИ: пересчет 79 позиций, дающих 80% стоимости, часть 1", "Москва", "критично"),
    (3, "Опись ключей: от каких помещений, у кого дубликаты", "Москва", "критично"),
    (3, "Подотчетные деньги: остаток и за что платили наличными", "Москва", "критично"),

    (4, "ОСТАТКИ: пересчет 79 позиций, часть 2", "Москва", "критично"),
    (4, "Порядок подачи показаний по Светланской: куда и в какие числа", "Москва", "критично"),
    (4, "Лицевой счет по Аксаковской и ссылка", "Москва", "критично"),
    (4, "Договоры и сканы: энергоснабжение, водоснабжение, аренда", "Москва", "критично"),

    (5, "ОСТАТКИ: свести пересчет и список отгрузок, объяснить расхождения", "Москва", "критично"),
    (5, "Сверка расчетов с поставщиками: кому должны, кто должен нам", "Москва", "критично"),

    (6, "Заполнить реестр контактов: все, с кем общались за год", "Москва", "критично"),
    (6, "Отдельно: мастер по ремонту арматуры, сварщик, электрик", "Москва", "важно"),
    (6, "Отдельно: воровайка, кран, вывоз мусора, приемка металла", "Москва", "важно"),
    (6, "Отдельно: таможенный брокер и перевозчики по контейнерам", "Москва", "критично"),

    (7, "Бренды и заводы по косметике: контакты, цены, скидки, минимальная партия", "Москва", "критично"),
    (7, "Байер, консолидатор, склад в Корее, логистика", "Москва", "критично"),
    (7, "Документы: CPNP, декларации, Честный знак", "Москва", "важно"),
    (7, "Журнал обзвона покупателей со статусами по каждой компании", "Олеся", "критично"),
    (7, "Кому отправлены КП и фото, кто что просил, кто в переговорах", "Олеся", "критично"),

    (8, "Обход Калинина вместе с новым сотрудником: счетчики, ключи, склад", "новый сотрудник", "критично"),
    (8, "Показать новому, где что лежит по выверенным остаткам", "новый сотрудник", "критично"),
    (8, "Знакомство со сторожами, порядок приема смены и оплаты", "новый сотрудник", "критично"),

    (9, "Обход Светланской, Фокино, Аксаковской: где счетчики", "новый сотрудник", "критично"),
    (9, "Новый сотрудник сам снимает показания, вы проверяете", "новый сотрудник", "критично"),
    (9, "Знакомство с подрядчиками и с экономистом ДЭК", "новый сотрудник", "важно"),
    (9, "Совместные звонки покупателям, представить Олесю", "Олеся", "важно"),

    (10, "Новый сотрудник сам оплачивает воду по четырем объектам", "новый сотрудник", "критично"),
    (10, "Подписать акт сверки остатков: что числится, что есть по факту", "Москва", "критично"),
    (10, "Акт приема-передачи, передача ключей", "Москва", "критично"),
]

INSTRUCTIONS = [
    "Как ведется учет отгрузок: где отмечать, что продано и кому",
    "Как снять показания на каждом объекте: где стоит счетчик, что записывать",
    "Куда и в какие числа передавать показания, по каждому объекту отдельно",
    "Как оплатить воду и свет: с какого счета, где брать квитанции",
    "Как принять смену сторожа и как ему платят",
    "Кому звонить при аварии на каждом объекте",
    "Порядок работы с ДЭК и Водоканалом, если в счете расхождение",
    "Как размещается заказ на косметику: от заявки до отгрузки",
    "Как проходит растаможка контейнера: кто что делает и в какие сроки",
]

CHECKLIST = [
    ("Остатки", "Список отгрузок с 16.06.2025: что, сколько, кому, когда"),
    ("Остатки", "Пересчитаны 79 позиций, дающих 80% стоимости"),
    ("Остатки", "Расхождения объяснены: продано, списано, не найдено"),
    ("Остатки", "Акт сверки остатков подписан"),
    ("Доступы", "Личные кабинеты ДЭК и Водоканала"),
    ("Доступы", "Рабочая почта"),
    ("Доступы", "Фарпост и другие площадки объявлений"),
    ("Доступы", "Лицевые счета по четырем объектам"),
    ("Документы", "Договоры энергоснабжения и водоснабжения"),
    ("Документы", "Договоры аренды и документы на объекты"),
    ("Документы", "Последние показания и оплаты"),
    ("Деньги", "Подотчетные средства: остаток"),
    ("Деньги", "Сверка расчетов с поставщиками косметики"),
    ("Деньги", "Кому мы должны и кто должен нам"),
    ("Ключи", "Ключи от помещений, гаража, контейнеров, будки"),
    ("Ключи", "Список: у кого еще есть ключи и дубликаты"),
    ("Контакты", "Реестр контактов заполнен полностью"),
    ("Контакты", "Мастер по ремонту арматуры"),
    ("Контакты", "Подрядчики: воровайка, кран, мусор, металл, сварщик, электрик"),
    ("Контакты", "Сторожа: график и порядок оплаты"),
    ("Контакты", "Таможенный брокер и перевозчики"),
    ("Косметика", "Открытые заказы: оплачено и не отгружено"),
    ("Косметика", "Заказы в пути"),
    ("Косметика", "Бренды и заводы: условия, цены, минимальные партии"),
    ("Косметика", "Байер, консолидатор, склад в Корее"),
    ("Косметика", "Сертификация: CPNP, декларации, Честный знак"),
    ("Оборудование", "Журнал обзвона со статусами"),
    ("Оборудование", "Кому отправлены КП, кто в переговорах"),
    ("Оборудование", "Фотоархив"),
    ("Инструкции", "Написаны все инструкции со второго листа"),
]


def workdays(n):
    """Список рабочих дат, начиная со START."""
    out, d = [], START
    while len(out) < n:
        if d.weekday() < 5:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def head(ws, row, ncols, height=28):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_schedule(wb):
    ws = wb.create_sheet("График по дням")
    ws["A1"] = "Передача дел: график на две недели"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Последний рабочий день — 23 сентября. Главная задача — привести "
                "остатки в порядок: только вы знаете, что отгружено с июня 2025. "
                "Новому человеку передаем уже выверенный склад.")
    ws["A2"].font = NOTE

    days = workdays(10)
    headers = ["День", "Дата", "Что передать", "Кому", "Важность",
               "Сделано", "Комментарий"]
    start = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(headers))

    row = start
    for n, (day, task, who, prio) in enumerate(SCHEDULE, start=1):
        row += 1
        d = days[day - 1]
        date_text = f"{d.strftime('%d.%m')} {WEEKDAYS[d.weekday()]}"
        for i, v in enumerate([day, date_text, task, who, prio, "Нет", ""], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i in (3, 7)))
            if i <= 2:
                cell.fill = DAY_FILL
        if prio == "критично":
            c = ws.cell(row=row, column=5)
            c.fill = KEY_FILL
            c.font = Font(name=FONT, size=10, bold=True, color="C00000")
        ws.cell(row=row, column=6).fill = INPUT_FILL
        ws.cell(row=row, column=7).fill = INPUT_FILL

    last = row
    dv = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F{start + 1}:F{last}")

    s = last + 2
    ws.cell(row=s, column=3, value="Всего пунктов").font = BASE
    ws.cell(row=s, column=4, value=f"=COUNTA(C{start + 1}:C{last})").font = BASE
    ws.cell(row=s + 1, column=3, value="Сделано").font = BASE
    ws.cell(row=s + 1, column=4,
            value=f'=COUNTIF(F{start + 1}:F{last},"Да")').font = BASE
    ws.cell(row=s + 2, column=3, value="Осталось критичных").font = BASE
    ws.cell(row=s + 2, column=4,
            value=f'=COUNTIFS(E{start + 1}:E{last},"критично",'
                  f'F{start + 1}:F{last},"Нет")').font = Font(name=FONT, bold=True,
                                                              color="C00000")
    ws.cell(row=s + 3, column=3, value="Готовность").font = BASE
    cell = ws.cell(row=s + 3, column=4, value=f"=IFERROR(D{s + 1}/D{s},0)")
    cell.number_format = "0%"
    cell.font = Font(name=FONT, bold=True)

    ws.auto_filter.ref = f"A{start}:G{last}"
    ws.freeze_panes = f"A{start + 1}"
    widths(ws, [7, 12, 66, 20, 13, 11, 34])
    return ws


def sheet_instructions(wb):
    ws = wb.create_sheet("Написать инструкции")
    ws["A1"] = "Инструкции, которые нужно написать"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Устная передача на расстоянии не работает. По каждому пункту "
                "нужна короткая инструкция на страницу: что делать по шагам.")
    ws["A2"].font = NOTE

    headers = ["№", "Про что инструкция", "Написана", "Где лежит"]
    start = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(headers))

    for n, text in enumerate(INSTRUCTIONS, start=1):
        row = start + n
        for i, v in enumerate([n, text, "Нет", ""], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 2))
        ws.cell(row=row, column=3).fill = INPUT_FILL
        ws.cell(row=row, column=4).fill = INPUT_FILL

    last = start + len(INSTRUCTIONS)
    dv = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{start + 1}:C{last}")

    ws.cell(row=last + 2, column=2, value="Написано инструкций").font = BASE
    ws.cell(row=last + 2, column=3,
            value=f'=COUNTIF(C{start + 1}:C{last},"Да")').font = BASE
    widths(ws, [5, 72, 12, 34])
    return ws


def sheet_checklist(wb):
    ws = wb.create_sheet("Чек-лист приемки")
    ws["A1"] = "Что должно быть передано к 23 сентября"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Итоговая сверка перед подписанием акта приема-передачи."
    ws["A2"].font = NOTE

    headers = ["№", "Блок", "Что передано", "Передано", "Принял"]
    start = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    head(ws, start, len(headers))

    for n, (block, item) in enumerate(CHECKLIST, start=1):
        row = start + n
        for i, v in enumerate([n, block, item, "Нет", ""], start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BASE
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 3))
        ws.cell(row=row, column=4).fill = INPUT_FILL
        ws.cell(row=row, column=5).fill = INPUT_FILL

    last = start + len(CHECKLIST)
    dv = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D{start + 1}:D{last}")

    ws.cell(row=last + 2, column=3, value="Передано пунктов").font = BASE
    ws.cell(row=last + 2, column=4,
            value=f'=COUNTIF(D{start + 1}:D{last},"Да")').font = BASE
    ws.cell(row=last + 3, column=3, value="Всего пунктов").font = BASE
    ws.cell(row=last + 3, column=4,
            value=f"=COUNTA(C{start + 1}:C{last})").font = BASE

    ws.auto_filter.ref = f"A{start}:E{last}"
    widths(ws, [5, 18, 62, 12, 22])
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_schedule(wb)
    sheet_instructions(wb)
    sheet_checklist(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    days = workdays(10)
    print(f"готово: {OUT}")
    print(f"рабочих дней: {days[0].strftime('%d.%m')} - {days[-1].strftime('%d.%m')}")
    print(f"пунктов в графике: {len(SCHEDULE)}, "
          f"инструкций: {len(INSTRUCTIONS)}, в чек-листе: {len(CHECKLIST)}")


if __name__ == "__main__":
    main()
