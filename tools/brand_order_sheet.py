"""Order sheet бренда: что почем и с какого завода.

Бренд отдает прайс собственной формой: штрихкод, название, объем, штук в
коробе, розница, оптовая цена и — главное — завод-изготовитель по каждой
позиции. Позиции одного бренда делают разные заводы, поэтому заказ
собирается не одной отгрузкой, а по производствам: у каждого свой короб,
свой вес и свой срок.

Рядом ставим цены посредников по тому же штрихкоду: разница с ценой
бренда и есть их наценка.

Запуск:
    python3 tools/brand_order_sheet.py data/price_lists/medipeel/<файл>.xlsx
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rates import IMPORT_MULTIPLIER, KRW_RUB

PRICES_TABLE = "outputs/prices_normalized.xlsx"
# Цена втрое ниже цены бренда — это не скидка, а другая фасовка: маски
# бренд считает пачкой по 10, а посредник продает штуками.
PRICE_LIMIT = 3
TARGET = "outputs/Order sheet по заводам.xlsx"

HEAD = PatternFill("solid", fgColor="1F3864")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
WARN = PatternFill("solid", fgColor="FFC7CE")
WHITE = Font(color="FFFFFF", bold=True)

# Как подписаны нужные колонки в форме бренда.
FIELDS = {
    "Штрихкод": [r"bar\s*code", r"штрихкод"],
    "Код": [r"product\s*code"],
    # Английское название стоит второй колонкой под общей шапкой
    # Product Name, первая — корейская.
    "Название": [r"english", r"product\s*name"],
    "Объем": [r"^volume$", r"^size$"],
    "Штук в коробе": [r"pcs\s*/\s*ct", r"pcs\s*per", r"q'?ty\s*/?\s*box"],
    "Розница, KRW": [r"retail\s*price"],
    "Оптовая, KRW": [r"wholesale\s*price", r"supply\s*price"],
    "Вес короба, кг": [r"weight"],
    "Завод": [r"manufacturer"],
    "Адрес завода": [r"address"],
}


def norm(value):
    """Текст ячейки шапки. Пустая ячейка приходит как NaN, а не как None."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def first_item_row(rows):
    """Первая строка товара: та, где штрихкод это число."""
    for index, row in enumerate(rows):
        if row and pd.notna(pd.to_numeric(pd.Series([row[0]]), errors="coerce")[0]):
            return index
    return 0


def find_columns(rows):
    """Номера нужных колонок.

    Шапка занимает несколько строк: подпись колонки может стоять и в первой
    строке, и во второй, поэтому все строки до первого товара склеиваем по
    колонкам. Шаблоны поля пробуем по очереди и слева направо: в конце
    формы стоят пустые колонки заказа с теми же словами (ORDER, WEIGHT), и
    брать надо первую подходящую, а не последнюю.
    """
    header_rows = rows[:first_item_row(rows)] or rows[:2]
    width = max(len(row) for row in header_rows)
    joined = [" ".join(norm(row[index]) for row in header_rows
                       if index < len(row) and norm(row[index]))
              for index in range(width)]

    mapping, used = {}, set()
    for field, patterns in FIELDS.items():
        for pattern in patterns:
            spot = next((index for index, title in enumerate(joined)
                         if title and index not in used
                         and re.search(pattern, title)), None)
            if spot is not None:
                mapping[field], _ = spot, used.add(spot)
                break
    return mapping


def read_sheet(path):
    table = pd.read_excel(path, sheet_name=0, header=None)
    rows = table.values.tolist()
    mapping = find_columns(rows)
    missing = [field for field in ("Штрихкод", "Оптовая, KRW", "Завод")
               if field not in mapping]
    if missing:
        raise SystemExit(f"В файле не нашлись колонки: {', '.join(missing)}")

    data = pd.DataFrame({field: table[index] for field, index in mapping.items()})
    # Строка товара — та, где штрихкод это число: шапка и подписи отпадают.
    data = data[pd.to_numeric(data["Штрихкод"], errors="coerce").notna()]
    data["Штрихкод"] = (pd.to_numeric(data["Штрихкод"], errors="coerce")
                        .astype("Int64").astype(str))
    for column in ("Розница, KRW", "Оптовая, KRW", "Штук в коробе", "Вес короба, кг"):
        if column in data:
            data[column] = pd.to_numeric(data[column], errors="coerce")
    data["Завод"] = data["Завод"].fillna("не указан").astype(str).str.strip()
    data["Цена, руб"] = (data["Оптовая, KRW"] * KRW_RUB).round()
    # Себестоимость на нашем складе: закупка плюс доставка, пошлина и приемка.
    data["Себестоимость, руб"] = (data["Цена, руб"] * IMPORT_MULTIPLIER).round()
    data["Наценка бренда к рознице, раз"] = (
        data["Розница, KRW"] / data["Оптовая, KRW"]).round(2)
    # Короб — реальная минимальная единица заказа, штуками завод не отгружает.
    data["Короб, руб"] = (data["Цена, руб"] * data["Штук в коробе"]).round()
    return data.reset_index(drop=True)


def suppliers(brand):
    """Цены посредников по тому же бренду, по штрихкоду."""
    if not os.path.exists(PRICES_TABLE):
        return pd.DataFrame()
    table = pd.read_excel(PRICES_TABLE, dtype={"Штрихкод": str})
    mark = re.sub(r"[^A-Z0-9]", "", brand.upper())
    table = table[table["Бренд"].astype(str).str.upper()
                  .str.replace(r"[^A-Z0-9]", "", regex=True) == mark]
    table = table[table["Поставщик"] != f"{brand} (бренд)"]
    return table[table["Цена за штуку, руб"].notna()]


def write(book, title, columns, rows, widths, money=(), scale=None, freeze="A2"):
    ws = book.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in money:
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    if scale and ws.max_row > 1:
        ws.conditional_formatting.add(
            f"{scale}2:{scale}{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="C6EFCE",
                           end_type="max", end_color="FFC7CE"))
    ws.freeze_panes = freeze
    return ws


def by_factory(book, data):
    rows = []
    for factory, part in data.groupby("Завод"):
        boxes = part["Короб, руб"].sum()
        rows.append([
            factory, str(part["Адрес завода"].iloc[0])[:70] if "Адрес завода" in part else "",
            len(part), int(part["Штук в коробе"].sum())
            if part["Штук в коробе"].notna().all() else None,
            int(boxes) if pd.notna(boxes) else None,
            round(part["Вес короба, кг"].sum(), 1)
            if part["Вес короба, кг"].notna().any() else None,
            f"{part['Название'].iloc[0][:44]}"])
    rows.sort(key=lambda row: -row[2])
    total = [
        "ВСЕГО", "", len(data), int(data["Штук в коробе"].sum()),
        int(data["Короб, руб"].sum()), round(data["Вес короба, кг"].sum(), 1), ""]
    ws = write(book, "ПО ЗАВОДАМ",
               ["Завод", "Адрес", "Позиций", "Штук в коробах",
                "Заказ по коробу на позицию, руб", "Вес коробов, кг",
                "Первая позиция"],
               rows + [total], [40, 60, 10, 16, 22, 16, 46], money="CDEF")
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = TOTAL, Font(bold=True)
    return ws


def positions(book, data):
    rows = []
    for _, item in data.sort_values(["Завод", "Оптовая, KRW"],
                                    ascending=[True, False]).iterrows():
        rows.append([
            item["Штрихкод"], item.get("Код"), item["Название"], item.get("Объем"),
            item["Штук в коробе"], item["Оптовая, KRW"], item["Цена, руб"],
            item["Себестоимость, руб"], item["Розница, KRW"],
            item["Наценка бренда к рознице, раз"], item["Короб, руб"],
            item["Вес короба, кг"], item["Завод"]])
    return write(book, "ПОЗИЦИИ",
                 ["Штрихкод", "Код", "Название", "Объем", "Штук в коробе",
                  "Оптовая, KRW", "Цена, руб", "Себестоимость у нас, руб",
                  "Розница, KRW", "Розница/опт, раз", "Короб, руб",
                  "Вес короба, кг", "Завод"],
                 rows, [15, 13, 56, 10, 12, 13, 12, 16, 13, 13, 13, 12, 34],
                 money="EFGHIK", freeze="C2")


def markup(book, data, prices):
    """Сколько посредники накидывают к цене бренда."""
    if prices.empty:
        return None
    rows, summary, mixed = [], {}, []
    for _, item in data.iterrows():
        same = prices[prices["Штрихкод"] == item["Штрихкод"]]
        for _, offer in same.iterrows():
            over = offer["Цена за штуку, руб"] / item["Цена, руб"] - 1
            line = [offer["Поставщик"], offer["Страна"], item["Название"],
                    item["Цена, руб"], round(offer["Цена за штуку, руб"]),
                    round(offer["Цена за штуку, руб"] - item["Цена, руб"]),
                    round(over * 100, 1)]
            if offer["Цена за штуку, руб"] * PRICE_LIMIT < item["Цена, руб"]:
                mixed.append(line + [item.get("Объем")])
                continue
            rows.append(line)
            summary.setdefault(offer["Поставщик"], []).append(over * 100)
    rows.sort(key=lambda row: (row[0], -row[6]))
    write(book, "НАЦЕНКА ПОСРЕДНИКОВ",
          ["Поставщик", "Страна", "Название", "Цена бренда, руб",
           "Цена посредника, руб", "Разница, руб", "Наценка, %"],
          rows, [16, 9, 56, 15, 18, 14, 12], money="DEF", scale="G", freeze="C2")

    if mixed:
        mixed.sort(key=lambda row: (row[0], row[2]))
        write(book, "РАЗНАЯ ФАСОВКА",
              ["Поставщик", "Страна", "Название", "Цена бренда, руб",
               "Цена посредника, руб", "Разница, руб", "Разница, %", "Объем"],
              mixed, [16, 9, 56, 15, 18, 14, 12, 12], money="DEF", freeze="C2")

    order = sorted(summary.items(), key=lambda pair: pd.Series(pair[1]).median())
    write(book, "ИТОГО ПО ПОСТАВЩИКАМ",
          ["Поставщик", "Сверено позиций", "Медиана наценки, %",
           "Минимум, %", "Максимум, %"],
          [[name, len(values), round(pd.Series(values).median(), 1),
            round(min(values), 1), round(max(values), 1)]
           for name, values in order],
          [18, 16, 20, 12, 12], scale="C")
    return summary


def main(source, brand, target):
    data = read_sheet(source)
    book = Workbook()
    book.remove(book.active)
    by_factory(book, data)
    positions(book, data)
    summary = markup(book, data, suppliers(brand))
    os.makedirs("outputs", exist_ok=True)
    book.save(target)

    print(f"Позиций: {len(data)}   заводов: {data['Завод'].nunique()}")
    print(f"Заказ по одному коробу на позицию: "
          f"{int(data['Короб, руб'].sum()):,} руб, "
          f"{int(data['Штук в коробе'].sum()):,} штук, "
          f"{round(data['Вес короба, кг'].sum()):,} кг".replace(",", " "))
    top = data.groupby("Завод").size().sort_values(ascending=False).head(5)
    print("Крупнейшие заводы:")
    for factory, count in top.items():
        print(f"   {factory[:44]:<46}{count:>4} поз.")
    if summary:
        print("Наценка к цене бренда, медиана:")
        for name, values in sorted(summary.items(),
                                   key=lambda pair: pd.Series(pair[1]).median()):
            print(f"   {name:<16}{pd.Series(values).median():>6.1f}%   "
                  f"позиций {len(values)}")
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Order sheet бренда по заводам")
    parser.add_argument("source")
    parser.add_argument("--brand", default="MEDI-PEEL")
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main(args.source, args.brand, args.out)
