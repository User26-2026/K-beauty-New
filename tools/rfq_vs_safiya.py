"""Какие позиции заявки в Корею есть у SAFIYA.

Заявку мы отправляем в Корею, но часть тех же товаров лежит на складе в
Екатеринбурге. Смысл сверки простой: что можно взять сразу и не ждать
поставку, и во сколько это обойдется против своего ввоза.

Сводим по штрихкоду, а где коды у поставщиков разошлись — по названию
внутри одного бренда. Названия у SAFIYA русские с латинским хвостом
("СЫВОРОТКА ДЛЯ ЛИЦА - VT COSMETICS ..."), поэтому для сверки берем то,
что идет после тире.

Цены SAFIYA переводим по их курсу, корейскую цену — по рыночному плюс
наши расходы на ввоз.

Запуск:
    python3 tools/rfq_vs_safiya.py "outputs/Заявки в Корею.xlsx"
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import name_match
from rates import KRW_RUB, SAFIYA_USD_RUB

ROOT = "data/price_lists/safiya"
B2B = f"{ROOT}/safiya_2026-09-01_b2b.xlsx"
PERSONAL = f"{ROOT}/archive/safiya_2026-08-12_znakomomu.xlsx"
PRICES = "outputs/prices_normalized.xlsx"
TARGET = "outputs/Заявки — что есть у SAFIYA.xlsx"

HEAD = PatternFill("solid", fgColor="1F3864")
YES = PatternFill("solid", fgColor="E2EFDA")
TOTAL = PatternFill("solid", fgColor="FCE4D6")
WHITE = Font(color="FFFFFF", bold=True)

COLUMNS = ["№", "Бренд", "Штрихкод", "Товар", "Количество, шт", "Есть у SAFIYA",
           "Наименование SAFIYA", "Объем", "Цена B2B, $", "Цена B2B, руб",
           "Спец. цена, $", "Сумма по заявке, руб", "Корея EXW, руб",
           "Своим ввозом, руб", "Дороже своего ввоза, %"]


def codes(column):
    """Штрихкоды в выгрузках лежат числом, и .0 на конце портит сверку."""
    numbers = pd.to_numeric(column, errors="coerce")
    return numbers.astype("Int64").astype(str).str.replace("<NA>", "", regex=False)


def key(brand):
    return re.sub(r"[^A-ZА-Я0-9]", "", str(brand).upper())


def same_brand(left, right):
    """Короткое написание считаем за полное: VT и VT COSMETICS."""
    one, two = key(left), key(right)
    if not one or not two:
        return False
    short, long = sorted((one, two), key=len)
    return len(short) >= 2 and long.startswith(short)


# Слова, которые задают вид товара. Если они расходятся, это разные
# товары, как бы ни совпадало остальное: маска против сыворотки, крем для
# глаз против обычного крема.
KIND = {"MASK", "PAD", "PADS", "EYE", "SET", "KIT", "SERUM", "CREAM", "TONER",
        "CLEANSER", "FOAM", "AMPOULE", "ESSENCE", "OIL", "BALM", "STICK",
        "PATCH", "MIST", "GEL", "SUNSCREEN", "SUN", "SHAMPOO", "LOTION",
        "EMULSION", "2STEP", "BOOSTER", "PLUMPER", "PEELING"}


def kind(name):
    return {word for word in re.findall(r"[A-Z0-9]+", str(name).upper()) if word in KIND}


def read_b2b():
    """B2B-прайс: бренд стоит заголовком секции, товары — строками ниже."""
    table = pd.read_excel(B2B, header=0)
    table.columns = ["Код", "Наименование", "Объем", "Цена, $", "Кол-во", "Сумма"]
    goods, brand = [], ""
    for _, row in table.iterrows():
        code = str(row["Код"]).strip()
        if code.isdigit() and len(code) >= 12:
            goods.append({"Бренд": brand, "Штрихкод": code, "Объем": row["Объем"],
                          "Наименование": str(row["Наименование"]).strip(),
                          "Цена B2B, $": float(row["Цена, $"])})
        elif code and code != "nan" and pd.isna(row["Цена, $"]):
            brand = code
    return pd.DataFrame(goods)


def read_personal():
    table = pd.read_excel(PERSONAL, skiprows=4, header=None)
    table.columns = ["Бренд", "Категория", "Наименование", "Штрихкод",
                     "Объем", "Остаток, шт", "Спец. цена, $"]
    table = table.dropna(subset=["Спец. цена, $"])
    table["Штрихкод"] = codes(table["Штрихкод"])
    # Русское название с латинским хвостом: для сверки нужен хвост.
    table["Наименование"] = table["Наименование"].astype(str).str.split(" - ").str[-1]
    return table[["Бренд", "Штрихкод", "Наименование", "Объем", "Спец. цена, $"]]


def safiya():
    """Каталог SAFIYA: действующая цена и цена спецпредложения рядом."""
    b2b, personal = read_b2b(), read_personal()
    catalog = b2b.merge(personal[["Штрихкод", "Спец. цена, $"]], on="Штрихкод", how="outer")
    # Позиции, которых в действующем прайсе нет, берем из спецпредложения.
    missing = catalog["Наименование"].isna()
    filler = personal.set_index("Штрихкод")
    for field in ("Бренд", "Наименование", "Объем"):
        catalog.loc[missing, field] = catalog.loc[missing, "Штрихкод"].map(filler[field])
    return catalog.dropna(subset=["Наименование"]).reset_index(drop=True)


def korea():
    """Самая дешевая корейская цена по штрихкоду, в рублях EXW."""
    table = pd.read_excel(PRICES, dtype={"Штрихкод": str})
    table = table[table["Страна"] == "KR"].copy()
    table["Штрихкод"] = (table["Штрихкод"].fillna("").astype(str)
                         .str.replace(r"\D", "", regex=True))
    table = table[table["Штрихкод"].str.len() >= 8]
    table = table[table["Закупка, KRW"].notna() & (table["Закупка, KRW"] > 0)]
    table["Корея EXW, руб"] = (table["Закупка, KRW"] * KRW_RUB).round()
    best = table.sort_values("Корея EXW, руб").drop_duplicates("Штрихкод")
    return best.set_index("Штрихкод")["Корея EXW, руб"]


def match_rest(request, catalog, found):
    """Позиции без совпадения по коду сводим по названию внутри бренда.

    Строки каталога, уже занятые сверкой по штрихкоду, второй раз не
    предлагаем: иначе один товар SAFIYA встанет против двух наших.
    """
    left = request[found.isna() & request["Бренд"].notna()]
    free = catalog.drop(index=[i for i in found.dropna()], errors="ignore")
    pairs = {}
    for brand in left["Бренд"].unique():
        part = left[left["Бренд"] == brand]
        side = free[[same_brand(brand, other) for other in free["Бренд"]]]
        if part.empty or side.empty:
            continue
        for index, other in name_match.match(part["Товар"], side["Наименование"]).items():
            if kind(part.loc[index, "Товар"]) == kind(side.loc[other, "Наименование"]):
                pairs[index] = other
    return pairs


def compare(request, catalog, prices, import_cost):
    request = request.copy()
    request["Штрихкод"] = codes(request["Штрихкод"])
    by_code = catalog.set_index("Штрихкод")
    found = request["Штрихкод"].map(
        pd.Series(catalog.index.values, index=catalog["Штрихкод"]))
    source = pd.Series(["по штрихкоду" if pd.notna(v) else None for v in found],
                       index=request.index)

    for index, other in match_rest(request, catalog, found).items():
        found[index] = other
        source[index] = "по названию"

    picked = found.map(lambda i: catalog.loc[i] if pd.notna(i) else None)
    request["Есть у SAFIYA"] = source.fillna("нет")
    for field, column in (("Наименование", "Наименование SAFIYA"), ("Объем", "Объем"),
                          ("Цена B2B, $", "Цена B2B, $"), ("Спец. цена, $", "Спец. цена, $")):
        request[column] = [row[field] if row is not None else None for row in picked]
    request["Цена B2B, руб"] = (request["Цена B2B, $"] * SAFIYA_USD_RUB).round()
    request["Сумма по заявке, руб"] = (request["Цена B2B, руб"]
                                       * request["Количество, шт"]).round()
    request["Корея EXW, руб"] = request["Штрихкод"].map(prices)
    request["Своим ввозом, руб"] = (request["Корея EXW, руб"]
                                    * (1 + import_cost / 100)).round()
    request["Дороже своего ввоза, %"] = (
        request["Цена B2B, руб"] / request["Своим ввозом, руб"] * 100 - 100).round(1)
    # Разрыв в разы — не наценка, а разная фасовка: пачка масок против
    # одного саше. Такие позиции считать нельзя, их надо уточнять.
    request.loc[request["Дороже своего ввоза, %"] > 150, "Есть у SAFIYA"] += ", сверить фасовку"
    return request


def write(ws, columns, rows, widths, total=None):
    ws.append(columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
        if len(row) > 5 and row[5] not in (None, "нет"):
            for cell in ws[ws.max_row]:
                cell.fill = YES
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return ws


def money(ws, letters, fmt="# ##0"):
    for letter in letters:
        for cell in ws[letter][1:]:
            cell.number_format = fmt


def main(source, import_cost):
    catalog, prices = safiya(), korea()
    book = Workbook()
    book.remove(book.active)
    summary, everything = [], []
    widths = [5, 14, 15, 52, 12, 14, 46, 11, 11, 12, 12, 15, 13, 13, 14]

    for sheet in pd.ExcelFile(source).sheet_names:
        if not sheet.upper().startswith("ЗАЯВКА"):
            continue
        request = pd.read_excel(source, sheet_name=sheet, header=0)
        request = request[request["Бренд"].notna() & (request["Бренд"] != "ИТОГО")]
        table = compare(request, catalog, prices, import_cost)
        table = table.reindex(columns=COLUMNS)
        есть = table[table["Есть у SAFIYA"] != "нет"]
        everything.append(есть.assign(Заявка=sheet))
        summary.append([sheet, len(table), len(есть),
                        int(есть["Сумма по заявке, руб"].sum(skipna=True)),
                        round(есть["Дороже своего ввоза, %"].median(), 1)
                        if есть["Дороже своего ввоза, %"].notna().any() else ""])
        ws = write(book.create_sheet(sheet[:31]), COLUMNS,
                   table.where(table.notna(), None).values.tolist(), widths,
                   ["", "ИТОГО", "", f"есть у SAFIYA: {len(есть)} из {len(table)}",
                    int(table["Количество, шт"].sum()), "", "", "", "", "",
                    "", int(есть["Сумма по заявке, руб"].sum(skipna=True)), "", "", ""])
        money(ws, "ЕJLMN".replace("Е", "E"))
        money(ws, "IK", "0.00")

    combined = pd.concat(everything) if everything else pd.DataFrame(columns=COLUMNS)
    ws = write(book.create_sheet("ЕСТЬ У SAFIYA"), ["Заявка"] + COLUMNS,
               combined.reindex(columns=["Заявка"] + COLUMNS)
               .where(combined.notna(), None).values.tolist(), [22] + widths)
    money(ws, "FKMN")

    ws = write(book.create_sheet("ИТОГО"),
               ["Заявка", "Позиций", "Есть у SAFIYA", "Сумма по заявке, руб",
                "Дороже своего ввоза, %"], summary, [46, 10, 14, 18, 16],
               ["ВСЕГО", sum(row[1] for row in summary), sum(row[2] for row in summary),
                sum(row[3] for row in summary), ""])
    money(ws, "D")
    book.move_sheet("ИТОГО", -(len(book.sheetnames) - 1))

    os.makedirs("outputs", exist_ok=True)
    book.save(TARGET)
    for row in summary:
        print(f"{row[0]:<12} есть у SAFIYA {row[2]:>3} из {row[1]:>3} позиций   "
              f"на {row[3]:>12,} руб   дороже своего ввоза на {row[4]}%".replace(",", " "))
    print(f"\nСохранено: {TARGET}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Заявка против каталога SAFIYA")
    parser.add_argument("source", nargs="?", default="outputs/Заявки в Корею.xlsx")
    parser.add_argument("--import-cost", type=float, default=10.0,
                        help="наши расходы на ввоз из Кореи, %% к цене")
    args = parser.parse_args()
    main(args.source, args.import_cost)
