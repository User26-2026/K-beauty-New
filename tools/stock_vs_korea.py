"""Наши позиции против корейских цен с доставкой.

По каждой позиции склада ищем самое дешевое корейское предложение —
и в ответах на заявку, и в прайсах поставщиков, — приводим к цене на
складе в Москве и сравниваем с нашей себестоимостью.

Себестоимость в складском файле лежит с наценкой 10%, поэтому делим.
Позицию берем, только если совпал бренд и вид товара: маска и сыворотка
с похожими словами — разные вещи.

Запуск:
    python3 tools/stock_vs_korea.py <остатки.xlsx> --offer "FINESKIN=инвойс.xls" \
        --delivery 15
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import add_barcodes
import brand_names
from brand_names import brand_of
from check_customer_order import read_order
from compare_offers import read_any
from offer_to_customer import SCORE, kind
from rates import KRW_RUB

TARGET = "outputs/Склад против Кореи.xlsx"
# Во сколько раз корейская цена может быть ниже нашей себестоимости,
# прежде чем это перестанет быть ценой того же товара.
PRICE_LIMIT = 3.0
PRICES = "outputs/prices_normalized.xlsx"

HEAD = PatternFill("solid", fgColor="1F3864")
BRAND = PatternFill("solid", fgColor="DDEBF7")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
GOOD = PatternFill("solid", fgColor="E2EFDA")
BAD = PatternFill("solid", fgColor="FCE4D6")
WHITE = Font(color="FFFFFF", bold=True)

SHOWN = ["№", "Бренд", "Наименование", "Остаток, шт", "Наша себестоимость, руб",
         "Цена в Корее, KRW", "Цена EXW, руб", "С доставкой, руб", "У кого",
         "Разница, руб", "Разница, %", "Разница на остаток, руб"]
WIDTHS = [5, 16, 58, 12, 18, 15, 13, 15, 16, 12, 11, 18]


def korean(offers):
    """Все корейские предложения одной таблицей: ответы плюс прайсы."""
    parts = []
    for name, path in offers:
        table = read_any(path)[["Бренд", "Товар", "Цена, KRW"]].copy()
        table["Поставщик"] = name
        parts.append(table)

    prices = pd.read_excel(PRICES, dtype={"Штрихкод": str})
    prices = prices[(prices["Страна"] == "KR")].copy()
    prices = prices[prices["Закупка, KRW"].notna() & (prices["Закупка, KRW"] > 0)]
    prices["Бренд в прайсе"] = prices["Бренд"]
    prices["Бренд"] = brand_names.resolve(prices).fillna("")
    parts.append(prices.rename(columns={"Название EN": "Товар",
                                        "Закупка, KRW": "Цена, KRW"})
                 [["Бренд", "Товар", "Цена, KRW", "Поставщик"]])

    table = pd.concat(parts, ignore_index=True)
    table["Марка"] = (table["Бренд"].astype(str).str.upper()
                      .str.replace(r"[^A-Z0-9]", "", regex=True))
    table["Слова"] = table["Товар"].map(add_barcodes.words)
    table["Тон"] = table["Товар"].map(add_barcodes.tone)
    return table.reset_index(drop=True)


def cheapest(name, brand, offers, cost, delivery):
    """Самое дешевое корейское предложение по позиции.

    Цену ниже трети нашей себестоимости не берем: столько стоит пробник
    или саше, а не та же банка. Если ничего кроме таких нет, говорим об
    этом отдельно, а не подставляем цифру, на которую нельзя опереться.
    """
    fits = []
    for mark in (brand, brand.split()[0] if brand.split() else brand):
        fits = [(score, index) for score, index in
                add_barcodes.candidates(name, mark, offers)
                if score >= SCORE and kind(name) == kind(offers.at[index, "Товар"])]
        if fits:
            break
    if not fits:
        return None, ""
    floor = cost / PRICE_LIMIT / (1 + delivery / 100)
    sane = [pair for pair in fits if offers.at[pair[1], "Цена, KRW"] * KRW_RUB >= floor]
    if not sane:
        return None, "нашлось только в разы дешевле — похоже, пробник"
    return min(sane, key=lambda pair: offers.at[pair[1], "Цена, KRW"])[1], ""


def build(stock, offers, delivery):
    rows = []
    for _, item in stock.iterrows():
        brand = brand_of(item["Наименование"])
        cost = item["Себестоимость, руб"] / 1.1
        found, note = cheapest(item["Наименование"], brand, offers, cost, delivery)
        if found is None:
            rows.append([0, brand, item["Наименование"], int(item["Остаток, шт"]),
                         round(cost), "", "", "", note or "нет предложения",
                         "", "", ""])
            continue
        won = offers.at[found, "Цена, KRW"]
        price = won * KRW_RUB
        landed = price * (1 + delivery / 100)
        gap = landed - cost
        rows.append([0, brand, item["Наименование"], int(item["Остаток, шт"]),
                     round(cost), round(won), round(price), round(landed),
                     offers.at[found, "Поставщик"], round(gap),
                     round(gap / cost * 100, 1), round(gap * item["Остаток, шт"])])
    return pd.DataFrame(rows, columns=SHOWN)


def write(book, title, rows, total=None):
    ws = book.create_sheet(title)
    ws.append(SHOWN)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    brand, number = None, 0
    for row in rows:
        if row[1] != brand:
            brand = row[1]
            ws.append([None, brand])
            for cell in ws[ws.max_row]:
                cell.fill, cell.font = BRAND, Font(bold=True)
        number += 1
        ws.append([number] + list(row[1:]))
        gap = row[10]
        fill = GOOD if gap != "" and gap < 0 else BAD if gap != "" and gap > 0 else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in "DEFGHJL":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    ws.freeze_panes = "C2"
    return ws


def main(source, offers_list, delivery, target):
    stock = read_order(source, everything=True)
    table = build(stock, korean(offers_list), delivery)
    table = table.sort_values(["Бренд", "Наименование"])
    known = table[table["Разница, %"] != ""]

    book = Workbook()
    book.remove(book.active)
    write(book, "ВСЕ ПОЗИЦИИ", table.values.tolist(),
          ["", "ИТОГО", f"позиций: {len(table)}", int(table["Остаток, шт"].sum()),
           "", "", "", "", "", "", "",
           int(pd.to_numeric(table["Разница на остаток, руб"],
                             errors="coerce").sum())])
    дешевле = known[known["Разница, %"] < 0].sort_values("Разница, %")
    write(book, "В КОРЕЕ ДЕШЕВЛЕ", дешевле.values.tolist(),
          ["", "ИТОГО", f"позиций: {len(дешевле)}",
           int(дешевле["Остаток, шт"].sum()), "", "", "", "", "", "", "",
           int(дешевле["Разница на остаток, руб"].sum())])
    дороже = known[known["Разница, %"] > 0].sort_values("Разница, %",
                                                        ascending=False)
    write(book, "В КОРЕЕ ДОРОЖЕ", дороже.values.tolist(),
          ["", "ИТОГО", f"позиций: {len(дороже)}", int(дороже["Остаток, шт"].sum()),
           "", "", "", "", "", "", "",
           int(дороже["Разница на остаток, руб"].sum())])

    os.makedirs("outputs", exist_ok=True)
    book.save(target)
    print(f"Позиций: {len(table)}   сверено с Кореей: {len(known)}")
    print(f"Дешевле в Корее: {len(дешевле)}   дороже: {len(дороже)}")
    print(f"Медиана разницы: {known['Разница, %'].median():.1f}%")
    print(f"На весь остаток: "
          f"{int(known['Разница на остаток, руб'].sum()):,} руб".replace(",", " "))
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Склад против корейских цен")
    parser.add_argument("source")
    parser.add_argument("--offer", action="append", default=[], help="ИМЯ=файл")
    parser.add_argument("--delivery", type=float, default=15.0,
                        help="доставка до нашего склада, %% к цене")
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main(args.source, [offer.split("=", 1) for offer in args.offer],
         args.delivery, args.out)
