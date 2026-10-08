#!/usr/bin/env python3
"""Воронка продаж по SKU из OZON-аналитики «По товарам» (.xlsx).

Считает конверсии из сырых счётчиков (показы, посещения карточки, заказано,
выкуплено) и определяет узкое место воронки + действие. Двухуровневая шапка:
группы на строке 9, подзаголовки на 10, описания на 11, «Итого» на 12,
товары с 13-й. Читаем через calamine (устойчив к «кривым» стилям OZON).

Использование:
  python tools/ozon/funnel.py <аналитика_по_товарам.xlsx> [out.xlsx]
"""
import sys, os, openpyxl
from openpyxl.styles import Font, PatternFill
from python_calamine import CalamineWorkbook

# фиксированные индексы колонок листа «По товарам»
C = dict(name=0, sku=7, pos=17, pokazy=19, posesh=31, zakaz=43, vykup=49,
         cena=59, drr=65, ostatok=69, postavit=71, otzyvy=72)

def num(v):
    if v is None: return None
    s = str(v).replace('%','').replace('\xa0','').replace(' ','').replace(',','.').strip()
    if s in ('','—','-','None'): return None
    try: return float(s)
    except Exception: return None

def pct(a, b): return round(100*a/b, 1) if a and b else 0

def bottleneck(pk, kz, vy, pokazy, zakaz, otz):
    if zakaz < 3 and pokazy < 8000:
        return 'Мало показов — видимость', 'SEO-название + реклама (Трафареты), проверить ключи'
    if pk < 3 and pokazy > 15000:
        return 'Слабый CTR из показа', 'Главное фото/цена/позиция'
    if 0 < kz < 3:
        return 'Карточка не конвертит', f'Цена/контент/отзывы (сейчас {int(otz)})'
    if 0 < vy < 80:
        return 'Низкий выкуп', 'Ожидания (фото/описание), сроки, качество'
    return 'Воронка ОК', 'Масштабировать (реклама при ДРР<15%)'

def main(path, out=None):
    data = CalamineWorkbook.from_path(path).get_sheet_by_name('По товарам').to_python()
    rows = []
    for r in range(12, len(data)):
        d = data[r]
        name = str(d[C['name']]).strip(); sku = str(d[C['sku']]).strip()
        if not name or name.lower().startswith('итог') or not sku or sku == 'None': continue
        rows.append({k: (str(d[i]).strip() if k in ('name',) else num(d[i])) for k, i in C.items()})
    out = out or f"outputs/ozon_funnel/ozon_funnel_{os.path.splitext(os.path.basename(path))[0]}.xlsx"
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Воронка по SKU'
    bold = Font(bold=True)
    ws.append(['Товар','Показы','Пок→карточка %','Карт→заказ %','Заказано','Выкуп %','Отзывы','Позиция','Узкое место','Что делать'])
    for c in ws[1]: c.font = bold; c.fill = PatternFill('solid', fgColor='D9E1F2')
    fills = {'Мало показов — видимость':'FFF2CC','Слабый CTR из показа':'FCE4D6','Карточка не конвертит':'FFC7CE','Низкий выкуп':'E4DFEC','Воронка ОК':'C6EFCE'}
    for p in sorted(rows, key=lambda x: (x['zakaz'] or 0), reverse=True):
        pk = pct(p['posesh'], p['pokazy']); kz = pct(p['zakaz'], p['posesh']); vy = pct(p['vykup'], p['zakaz'])
        bn, act = bottleneck(pk, kz, vy, p['pokazy'] or 0, p['zakaz'] or 0, p['otzyvy'] or 0)
        ws.append([p['name'][:60], p['pokazy'] or 0, pk, kz, p['zakaz'] or 0, vy, p['otzyvy'] or 0, p['pos'] or 0, bn, act])
        ws.cell(ws.max_row, 9).fill = PatternFill('solid', fgColor=fills.get(bn, 'FFFFFF'))
    ws.freeze_panes = 'A2'
    wb.save(out)
    from collections import Counter
    cnt = Counter(bottleneck(pct(p['posesh'],p['pokazy']), pct(p['zakaz'],p['posesh']), pct(p['vykup'],p['zakaz']),
                             p['pokazy'] or 0, p['zakaz'] or 0, p['otzyvy'] or 0)[0] for p in rows)
    print(f"SKU: {len(rows)} | узкие места: {dict(cnt)} | -> {out}")

if __name__ == '__main__':
    if len(sys.argv) < 2: sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
