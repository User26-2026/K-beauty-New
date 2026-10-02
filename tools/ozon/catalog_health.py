#!/usr/bin/env python3
"""Health-check каталога OZON из выгрузки товаров (CSV, разделитель ';').

Находит товары «Не продается» и группирует по причине блокировки с
рекомендуемым действием; отдельным листом — предупреждения. Колонки ищем по
имени заголовка (первая строка), заголовки OZON бывают многострочные.

Использование:
  python tools/ozon/catalog_health.py <каталог.csv> [out.xlsx]
"""
import sys, os, csv, openpyxl
from openpyxl.styles import Font, PatternFill

def gi(hdr, name):
    return next((i for i, h in enumerate(hdr) if name.lower() in h.lower()), None)

def classify(pr, err, art, name):
    t = (pr + ' ' + err).lower(); a = art.lower()
    if any(x in a for x in ('оаррп','рапор','оплгп')) or not name.strip():
        return 'Тестовая/мусорная карточка', 'Удалить карточку'
    if 'копи' in t or 'реплик' in t:
        return 'Копии/реплики запрещены', 'Документы на бренд + апелляция в поддержку OZON'
    if 'документы на бренд' in t:
        return 'Нет документов на бренд', 'Загрузить документы бренда (сертификат/декларация/письмо-разрешение)'
    if 'документ' in t:
        return 'Нет документов качества', 'Загрузить декларацию соответствия / СГР'
    if 'фото' in t or 'изображени' in t:
        return 'Нет/некорректное фото', 'Добавить главное фото; убрать надписи (кэшбэк/розыгрыш) с изображений'
    return 'Другое', pr or err

def main(path, out=None):
    rows = list(csv.reader(open(path, encoding='utf-8-sig'), delimiter=';'))
    hdr = [h.replace('\n', ' ').strip() for h in rows[0]]
    data = [r for r in rows[1:] if any(c.strip() for c in r)]
    I = {k: gi(hdr, k) for k in ('Артикул','Название','Статус','Категория','Причины','Ошибки',
                                 'Предупреждения','Бренд','FBO','FBS','Цена на сайте','Контент-рейтинг')}
    g = lambda r, k: r[I[k]].strip() if I[k] is not None and I[k] < len(r) else ''
    def fnum(x):
        x = (x or '').strip().lstrip("'\"").replace(' ', '').replace('\xa0', '').replace(',', '.')
        try: return float(x)
        except: return 0.0
    stock = lambda r: fnum(g(r, 'FBO')) + fnum(g(r, 'FBS'))
    notsell = [r for r in data if g(r, 'Статус') == 'Не продается']
    out = out or f"outputs/ozon_catalog_health/ozon_blockers_{os.path.splitext(os.path.basename(path))[0]}.xlsx"
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Не продается'
    bold = Font(bold=True)
    ws.append(['Артикул','Название','Категория','Причина блокировки','Что делать'])
    for c in ws[1]: c.font = bold; c.fill = PatternFill('solid', fgColor='FFC7CE')
    from collections import Counter
    cnt = Counter()
    for r in sorted(notsell, key=lambda x: classify(g(x,'Причины'), g(x,'Ошибки'), g(x,'Артикул'), g(x,'Название'))[0]):
        pr, act = classify(g(r,'Причины'), g(r,'Ошибки'), g(r,'Артикул'), g(r,'Название'))
        cnt[pr] += 1
        ws.append([g(r,'Артикул'), g(r,'Название'), g(r,'Категория'), pr, act])
    ws2 = wb.create_sheet('Предупреждения'); ws2.append(['Артикул','Название','Категория','Предупреждение'])
    for c in ws2[1]: c.font = bold; c.fill = PatternFill('solid', fgColor='FFF2CC')
    for r in data:
        if g(r, 'Предупреждения'): ws2.append([g(r,'Артикул'), g(r,'Название'), g(r,'Категория'), g(r,'Предупреждения')])
    # Приоритет пополнения: карточка НЕ заблокирована, но остаток = 0 (продавала бы, если бы был сток)
    ws3 = wb.create_sheet('Нет стока (пополнить)')
    ws3.append(['Артикул','Название','Бренд','Категория','Статус','FBO','FBS','Цена, ₽'])
    for c in ws3[1]: c.font = bold; c.fill = PatternFill('solid', fgColor='DDEBF7')
    restock = [r for r in data if g(r,'Статус') != 'Не продается' and stock(r) == 0]
    for r in sorted(restock, key=lambda x: (g(x,'Бренд'), g(x,'Название'))):
        ws3.append([g(r,'Артикул'), g(r,'Название'), g(r,'Бренд'), g(r,'Категория'),
                    g(r,'Статус'), fnum(g(r,'FBO')), fnum(g(r,'FBS')), fnum(g(r,'Цена на сайте'))])
    # Сводка (KPI)
    from collections import Counter
    ws0 = wb.create_sheet('Сводка', 0)
    st = Counter(g(r,'Статус') for r in data)
    n = len(data); nostock = sum(1 for r in data if stock(r) == 0)
    crs = [fnum(g(r,'Контент-рейтинг')) for r in data if g(r,'Контент-рейтинг')]
    lowcr = sum(1 for v in crs if v < 70)
    kpi = [
        ['Показатель', 'Значение'],
        ['Товаров в каталоге', n],
        ['— Продаётся', st.get('Продается', 0)],
        ['— Готов к продаже', st.get('Готов к продаже', 0)],
        ['— Не продаётся (заблокировано)', st.get('Не продается', 0)],
        ['Без остатка (FBO+FBS = 0)', f'{nostock} ({round(100*nostock/n)}%)'],
        ['Готовы, но без стока (приоритет пополнения)', len(restock)],
        ['Контент-рейтинг < 70 (карточки дозаполнить)', lowcr],
    ]
    for row in kpi: ws0.append(row)
    for c in ws0[1]: c.font = bold; c.fill = PatternFill('solid', fgColor='1F4E78'); c.font = Font(bold=True, color='FFFFFF')
    ws0.append([]); ws0.append(['Блокировки «Не продаётся» по причинам:'])
    ws0[ws0.max_row][0].font = bold
    for k, v in cnt.most_common(): ws0.append([k, v])
    ws0.column_dimensions['A'].width = 46; ws0.column_dimensions['B'].width = 16
    wb.save(out)
    from collections import Counter as C2
    st = C2(g(r, 'Статус') for r in data)
    print(f"Товаров: {len(data)} | статусы: {dict(st)} | не продаётся по причинам: {dict(cnt)} | -> {out}")

if __name__ == '__main__':
    if len(sys.argv) < 2: sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
