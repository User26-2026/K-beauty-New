"""Сколько SAFIYA накидывает к корейской цене.

SAFIYA держит склад в Екатеринбурге, а товар везет из Кореи. Значит, в
ее цене сидит все сразу: доставка, растаможка и собственный заработок.
Разделить их по прайсу нельзя, поэтому считаем две вещи:

- наценку к корейской цене EXW — весь разрыв целиком;
- переплату к нашему импорту — корейская цена плюс наши расходы на
  ввоз. Вот это и есть ответ, выгоднее ли брать у SAFIYA или везти
  самим. По умолчанию 10%: столько выходит по машине.

Счет SAFIYA выставляет по своему курсу доллара, выше рыночного, поэтому
в рубли переводим по нему: иначе наценка выйдет меньше, чем на деле.

Сравниваем цену, как она напечатана в прайсе: приведение к штуке делит
цену на число пэдов в банке у одной стороны и не делит у другой. Где
разрыв невозможный, позиция уходит на сверку, а не в расчет.

Запуск:
    python3 tools/safiya_vs_korea.py --import-cost 10
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brand_names
from rates import KRW_RUB, SAFIYA_USD_RUB, USD_RUB
from supplier_brand_matrix import split_conflicts

SRC = "outputs/prices_normalized.xlsx"
ROOT = "data/price_lists/safiya"
OFFERS = {
    "Спец. предложение 12.08": f"{ROOT}/archive/safiya_2026-08-12_znakomomu.xlsx",
    "B2B-прайс 01.09": f"{ROOT}/safiya_2026-09-01_b2b.xlsx",
}
TARGET = "outputs/SAFIYA_против_Кореи.xlsx"
# Разрыв больше этого — не наценка, а разная фасовка: цена за банку
# против цены за пэд. Такие позиции считаем спорными.
LIMIT = 200

HEAD = PatternFill("solid", fgColor="1F3864")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)


def key(brand):
    return re.sub(r"[^A-ZА-Я0-9]", "", str(brand).upper())


def same_brand(left, right):
    """Один бренд пишут по-разному: VT и VT COSMETICS, CENTELLIAN24 и
    CENTELLIAN 24 MADECA. Считаем совпадением, если короткое написание —
    начало длинного."""
    one, two = key(left), key(right)
    if not one or not two:
        return False
    short, long = sorted((one, two), key=len)
    return len(short) >= 2 and long.startswith(short)


def korea():
    """Самая дешевая корейская цена по каждому штрихкоду, в рублях EXW."""
    table = pd.read_excel(SRC, dtype={"Штрихкод": str})
    table = table[table["Страна"] == "KR"].copy()
    table["Штрихкод"] = (table["Штрихкод"].fillna("").astype(str)
                         .str.replace(r"\D", "", regex=True))
    table = table[table["Штрихкод"].str.len() >= 8]
    table = table[table["Закупка, KRW"].notna() & (table["Закупка, KRW"] > 0)]
    table["Бренд в прайсе"] = table["Бренд"]
    table["Бренд"] = brand_names.resolve(table).fillna("БЕЗ БРЕНДА")
    table, _ = split_conflicts(table)
    table["Корея EXW, руб"] = (table["Закупка, KRW"] * KRW_RUB).round()
    best = table.sort_values("Корея EXW, руб").drop_duplicates("Штрихкод")
    return best.set_index("Штрихкод")[
        ["Поставщик", "Бренд", "Закупка, KRW", "Корея EXW, руб"]
    ].rename(columns={"Поставщик": "Корея, поставщик", "Бренд": "Бренд KR",
                      "Закупка, KRW": "Корея, KRW"})


def read_personal(path):
    table = pd.read_excel(path, skiprows=4, header=None)
    table.columns = ["Бренд", "Категория", "Наименование", "Штрихкод",
                     "Объем", "Остаток, шт", "Цена, $"]
    return table.dropna(subset=["Цена, $"])


def read_b2b(path):
    """B2B-прайс: бренд стоит заголовком секции, товары — строками ниже."""
    table = pd.read_excel(path, header=0)
    table.columns = ["Код", "Наименование", "Объем", "Цена, $", "Кол-во", "Сумма"]
    goods, brand = [], ""
    for _, row in table.iterrows():
        code = str(row["Код"]).strip()
        if code.isdigit() and len(code) >= 12:
            goods.append({"Бренд": brand, "Наименование": str(row["Наименование"]).strip(),
                          "Штрихкод": code, "Объем": row["Объем"],
                          "Цена, $": float(row["Цена, $"])})
        elif code and code != "nan" and pd.isna(row["Цена, $"]):
            brand = code
    return pd.DataFrame(goods)


def read_offer(path):
    table = read_b2b(path) if "b2b" in os.path.basename(path) else read_personal(path)
    table["Штрихкод"] = table["Штрихкод"].astype(str).str.replace(r"\D", "", regex=True)
    return table


def compare(offer, best, import_cost):
    table = offer.join(best, on="Штрихкод")
    table["Цена SAFIYA, руб"] = (table["Цена, $"] * SAFIYA_USD_RUB).round()
    table["Своим импортом, руб"] = (table["Корея EXW, руб"] * (1 + import_cost / 100)).round()
    table["Наценка к Корее, %"] = (
        table["Цена SAFIYA, руб"] / table["Корея EXW, руб"] * 100 - 100).round(1)
    table["К своему импорту, %"] = (
        table["Цена SAFIYA, руб"] / table["Своим импортом, руб"] * 100 - 100).round(1)
    # Бренд подтверждаем с двух сторон: один штрихкод у разных поставщиков
    # иногда стоит на разном товаре.
    table["Замечание"] = [
        "нет в корейских прайсах" if pd.isna(kr) else
        "бренд не совпал" if not same_brand(kr, ru) else
        "на сверку: разная фасовка" if markup > LIMIT or markup < -50 else ""
        for kr, ru, markup in zip(table["Бренд KR"], table["Бренд"],
                                  table["Наценка к Корее, %"].fillna(0))
    ]
    return table


def write(ws, columns, rows, widths, total=None):
    ws.append(columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return ws


def money(ws, letters, fmt="# ##0"):
    for letter in letters:
        for cell in ws[letter][1:]:
            cell.number_format = fmt


def scale(ws, letter, rows, low, high):
    if rows < 2:
        return
    ws.conditional_formatting.add(
        f"{letter}2:{letter}{rows}",
        ColorScaleRule(start_type="num", start_value=low, start_color="63BE7B",
                       mid_type="num", mid_value=(low + high) / 2, mid_color="FFEB84",
                       end_type="num", end_value=high, end_color="F8696B"))


def main(import_cost):
    best = korea()
    book = None
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    summary = []

    for label, path in OFFERS.items():
        table = compare(read_offer(path), best, import_cost)
        clean = table[table["Замечание"] == ""]
        summary.append([
            label, len(table), len(clean),
            round(clean["Наценка к Корее, %"].median(), 1),
            round(clean["К своему импорту, %"].median(), 1),
            int((clean["К своему импорту, %"] < 0).sum()),
        ])

        columns = ["Бренд", "Наименование", "Объем", "Штрихкод", "Корея, поставщик",
                   "Корея, KRW", "Корея EXW, руб", "Своим импортом, руб", "Цена, $",
                   "Цена SAFIYA, руб", "Наценка к Корее, %", "К своему импорту, %"]
        part = clean.sort_values("Наценка к Корее, %")
        ws = write(book.create_sheet(label[:31]), columns,
                   part[columns].values.tolist(),
                   [18, 56, 12, 15, 15, 12, 13, 14, 10, 14, 13, 13],
                   ["ИТОГО", f"позиций: {len(part)}", "", "", "", "",
                    int(part["Корея EXW, руб"].median()),
                    int(part["Своим импортом, руб"].median()), "",
                    int(part["Цена SAFIYA, руб"].median()),
                    round(part["Наценка к Корее, %"].median(), 1),
                    round(part["К своему импорту, %"].median(), 1)])
        money(ws, "FGHJ")
        money(ws, "I", "0.00")
        scale(ws, "K", ws.max_row - 1, 0, 150)
        scale(ws, "L", ws.max_row - 1, -20, 80)

        brands = []
        for brand, group in clean.groupby("Бренд", sort=True):
            brands.append([brand, len(group),
                           int(group["Корея EXW, руб"].median()),
                           int(group["Цена SAFIYA, руб"].median()),
                           round(group["Наценка к Корее, %"].median(), 1),
                           round(group["К своему импорту, %"].median(), 1)])
        ws = write(book.create_sheet(f"{label[:22]} бренды"),
                   ["Бренд", "Позиций", "Корея EXW, руб", "Цена SAFIYA, руб",
                    "Наценка к Корее, %", "К своему импорту, %"],
                   brands, [22, 9, 15, 16, 15, 16],
                   ["ВСЕГО", len(clean), int(clean["Корея EXW, руб"].median()),
                    int(clean["Цена SAFIYA, руб"].median()),
                    round(clean["Наценка к Корее, %"].median(), 1),
                    round(clean["К своему импорту, %"].median(), 1)])
        money(ws, "CD")
        scale(ws, "E", ws.max_row - 1, 0, 150)
        scale(ws, "F", ws.max_row - 1, -20, 80)

        spor = table[table["Замечание"] != ""]
        write(book.create_sheet(f"{label[:20]} на сверку"),
              ["Бренд", "Наименование", "Штрихкод", "Бренд KR", "Цена, $",
               "Корея EXW, руб", "Наценка к Корее, %", "Замечание"],
              spor[["Бренд", "Наименование", "Штрихкод", "Бренд KR", "Цена, $",
                    "Корея EXW, руб", "Наценка к Корее, %", "Замечание"]].values.tolist(),
              [18, 56, 15, 16, 10, 13, 13, 26])

    ws = write(book.create_sheet("ИТОГО"),
               ["Прайс SAFIYA", "Позиций", "Сверено с Кореей",
                "Наценка к Корее, %", "К своему импорту, %", "Дешевле нашего импорта, поз."],
               summary, [26, 10, 16, 15, 17, 18])
    scale(ws, "D", ws.max_row, 0, 150)
    scale(ws, "E", ws.max_row, -20, 80)
    book.move_sheet("ИТОГО", -(len(book.sheetnames) - 1))

    os.makedirs("outputs", exist_ok=True)
    book.save(TARGET)
    print(f"Курс SAFIYA: {SAFIYA_USD_RUB} руб/доллар "
          f"(рыночный {USD_RUB}, это +{SAFIYA_USD_RUB / USD_RUB * 100 - 100:.1f}%)")
    print(f"Наши расходы на импорт из Кореи: +{import_cost:.0f}%\n")
    for row in summary:
        print(f"{row[0]:<24} сверено {row[2]:>3} поз.  "
              f"наценка к Корее {row[3]:>6.1f}%   к своему импорту {row[4]:>6.1f}%   "
              f"дешевле нашего импорта: {row[5]} поз.")
    print(f"\nСохранено: {TARGET}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="SAFIYA против корейских цен")
    parser.add_argument("--import-cost", type=float, default=10.0,
                        help="наши расходы на ввоз из Кореи, %% к цене")
    args = parser.parse_args()
    main(args.import_cost)
