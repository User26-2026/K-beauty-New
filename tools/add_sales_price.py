"""Проставить наши продажные цены к списку остатков.

В остатках есть себестоимость, но нет цены, по которой мы продаем. Цену
берем из нашего оптового прайса и считаем, на какую сумму лежит склад в
продажных ценах.

С ключом --sales добавляем продажи из отчета WB: сколько ушло за месяц,
на какую сумму и на сколько месяцев хватит остатка. Один товар на WB
может идти несколькими карточками, поэтому продажи сначала складываем по
названию.

Общего кода у файлов нет, сводим по названию: в прайсе и в остатках оно
пишется одинаково, поэтому сначала сравниваем строку целиком, а остаток
догоняем разбором по словам.

Запуск:
    python3 tools/add_sales_price.py <остатки.xlsx> [--out файл.xlsx]
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import name_match

PRICES = "data/pricing/pricelist_2026-08-24.xlsx"
# Колонки отчета WB: название, чистые продажи в штуках и в рублях с СПП.
SALES_COLUMNS = {1: "Товар", 15: "Продано, шт", 18: "Выручка, руб"}
SALES_SKIP = 5   # первые строки отчета — шапка в три этажа и две строки итогов
TARGET = "outputs/Остатки с продажными ценами.xlsx"

HEAD = PatternFill("solid", fgColor="1F3864")
MISS = PatternFill("solid", fgColor="FCE4D6")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)

COLUMNS = ["Наименование", "Остаток, шт", "Себестоимость, руб", "Срок годности",
           "Продажная цена, руб", "Продажа х остаток, руб"]
SALES_ADDED = ["Продано за месяц, шт", "Выручка за месяц, руб", "Запас, месяцев"]


def normal(name):
    return re.sub(r"[^0-9a-zа-яё]+", "", str(name).lower())


def read_prices():
    table = pd.read_excel(PRICES, skiprows=4)
    table.columns = ["№", "Наименование", "Штрихкод", "В наличии, шт", "Цена, ₽"]
    table = table.dropna(subset=["Цена, ₽"])
    return table.reset_index(drop=True)


def read_stock(path):
    table = pd.read_excel(path, header=0)
    table = table.iloc[:, :4]
    table.columns = COLUMNS[:4]
    table = table[table["Наименование"].notna()]
    table = table[table["Остаток, шт"].notna()]
    table["Остаток, шт"] = pd.to_numeric(
        table["Остаток, шт"].astype(str).str.replace(r"\s", "", regex=True),
        errors="coerce")
    return table[table["Остаток, шт"].notna()].reset_index(drop=True)


def read_sales(path):
    """Отчет WB: продажи по карточкам, сложенные по названию товара."""
    table = pd.read_excel(path, sheet_name=0, header=None).iloc[SALES_SKIP:]
    table = table[list(SALES_COLUMNS)].rename(columns=SALES_COLUMNS)
    table = table[table["Товар"].notna()]
    for column in ("Продано, шт", "Выручка, руб"):
        table[column] = pd.to_numeric(table[column], errors="coerce").fillna(0)
    return table.groupby("Товар", as_index=False).sum()


def sales_by_stock(stock, sales):
    """Продажи, сложенные по позициям остатков.

    Один товар идет на WB несколькими карточками: "…BOOSTER",
    "…BOOSTER [15ml]" и "…BOOSTER [15ml] N" — это одна и та же банка.
    Поэтому идем от карточки к остаткам, а не наоборот, и складываем;
    карточку берем, только если подходит ровно одна строка остатков.
    """
    rest_words = {index: name_match.words(name, True)
                  for index, name in stock["Наименование"].items()}
    rest_text = {index: normal(name) for index, name in stock["Наименование"].items()}
    totals = {index: [0.0, 0.0, False] for index in stock.index}
    taken = set()

    for _, card in sales.iterrows():
        text, parts = normal(card["Товар"]), name_match.words(card["Товар"], True)
        fits = [index for index in stock.index if rest_text[index] == text]
        if not fits and parts:
            fits = [index for index in stock.index
                    if parts and parts <= rest_words[index]]
        if len(fits) != 1:
            continue
        found = totals[fits[0]]
        found[0] += card["Продано, шт"]
        found[1] += card["Выручка, руб"]
        found[2] = True
        taken.add(card.name)

    # Что не разобралось по словам, догоняем разбором названий: там свои
    # заходы — начало строки, набор слов без фасовки, похожесть.
    rest = stock.loc[[not totals[index][2] for index in stock.index], "Наименование"]
    free = sales.drop(index=list(taken))
    for index, other in name_match.match(rest, free["Товар"]).items():
        card = free.loc[other]
        totals[index] = [card["Продано, шт"], card["Выручка, руб"], True]

    return pd.DataFrame(
        [totals[index] for index in stock.index],
        index=stock.index, columns=["Продано, шт", "Выручка, руб", "Нашлось"])


def attach_sales(stock, sales):
    found = sales_by_stock(stock, sales)
    for column, field in (("Продано за месяц, шт", "Продано, шт"),
                          ("Выручка за месяц, руб", "Выручка, руб")):
        stock[column] = [value if seen else None
                         for value, seen in zip(found[field], found["Нашлось"])]
    # Запас в месяцах: сколько еще продержится склад при том же темпе.
    stock["Запас, месяцев"] = [
        round(rest / sold, 1) if sold else None
        for rest, sold in zip(stock["Остаток, шт"], stock["Продано за месяц, шт"])
    ]
    return stock


def attach(stock, prices):
    """Цена из прайса: сначала точное совпадение строки, потом по словам."""
    by_text = {}
    for index, name in prices["Наименование"].items():
        by_text.setdefault(normal(name), index)
    found = pd.Series([by_text.get(normal(name)) for name in stock["Наименование"]],
                      index=stock.index, dtype="object")

    left = stock.loc[found.isna(), "Наименование"]
    free = prices.drop(index=[i for i in found.dropna()])
    for index, other in name_match.match(left, free["Наименование"]).items():
        found[index] = other

    stock["Продажная цена, руб"] = [
        prices.loc[index, "Цена, ₽"] if pd.notna(index) else None for index in found]
    stock["Продажа х остаток, руб"] = (
        stock["Продажная цена, руб"] * stock["Остаток, шт"]).round()
    return stock


def main(source, target, sales_path):
    stock = attach(read_stock(source), read_prices())
    columns = list(COLUMNS)
    if sales_path:
        stock = attach_sales(stock, read_sales(sales_path))
        columns += SALES_ADDED
    book = Workbook()
    ws = book.active
    ws.title = "ОСТАТКИ"
    ws.append(columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in stock[columns].where(stock.notna(), None).values.tolist():
        ws.append(row)
        if row[4] is None:
            # Цены в прайсе нет — строка не попадает в сумму, это видно.
            for cell in ws[ws.max_row]:
                cell.fill = MISS

    total = ["ИТОГО", int(stock["Остаток, шт"].sum()), "", "", "",
             int(stock["Продажа х остаток, руб"].sum(skipna=True))]
    if sales_path:
        total += [int(stock["Продано за месяц, шт"].sum(skipna=True)),
                  int(stock["Выручка за месяц, руб"].sum(skipna=True)), ""]
    ws.append(total)
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = TOTAL, Font(bold=True)

    for index, width in enumerate([96, 12, 16, 13, 16, 18, 15, 16, 12], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in "BEFGH":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    for cell in ws["C"][1:]:
        cell.number_format = "# ##0.00"
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(target) or ".", exist_ok=True)
    book.save(target)
    priced = stock["Продажная цена, руб"].notna().sum()
    print(f"Позиций: {len(stock)}   с продажной ценой: {priced}   "
          f"без цены: {len(stock) - priced}")
    if sales_path:
        sold = stock["Продано за месяц, шт"]
        print(f"Нашлось в отчете продаж: {sold.notna().sum()}   "
              f"продано {int(sold.sum(skipna=True)):,} шт на "
              f"{int(stock['Выручка за месяц, руб'].sum(skipna=True)):,} руб"
              .replace(",", " "))
    print(f"Склад в продажных ценах: "
          f"{int(stock['Продажа х остаток, руб'].sum(skipna=True)):,} руб".replace(",", " "))
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Продажные цены к остаткам")
    parser.add_argument("source")
    parser.add_argument("--out", default=TARGET)
    parser.add_argument("--sales", help="отчет WB по продажам за месяц")
    args = parser.parse_args()
    main(args.source, args.out, args.sales)
