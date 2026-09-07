"""Продажи против остатков по всем товарам кабинета.

Отчет WB несет и продажи, и остаток по каждой карточке, поэтому берем
его целиком, а не только те позиции, что попали в складской файл. Дальше
списываем остаток по месяцам с сезонным коэффициентом и смотрим, на
сколько его хватит.

Пустые карточки — без остатка и без продаж — в расчет не идут, это
архив кабинета.

Запуск:
    python3 tools/sales_vs_stock.py --sales data/sales/wb_sales_2026-08.xls
"""

import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from stock_forecast import (HEAD, LATER, NONE_, PLAN, REST_COLUMN, SHOWN_MONTHS,
                            SOON, TOTAL, WHITE, run_out)

TARGET = "outputs/Продажи и остатки.xlsx"
# Колонки отчета WB: артикул, название, остатки FBO и FBS, продажи.
REPORT = {0: "Артикул ВБ", 1: "Товар", 7: "FBO", 8: "FBS",
          15: "Продано за месяц, шт", 18: "Выручка за месяц, руб"}
SKIP = 5

SHOWN = (["Артикул ВБ", "Товар", "Остаток, шт", "Продано за месяц, шт",
          "Выручка за месяц, руб"] + PLAN + [REST_COLUMN, "Хватит до", "Запас, месяцев"])
WIDTHS = [13, 62, 12, 15, 16] + [11] * SHOWN_MONTHS + [16, 22, 12]
LAST_NUMBER = 5 + SHOWN_MONTHS + 1     # по какую колонку идут числа


def read_report(path):
    table = pd.read_excel(path, sheet_name=0, header=None).iloc[SKIP:]
    table = table[list(REPORT)].rename(columns=REPORT)
    table = table[table["Товар"].notna()]
    for column in ("FBO", "FBS", "Продано за месяц, шт", "Выручка за месяц, руб"):
        table[column] = pd.to_numeric(table[column], errors="coerce").fillna(0)
    table["Остаток, шт"] = table["FBO"] + table["FBS"]
    return table.reset_index(drop=True)


def forecast(table):
    rows = []
    for _, item in table.iterrows():
        rest, monthly = item["Остаток, шт"], item["Продано за месяц, шт"]
        head = [item["Артикул ВБ"], item["Товар"], int(rest), int(monthly),
                round(item["Выручка за месяц, руб"])]
        if not monthly:
            rows.append(head + [""] * SHOWN_MONTHS + [int(rest), "нет продаж", ""])
            continue
        if not rest:
            rows.append(head + [0] * SHOWN_MONTHS + [0, "закончился", 0])
            continue
        spent, months, until = run_out(rest, monthly)
        spent += [0] * (SHOWN_MONTHS - len(spent))
        rows.append(head + spent + [int(max(rest - sum(spent), 0)), until,
                                    round(months, 1)])
    return pd.DataFrame(rows, columns=SHOWN)


def write(book, title, rows, total=None):
    ws = book.create_sheet(title)
    ws.append(SHOWN)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
        months = row[-1]
        fill = (NONE_ if months == "" else
                SOON if months <= 4 else LATER if months <= 8 else None)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for index in range(3, LAST_NUMBER + 2):
        for cell in ws[get_column_letter(index)][1:]:
            cell.number_format = "# ##0"
    ws.freeze_panes = "C2"
    return ws


def main(sales_path, target):
    report = read_report(sales_path)
    live = report[(report["Остаток, шт"] > 0) | (report["Продано за месяц, шт"] > 0)]
    table = forecast(live)

    book = Workbook()
    book.remove(book.active)
    order = table.sort_values("Запас, месяцев",
                              key=lambda col: col.replace("", 10 ** 6))
    write(book, "ВСЕ ТОВАРЫ", order.values.tolist(),
          ["", "ИТОГО", int(table["Остаток, шт"].sum()),
           int(table["Продано за месяц, шт"].sum()),
           int(table["Выручка за месяц, руб"].sum()),
           *[int(pd.to_numeric(table[column], errors="coerce").sum()) for column in PLAN],
           int(table[REST_COLUMN].sum()), "", ""])

    скоро = order[[value != "" and value <= SHOWN_MONTHS
                   for value in order["Запас, месяцев"]]]
    write(book, "КОНЧИТСЯ ДО МАЯ", скоро.values.tolist())
    write(book, "НЕТ ПРОДАЖ",
          order[order["Хватит до"] == "нет продаж"].values.tolist())
    # Продажи были, а товара нет — это уже упущенная выручка.
    write(book, "ЗАКОНЧИЛСЯ",
          order[order["Хватит до"] == "закончился"].values.tolist())

    os.makedirs("outputs", exist_ok=True)
    book.save(target)
    known = table[table["Запас, месяцев"] != ""]
    print(f"Карточек в отчете: {len(report)}   в работе: {len(table)}   "
          f"пустых: {len(report) - len(table)}")
    print(f"Остаток: {int(table['Остаток, шт'].sum()):,} шт   "
          f"продано за месяц: {int(table['Продано за месяц, шт'].sum()):,} шт"
          .replace(",", " "))
    print(f"Кончится до мая: {len(скоро)}   без продаж: "
          f"{(table['Хватит до'] == 'нет продаж').sum()}   "
          f"закончилось: {(table['Хватит до'] == 'закончился').sum()}")
    if len(known):
        print(f"Медиана запаса: {known['Запас, месяцев'].median():.1f} мес")
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Продажи против остатков")
    parser.add_argument("--sales", default="data/sales/wb_sales_2026-08.xls")
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main(args.sales, args.out)
