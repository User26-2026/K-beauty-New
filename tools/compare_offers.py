"""Ответы поставщиков на заявку в одной таблице.

Поставщики отвечают по-разному: кто-то присылает свой инвойс, кто-то
возвращает нашу же заявку с проставленными ценами. Приводим оба вида к
одному виду, сводим по штрихкоду, а где его нет — по названию внутри
бренда, и ставим колонку на поставщика.

Рядом наша текущая себестоимость: видно, стоит ли вообще брать.

Запуск:
    python3 tools/compare_offers.py "FINESKIN=инвойс.xls" "J2K=ответ.xlsx" \
        --stock <остатки.xlsx>
"""

import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import add_barcodes
import name_match
from parse_price_lists import normalize_brand
from brand_names import same
from check_customer_order import read_order
from check_offer import read_offer
from rates import KRW_RUB

TARGET = "outputs/Ответы поставщиков.xlsx"
IMPORT_COST = 1.10

HEAD = PatternFill("solid", fgColor="1F3864")
BRAND = PatternFill("solid", fgColor="DDEBF7")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
BEST = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)


def read_reply(path):
    """Наша заявка, вернувшаяся с ценами: цены стоят не по всем строкам."""
    parts = []
    for sheet in pd.ExcelFile(path).sheet_names:
        table = pd.read_excel(path, sheet_name=sheet, header=0)
        # Поставщик дописывает в шапку свой перевод: «Цена, KRW\n(공급가격)».
        column = next((name for name in table.columns
                       if str(name).startswith("Цена, KRW")), None)
        if column is None:
            continue
        table = table.rename(columns={column: "Цена, KRW"})
        table = table[table["Цена, KRW"].notna()]
        if table.empty:
            continue
        table = table[["Бренд", "Товар", "Количество, шт", "Цена, KRW", "Штрихкод"]]
        table.columns = ["Бренд", "Товар", "Кол-во, шт", "Цена, KRW", "Штрихкод"]
        parts.append(table)
    if not parts:
        return pd.DataFrame(columns=["Бренд", "Товар", "Кол-во, шт",
                                     "Цена, KRW", "Штрихкод"])
    table = pd.concat(parts, ignore_index=True)
    # Пустой штрихкод должен стать пустой строкой: NaN в словаре
    # схлопывает все безкодовые строки в одну.
    codes = pd.to_numeric(table["Штрихкод"], errors="coerce").astype("Int64")
    table["Штрихкод"] = [str(code) if pd.notna(code) else "" for code in codes]
    table["Бренд"] = table["Бренд"].map(same)
    return table


def read_any(path):
    """Инвойс поставщика или наша заявка с ценами — определяем по шапке."""
    head = pd.read_excel(path, sheet_name=0, header=None, nrows=1)
    return read_reply(path) if str(head.iat[0, 0]).strip() == "№" else read_offer(path)


# Цена одного товара у двух корейских поставщиков не может отличаться в
# разы: такой разрыв означает, что свелись разные фасовки.
PRICE_LIMIT = 3.0


def merge(offers):
    """Все позиции всех ответов в одну таблицу, цена — колонка на поставщика."""
    base = None
    for name, table in offers.items():
        table = table.rename(columns={"Цена, KRW": name})
        if base is None:
            base = table[["Бренд", "Товар", "Кол-во, шт", "Штрихкод", name]].copy()
            continue
        base[name] = None
        with_code = {code: index for index, code in
                     zip(base.index, base["Штрихкод"])
                     if isinstance(code, str) and code}
        free = []
        for _, item in table.iterrows():
            code = item["Штрихкод"] if isinstance(item["Штрихкод"], str) else ""
            found = with_code.get(code) if code else None
            if found is None:
                free.append(item)
                continue
            base.at[found, name] = item[name]
        if free:
            rest = pd.DataFrame(free)
            # match, а не to_one: у одного поставщика название короткое и
            # заглавными, у другого длинное с фасовкой, и вложенность
            # набора слов работает в обе стороны.
            pairs = name_match.match(rest["Товар"], base["Товар"])
            extra = []
            known = [column for column in base.columns
                     if column in offers and column != name]
            for index, item in rest.iterrows():
                target = pairs.get(index)
                if target is not None and sane(base, target, known, item[name]):
                    base.at[target, name] = item[name]
                else:
                    extra.append({"Бренд": item["Бренд"], "Товар": item["Товар"],
                                  "Кол-во, шт": item["Кол-во, шт"],
                                  "Штрихкод": item["Штрихкод"], name: item[name]})
            if extra:
                base = pd.concat([base, pd.DataFrame(extra)], ignore_index=True)
    return base


def sane(base, target, known, price):
    """Пара по названию правдоподобна, если цены сопоставимы."""
    for column in known:
        other = base.at[target, column]
        if pd.notna(other) and other:
            ratio = max(price, other) / min(price, other)
            if ratio > PRICE_LIMIT:
                return False
    return True


def as_price_rows(pairs, known, country="KR", brand=None):
    """Ответы на заявку в виде строк прайса.

    Инвойсы в сводную таблицу прайсов не попадают, поэтому поставщик,
    который ответил на заявку, в сравнении брендов не участвует. Штрихкод
    в ответах стоит не везде — недостающий берем по названию из прайсов
    того же бренда, но только при сопоставимой цене: рядом с банкой на 60
    патчей у PETITFEE идет пробник «1 pairs».
    """
    if not pairs:
        return pd.DataFrame()
    pool = known.copy()
    pool["Марка"] = (pool["Бренд"].astype(str).str.upper()
                     .str.replace(r"[^A-Z0-9]", "", regex=True))
    pool["Слова"] = pool["Название EN"].map(add_barcodes.words)
    pool["Тон"] = pool["Название EN"].map(add_barcodes.tone)

    rows = []
    for name, path in pairs:
        table = read_any(path)
        if brand:
            mask = (table["Бренд"].fillna("").astype(str).str.upper()
                    .str.contains(brand.upper())
                    | table["Товар"].fillna("").astype(str).str.upper()
                    .str.contains(brand.upper()))
            table = table[mask]
        for _, item in table.iterrows():
            mark = brand or item["Бренд"]
            if not mark or pd.isna(item["Цена, KRW"]):
                continue
            code = str(item["Штрихкод"] or "")
            if len(code) < 8:
                code = guess_code(item["Товар"], item["Цена, KRW"], mark, pool)
            if not code:
                continue
            rows.append({
                "Поставщик": name, "Страна": country, "Валюта": "KRW",
                # Бренд из инвойса приводим к общему написанию: иначе
                # ETUDE HOUSE и ETUDE разойдутся на два бренда.
                "Бренд": normalize_brand(str(mark).upper()), "Штрихкод": code,
                "Название EN": item["Товар"], "Объем": None,
                "Базис": "EXW", "Единица цены": "за шт", "Штук в упаковке": None,
                "MSRP, KRW": None, "Закупка, KRW": item["Цена, KRW"],
                "Цена за штуку, KRW": item["Цена, KRW"],
                "Цена за штуку, руб": item["Цена, KRW"] * KRW_RUB,
            })
    if not rows:
        return pd.DataFrame()
    # Одну позицию поставщик мог прислать дважды — банкой и пробником.
    return pd.DataFrame(rows).drop_duplicates(["Поставщик", "Штрихкод"])


def guess_code(name, price, brand, pool):
    """Штрихкод по названию, если цена сопоставима с прайсовой."""
    if pool.empty:
        return None
    fits = [index for _, index in add_barcodes.candidates(name, brand, pool)
            if pd.notna(pool.at[index, "Закупка, KRW"])
            and pool.at[index, "Закупка, KRW"] > 0
            and max(price, pool.at[index, "Закупка, KRW"])
            / min(price, pool.at[index, "Закупка, KRW"]) <= PRICE_LIMIT]
    codes = set(pool.loc[fits, "Штрихкод"])
    return fits and len(codes) == 1 and pool.at[fits[0], "Штрихкод"] or None


def read_rfq(path):
    """Наша заявка как основа таблицы: строки брендов в ней пропускаем."""
    table = pd.read_excel(path, sheet_name=0, header=0)
    table = table[table["Товар"].notna()]
    codes = pd.to_numeric(table["Штрихкод"], errors="coerce").astype("Int64")
    return pd.DataFrame({
        "Бренд": table["Бренд"].map(same),
        "Товар": table["Товар"],
        "Кол-во, шт": table["Количество, шт"],
        "Штрихкод": [str(code) if pd.notna(code) else "" for code in codes],
        # Своей цены у заявки нет, но merge ждет колонку у каждого источника.
        "Цена, KRW": None,
    }).reset_index(drop=True)


def main(sources, stock_path, target, rfq_path=None):
    offers = {name: read_any(path) for name, path in sources}
    asked = 0
    if rfq_path:
        # Основа — наша заявка: тогда видно и то, что поставщик закрыл, и то,
        # чего он не предложил вовсе.
        rfq = read_rfq(rfq_path)
        asked = len(rfq)
        offers = {"ЗАЯВКА": rfq, **offers}
        table = merge(offers).drop(columns=["ЗАЯВКА"], errors="ignore")
    else:
        table = merge(offers)
    names = [name for name in offers if name != "ЗАЯВКА"]
    if asked:
        table["В заявке"] = ["да" if index < asked else "нет"
                             for index in range(len(table))]

    stock = read_order(stock_path, everything=True)
    stock["Себестоимость, руб"] = stock["Себестоимость, руб"] / 1.1
    pairs = name_match.to_one(table["Товар"], stock["Наименование"])
    table["Наша себестоимость, руб"] = [
        round(stock.at[pairs[index], "Себестоимость, руб"]) if index in pairs else ""
        for index in table.index]

    for name in names:
        table[f"{name}, руб"] = [
            round(value * KRW_RUB * IMPORT_COST) if pd.notna(value) else ""
            for value in table[name]]
    прайсы = table[names].apply(pd.to_numeric, errors="coerce")
    table["Лучший"] = [прайсы.loc[index].idxmin()
                       if прайсы.loc[index].notna().any() else ""
                       for index in table.index]
    table["Лучшая цена, руб"] = [
        round(прайсы.loc[index].min() * KRW_RUB * IMPORT_COST)
        if прайсы.loc[index].notna().any() else "" for index in table.index]

    columns = (["Бренд", "Товар", "Штрихкод", "Кол-во, шт"] + names
               + [f"{name}, руб" for name in names]
               + ["Лучший", "Лучшая цена, руб", "Наша себестоимость, руб"]
               + (["В заявке"] if asked else []))
    table = table.sort_values(["Бренд", "Товар"])

    book = Workbook()
    ws = book.active
    ws.title = "ОТВЕТЫ"
    ws.append(["№"] + columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    brand, number = None, 0
    for _, item in table.iterrows():
        if item["Бренд"] != brand:
            brand = item["Бренд"]
            ws.append([None, brand])
            for cell in ws[ws.max_row]:
                cell.fill, cell.font = BRAND, Font(bold=True)
        number += 1
        ws.append([number] + [item[column] if pd.notna(item[column]) else ""
                              for column in columns])
    for index, width in enumerate([5, 16, 56, 15, 11] + [11] * len(names)
                                  + [12] * len(names) + [14, 15, 18, 10], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "C2"

    os.makedirs("outputs", exist_ok=True)
    book.save(target)
    for name in names:
        got = table[name].notna()
        if asked:
            inside = int((got & (table["В заявке"] == "да")).sum())
            print(f"{name}: {inside} позиций заявки из {asked}, "
                  f"еще {int(got.sum()) - inside} предложил сверх нее")
        else:
            print(f"{name}: {int(got.sum())} позиций с ценой")
    if asked:
        covered = table[names].notna().any(axis=1) & (table["В заявке"] == "да")
        print(f"Заявка закрыта: {int(covered.sum())} позиций из {asked}")
    print(f"Всего строк: {len(table)}")
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ответы поставщиков в одной таблице")
    parser.add_argument("offers", nargs="+", help="ИМЯ=файл")
    parser.add_argument("--stock", required=True)
    parser.add_argument("--rfq", help="наша заявка: считать по ее позициям")
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main([offer.split("=", 1) for offer in args.offers], args.stock, args.out,
         args.rfq)
