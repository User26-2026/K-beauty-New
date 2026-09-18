"""Прайс по минимальным ценам: лист на бренд, цены в вонах.

По каждой позиции берем самую низкую цену среди корейских поставщиков и
рядом пишем компанию, которая ее дает. Получается рабочий прайс для
закупа: видно, к кому идти за каждой позицией бренда и сколько мы
переплатим, если возьмем все у второго.

Запуск:
    python3 tools/best_prices_by_brand.py
    python3 tools/best_prices_by_brand.py --brand CELIMAX --brand FARMSTAY
    python3 tools/best_prices_by_brand.py --offer "J2K=ответ.xlsx"
"""

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from price_unit import unify_packs
from rates import rub_per_unit
from supplier_brand_matrix import load, price_table, split_conflicts, with_offers

TARGET = "outputs/Прайс по минимальным ценам.xlsx"
COLUMN = "Цена за штуку (сводно)"

HEAD = PatternFill("solid", fgColor="1F3864")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
ALONE = PatternFill("solid", fgColor="FFF2CC")
WHITE = Font(color="FFFFFF", bold=True)

SUMMARY = ["Бренд", "Позиций", "Компаний", "Сумма по минимальным ценам, KRW",
           "Себестоимость, руб", "Основная компания", "Позиций у нее",
           "Остальные компании", "Лист"]
ITEMS = ["№", "Штрихкод", "Название", "Объем", "Мин. цена, KRW",
         "Компания в Корее", "Себестоимость, руб", "Вторая компания",
         "Ее цена, KRW", "Дороже, %"]


def sheet_name(brand, used):
    """Имя листа: Excel не дает больше 31 знака и запрещает : \\ / ? * [ ]."""
    name = re.sub(r"[:\\/?*\[\]]", " ", str(brand)).strip()[:31] or "БЕЗ БРЕНДА"
    if name.upper() in used:
        for number in range(2, 100):
            tail = f" {number}"
            name = name[:31 - len(tail)] + tail
            if name.upper() not in used:
                break
    used.add(name.upper())
    return name


def best_prices(df, suppliers):
    """По каждой позиции: минимальная цена, компания и второй по цене."""
    table, prices = price_table(df, COLUMN)
    columns = [name for name in suppliers if name in prices.columns]
    values = prices[columns]

    rows = table.copy()
    rows["Мин. цена, KRW"] = values.min(axis=1).round(0)
    rows["Компания в Корее"] = values.idxmin(axis=1)
    rows["Компаний с позицией"] = values.notna().sum(axis=1)

    # Второй по цене нужен как запасной вариант: у первой компании может не
    # быть остатка, а разница показывает цену такой замены.
    ranked = values.rank(axis=1, method="first")
    second = values.where(ranked == 2)
    rows["Вторая компания"] = pd.NA
    has_second = second.notna().any(axis=1)
    rows.loc[has_second, "Вторая компания"] = second[has_second].idxmin(axis=1)
    rows["Ее цена, KRW"] = pd.to_numeric(second.min(axis=1), errors="coerce").round(0)
    rows["Дороже, %"] = ((rows["Ее цена, KRW"] / rows["Мин. цена, KRW"] - 1)
                         * 100).round(1)
    return rows[rows["Мин. цена, KRW"].notna()]


def write_brand(book, brand, part, rate, used):
    ws = book.create_sheet(sheet_name(brand, used))
    ws.append(ITEMS)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    part = part.sort_values(["Компания в Корее", "Мин. цена, KRW"],
                            ascending=[True, False])
    for number, (_, item) in enumerate(part.iterrows(), start=1):
        ws.append([
            number, item.name, item["Название EN"], item["Объем"],
            item["Мин. цена, KRW"], item["Компания в Корее"],
            round(item["Мин. цена, KRW"] * rate),
            item["Вторая компания"] if pd.notna(item["Вторая компания"]) else "",
            item["Ее цена, KRW"] if pd.notna(item["Ее цена, KRW"]) else "",
            item["Дороже, %"] if pd.notna(item["Дороже, %"]) else "",
        ])
        # Желтым — позиции, которые есть только у одной компании: выбора нет.
        if item["Компаний с позицией"] == 1:
            ws[f"F{ws.max_row}"].fill = ALONE

    ws.append([None, "ИТОГО", f"позиций: {len(part)}", "",
               int(part["Мин. цена, KRW"].sum()), "",
               round(part["Мин. цена, KRW"].sum() * rate), "", "", ""])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate([5, 15, 58, 14, 14, 18, 16, 18, 13, 11], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in "EGI":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    ws.freeze_panes = "C2"
    return ws.title


def write_summary(book, rows):
    ws = book.create_sheet("ИТОГО ПО БРЕНДАМ", 0)
    ws.append(SUMMARY)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    ws.append(["ВСЕГО", sum(row[1] for row in rows), "",
               sum(row[3] for row in rows), sum(row[4] for row in rows),
               "", "", "", ""])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate([22, 10, 11, 26, 20, 20, 13, 46, 20], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in "BDE":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    ws.freeze_panes = "A2"


def main(country, only, offers, target):
    df = with_offers(load(country), offers, country)
    suppliers = sorted(df["Поставщик"].unique())
    df, _ = split_conflicts(df)
    df, _ = unify_packs(df)
    rate = rub_per_unit(df["Валюта"].mode().iat[0], imported=country == "KR")

    rows = best_prices(df, suppliers)
    if only:
        wanted = {name.upper() for name in only}
        rows = rows[rows["Бренд"].astype(str).str.upper().isin(wanted)]
        if rows.empty:
            raise SystemExit(f"Бренды не найдены: {', '.join(only)}")

    book = Workbook()
    book.remove(book.active)
    summary, used = [], set()
    order = (rows.groupby("Бренд")["Мин. цена, KRW"].sum()
             .sort_values(ascending=False).index)
    for brand in order:
        part = rows[rows["Бренд"] == brand]
        title = write_brand(book, brand, part, rate, used)
        companies = part["Компания в Корее"].value_counts()
        summary.append([
            brand, len(part), int(part["Компания в Корее"].nunique()),
            int(part["Мин. цена, KRW"].sum()),
            round(part["Мин. цена, KRW"].sum() * rate),
            companies.index[0], int(companies.iloc[0]),
            ", ".join(f"{name} — {count}" for name, count
                      in companies.items() if name != companies.index[0]),
            title,
        ])
    write_summary(book, summary)
    os.makedirs("outputs", exist_ok=True)
    book.save(target)

    print(f"Брендов: {len(summary)}   позиций: {len(rows)}")
    print("Кто дает минимальную цену, позиций:")
    print(rows["Компания в Корее"].value_counts().to_string())
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Прайс по минимальным ценам")
    parser.add_argument("--country", default="KR")
    parser.add_argument("--brand", action="append", default=[],
                        help="только эти бренды")
    parser.add_argument("--offer", action="append", default=[],
                        help="ответ на заявку: ИМЯ=файл")
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main(args.country, args.brand,
         [item.split("=", 1) for item in args.offer], args.out)
