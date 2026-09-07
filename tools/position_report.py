"""Наглядная таблица по каждой позиции склада.

Сводим в одну строку все, чем позиция описывается: сколько лежит,
сколько едет, сколько продали за месяц и на какую сумму, сколько с нее
заработали и в процентах, и на сколько месяцев хватит остатка с учетом
сезона.

Прибыль берем фактическую из отчета WB — после комиссии, логистики,
хранения, рекламы и налога. Проценты считаем двумя: к себестоимости
(сколько дает вложенный рубль) и к выручке (сколько остается с рубля
продаж).

Запуск:
    python3 tools/position_report.py <остатки.xlsx> --container <инвойс.xlsx>
"""

import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import name_match
from add_sales_price import sales_by_stock
from brand_names import brand_of
from check_customer_order import read_order
from stock_forecast import run_out
from stock_with_incoming import read_container, truck
from wholesale_vs_wb import MARGIN, SKIP

TARGET = "outputs/Позиции — продажи, остатки, прибыль.xlsx"

HEAD = PatternFill("solid", fgColor="1F3864")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
LOSS = PatternFill("solid", fgColor="F8696B")
IDLE = PatternFill("solid", fgColor="D9D9D9")
WHITE = Font(color="FFFFFF", bold=True)

SHOWN = ["№", "Бренд", "Наименование", "Остаток, шт", "В пути, шт",
         "Продано за август, шт", "Выручка за август, руб",
         "Себестоимость, руб/шт", "Прибыль на штуку, руб",
         "Прибыль за август, руб", "Рентабельность, %", "Маржинальность, %",
         "Хватит на, мес", "Прибыль на весь остаток, руб"]
WIDTHS = [5, 16, 68, 12, 12, 14, 16, 14, 14, 15, 13, 13, 12, 17]


def read_report(path, column, title):
    columns = {1: "Товар", 15: "Продано, шт", column: title}
    table = pd.read_excel(path, sheet_name=0, header=None).iloc[SKIP:]
    table = table[list(columns)].rename(columns=columns)
    table = table[table["Товар"].notna()]
    for name in ("Продано, шт", title):
        table[name] = pd.to_numeric(table[name], errors="coerce").fillna(0)
    return table.reset_index(drop=True)


def incoming(base, container_path):
    coming = pd.concat([read_container(container_path), truck()], ignore_index=True)
    pairs = name_match.to_one(coming["Товар"], base["Наименование"])
    totals = pd.Series(0.0, index=base.index)
    for index, item in coming.iterrows():
        if index in pairs:
            totals[pairs[index]] += item["Кол-во"]
    return totals


def build(stock, sales_path, container_path):
    money = read_report(sales_path, 21, "Выручка, руб")
    profit = read_report(sales_path, 123, "Выручка, руб")
    sold = sales_by_stock(stock, money)
    earned = sales_by_stock(stock, profit)

    stock = stock.copy()
    stock["Бренд"] = stock["Наименование"].map(brand_of)
    stock["В пути, шт"] = incoming(stock, container_path)
    stock["Продано за август, шт"] = sold["Продано, шт"]
    stock["Выручка за август, руб"] = sold["Выручка, руб"].round()
    stock["Прибыль за август, руб"] = earned["Выручка, руб"].round()

    rows = []
    for _, item in stock.iterrows():
        rest, coming = item["Остаток, шт"], item["В пути, шт"]
        monthly, cost = item["Продано за август, шт"], item["Себестоимость, руб"]
        revenue, gain = item["Выручка за август, руб"], item["Прибыль за август, руб"]
        unit = round(gain / monthly, 1) if monthly else ""
        rows.append([
            0, item["Бренд"], item["Наименование"], int(rest), int(coming),
            int(monthly), revenue, cost, unit, gain,
            # Рентабельность — к вложенному рублю, маржинальность — к выручке.
            round(gain / (cost * monthly) * 100, 1) if monthly and cost else "",
            round(gain / revenue * 100, 1) if revenue else "",
            round(run_out(rest + coming, monthly)[1], 1) if monthly else "",
            round(unit * (rest + coming)) if monthly else "",
        ])
    return pd.DataFrame(rows, columns=SHOWN)


def write(book, title, rows, total=None):
    ws = book.create_sheet(title)
    ws.append(SHOWN)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
        gain = row[9]
        fill = (IDLE if not row[5] else LOSS if isinstance(gain, (int, float))
                and gain < 0 else None)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for letter in "DEFGJN":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0"
    for letter in "HI":
        for cell in ws[letter][1:]:
            cell.number_format = "# ##0.00"
    if ws.max_row > 2:
        ws.conditional_formatting.add(
            f"K2:K{ws.max_row - 1}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=25, mid_color="FFEB84",
                           end_type="num", end_value=60, end_color="63BE7B"))
    ws.freeze_panes = "C2"
    return ws


def main(source, sales_path, container_path, target):
    table = build(read_order(source, everything=True), sales_path, container_path)
    order = table.sort_values("Прибыль за август, руб", ascending=False)
    order["№"] = range(1, len(order) + 1)

    book = Workbook()
    book.remove(book.active)
    numbers = ["Остаток, шт", "В пути, шт", "Продано за август, шт",
               "Выручка за август, руб", "Прибыль за август, руб",
               "Прибыль на весь остаток, руб"]
    totals = {name: int(pd.to_numeric(table[name], errors="coerce").sum())
              for name in numbers}
    write(book, "ПОЗИЦИИ", order.values.tolist(),
          ["", "ИТОГО", f"позиций: {len(table)}", totals["Остаток, шт"],
           totals["В пути, шт"], totals["Продано за август, шт"],
           totals["Выручка за август, руб"], "", "",
           totals["Прибыль за август, руб"],
           round(totals["Прибыль за август, руб"]
                 / (table["Себестоимость, руб/шт"]
                    * table["Продано за август, шт"]).sum() * 100, 1),
           round(totals["Прибыль за август, руб"]
                 / totals["Выручка за август, руб"] * 100, 1), "",
           totals["Прибыль на весь остаток, руб"]])

    свод = book.create_sheet("ПО БРЕНДАМ")
    свод.append(["Бренд", "Позиций", "Остаток, шт", "Продано за август, шт",
                 "Выручка за август, руб", "Прибыль за август, руб",
                 "Маржинальность, %", "Прибыль на весь остаток, руб"])
    for cell in свод[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for brand, part in table.groupby("Бренд", sort=True):
        revenue = part["Выручка за август, руб"].sum()
        gain = part["Прибыль за август, руб"].sum()
        свод.append([brand, len(part), int(part["Остаток, шт"].sum()),
                     int(part["Продано за август, шт"].sum()), int(revenue),
                     int(gain), round(gain / revenue * 100, 1) if revenue else "",
                     int(pd.to_numeric(part["Прибыль на весь остаток, руб"],
                                       errors="coerce").sum())])
    свод.append(["ВСЕГО", len(table), totals["Остаток, шт"],
                 totals["Продано за август, шт"], totals["Выручка за август, руб"],
                 totals["Прибыль за август, руб"],
                 round(totals["Прибыль за август, руб"]
                       / totals["Выручка за август, руб"] * 100, 1),
                 totals["Прибыль на весь остаток, руб"]])
    for cell in свод[свод.max_row]:
        cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate([20, 10, 13, 14, 16, 16, 14, 17], start=1):
        свод.column_dimensions[get_column_letter(index)].width = width
    for letter in "CDEFH":
        for cell in свод[letter][1:]:
            cell.number_format = "# ##0"
    свод.freeze_panes = "A2"

    os.makedirs("outputs", exist_ok=True)
    book.save(target)
    print(f"Позиций: {len(table)}   продано за август: "
          f"{totals['Продано за август, шт']:,} шт".replace(",", " "))
    print(f"Выручка: {totals['Выручка за август, руб']:,} руб   "
          f"прибыль: {totals['Прибыль за август, руб']:,} руб".replace(",", " "))
    print(f"Прибыль на весь остаток: "
          f"{totals['Прибыль на весь остаток, руб']:,} руб".replace(",", " "))
    print(f"Сохранено: {target}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Таблица по позициям склада")
    parser.add_argument("source")
    parser.add_argument("--sales", default="data/sales/wb_sales_2026-08.xls")
    parser.add_argument("--container", required=True)
    parser.add_argument("--out", default=TARGET)
    args = parser.parse_args()
    main(args.source, args.sales, args.container, args.out)
