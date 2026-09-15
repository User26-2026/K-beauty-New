"""Прайс SAFIYA по брендам и товарам.

Лист на каждый бренд плюс своды. Основа — действующий B2B-прайс. Рядом
ставим персональное предложение, которое компания дала знакомому, и
публичный прайс: видно, сколько SAFIYA готова уступать с прайса и
сколько мы выигрываем к публичной цене.

Публичная цена идет с НДС 22%, прайсы в долларах — без НДС, поэтому к
публичной считаем две разницы: как есть и к цене, очищенной от НДС.

Запуск:
    python3 tools/build_safiya_offer.py
"""

import csv
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rates import USD_RUB

ROOT = "data/price_lists/safiya"
B2B = f"{ROOT}/safiya_2026-09-01_b2b.xlsx"
PERSONAL = f"{ROOT}/archive/safiya_2026-08-12_znakomomu.xlsx"
PUBLIC = f"{ROOT}/official/safiya_2026-08-04_official.csv"
TARGET = "outputs/SAFIYA_прайс_по_брендам.xlsx"
VAT = 1.22

HEAD = PatternFill("solid", fgColor="1F3864")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)

COLUMNS = ["Наименование", "Объем", "Штрихкод", "Цена B2B, $", "Цена B2B, руб",
           "Знакомому 12.08, $", "Дороже, чем знакомому, %", "Остаток на 12.08, шт",
           "Публичная цена A, руб", "Дешевле публичной, %"]


def read_b2b():
    """B2B-прайс: бренд стоит заголовком секции, товары — строками ниже."""
    table = pd.read_excel(B2B, header=0)
    table.columns = ["Код", "Наименование", "Объем", "Цена, $", "Кол-во", "Сумма"]
    goods, brand = [], ""
    for _, row in table.iterrows():
        code = str(row["Код"]).strip()
        if code.isdigit() and len(code) >= 12:
            goods.append({"Бренд": brand, "Штрихкод": code,
                          "Наименование": str(row["Наименование"]).strip(),
                          "Объем": row["Объем"], "Цена B2B, $": float(row["Цена, $"])})
        elif code and code != "nan" and pd.isna(row["Цена, $"]):
            brand = code
    goods = pd.DataFrame(goods)
    goods["Цена B2B, руб"] = (goods["Цена B2B, $"] * USD_RUB).round()
    return goods


def read_personal():
    table = pd.read_excel(PERSONAL, skiprows=4, header=None)
    table.columns = ["Бренд", "Категория", "Наименование", "Штрихкод",
                     "Объем", "Остаток на 12.08, шт", "Знакомому 12.08, $"]
    table = table.dropna(subset=["Знакомому 12.08, $"])
    table["Штрихкод"] = table["Штрихкод"].astype(str).str.replace(r"\D", "", regex=True)
    return table.set_index("Штрихкод")[["Знакомому 12.08, $", "Остаток на 12.08, шт"]]


def read_public():
    with open(PUBLIC, encoding="utf-8") as handle:
        return {row["Штрихкод"]: int(row["Цена A, руб"])
                for row in csv.DictReader(handle, delimiter=";")}


def build():
    goods = read_b2b()
    personal, public = read_personal(), read_public()

    goods = goods.join(personal, on="Штрихкод")
    goods["Дороже, чем знакомому, %"] = [
        round(new / old * 100 - 100, 1) if pd.notna(old) else ""
        for new, old in zip(goods["Цена B2B, $"], goods["Знакомому 12.08, $"])
    ]
    goods["Публичная цена A, руб"] = [public.get(code, "") for code in goods["Штрихкод"]]
    # Скидка положительным числом: сколько процентов публичной цены мы не платим.
    goods["Дешевле публичной, %"] = [
        round(100 - price / base * 100, 1) if base else ""
        for price, base in zip(goods["Цена B2B, руб"], goods["Публичная цена A, руб"])
    ]
    goods["Без НДС, %"] = [
        round(100 - price / (base / VAT) * 100, 1) if base else ""
        for price, base in zip(goods["Цена B2B, руб"], goods["Публичная цена A, руб"])
    ]
    return goods.fillna({"Знакомому 12.08, $": "", "Остаток на 12.08, шт": ""})


def sheet_name(brand):
    """Имя листа: Excel не берет длиннее 31 знака и служебные символы."""
    return re.sub(r"[\\/*?:\[\]]", "-", brand)[:31]


def write(ws, columns, rows, widths, total=None):
    ws.append(columns)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = TOTAL, Font(bold=True)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return ws


def money(ws, letters, fmt="# ##0"):
    for letter in letters:
        for cell in ws[letter][1:]:
            cell.number_format = fmt


def scale(ws, letter, rows, low, high, reverse=False):
    if rows < 2:
        return
    colors = ("F8696B", "FFEB84", "63BE7B")
    if reverse:
        colors = colors[::-1]
    ws.conditional_formatting.add(
        f"{letter}2:{letter}{rows}",
        ColorScaleRule(start_type="num", start_value=low, start_color=colors[0],
                       mid_type="num", mid_value=(low + high) / 2, mid_color=colors[1],
                       end_type="num", end_value=high, end_color=colors[2]))


def median(values):
    values = [value for value in values if value != ""]
    return round(pd.Series(values).median(), 1) if values else ""


def main():
    goods = build()
    book = Workbook()
    book.remove(book.active)

    summary = []
    for brand, part in goods.groupby("Бренд", sort=True):
        summary.append([
            brand, len(part),
            int(part["Цена B2B, руб"].mean()),
            int(part["Цена B2B, руб"].min()),
            int(part["Цена B2B, руб"].max()),
            int((part["Знакомому 12.08, $"] != "").sum()),
            median(part["Дороже, чем знакомому, %"]),
            int((part["Дешевле публичной, %"] != "").sum()),
            median(part["Дешевле публичной, %"]),
        ])
    итого = write(
        book.create_sheet("ИТОГО"),
        ["Бренд", "Позиций", "Средняя цена, руб", "Мин, руб", "Макс, руб",
         "Есть в прайсе знакомому, поз.", "Дороже, чем знакомому, %",
         "Есть в публичном, поз.", "Дешевле публичной, %"],
        summary,
        [24, 9, 15, 11, 11, 14, 12, 14, 14],
        ["ВСЕГО", len(goods), int(goods["Цена B2B, руб"].mean()),
         int(goods["Цена B2B, руб"].min()), int(goods["Цена B2B, руб"].max()),
         int((goods["Знакомому 12.08, $"] != "").sum()), median(goods["Дороже, чем знакомому, %"]),
         int((goods["Дешевле публичной, %"] != "").sum()),
         median(goods["Дешевле публичной, %"])],
    )
    money(итого, "CDE")
    scale(итого, "G", итого.max_row - 1, 0, 20, reverse=True)
    scale(итого, "I", итого.max_row - 1, 0, 45)

    рост = goods[goods["Дороже, чем знакомому, %"] != ""].sort_values("Дороже, чем знакомому, %", ascending=False)
    лист = write(
        book.create_sheet("ЦЕНА ЗНАКОМОМУ"),
        ["Бренд", "Наименование", "Объем", "Штрихкод", "Знакомому 12.08, $",
         "Прайс 01.09, $", "Дороже, чем знакомому, %"],
        рост[["Бренд", "Наименование", "Объем", "Штрихкод",
              "Знакомому 12.08, $", "Цена B2B, $", "Дороже, чем знакомому, %"]].values.tolist(),
        [18, 58, 12, 16, 12, 13, 12],
    )
    money(лист, "EF", "0.00")
    scale(лист, "G", лист.max_row, 0, 20, reverse=True)

    сверка = goods[goods["Дешевле публичной, %"] != ""].sort_values(
        "Дешевле публичной, %", ascending=False)
    лист = write(
        book.create_sheet("СКИДКА К ПУБЛИЧНОМУ"),
        ["Бренд", "Наименование", "Объем", "Штрихкод", "Цена B2B, руб",
         "Публичная цена A, руб", "Дешевле публичной, %", "Без НДС, %"],
        сверка[["Бренд", "Наименование", "Объем", "Штрихкод", "Цена B2B, руб",
                "Публичная цена A, руб", "Дешевле публичной, %",
                "Без НДС, %"]].values.tolist(),
        [18, 58, 12, 16, 13, 15, 14, 12],
    )
    money(лист, "EF")
    scale(лист, "G", лист.max_row, 0, 45)

    for brand, part in goods.groupby("Бренд", sort=True):
        part = part.sort_values("Наименование")
        ws = write(
            book.create_sheet(sheet_name(brand)),
            COLUMNS,
            part[COLUMNS].values.tolist(),
            [58, 12, 16, 12, 13, 12, 12, 13, 15, 14],
            ["ИТОГО", f"позиций: {len(part)}", "", "",
             int(part["Цена B2B, руб"].mean()), "", median(part["Дороже, чем знакомому, %"]), "", "",
             median(part["Дешевле публичной, %"])],
        )
        money(ws, "EI")
        money(ws, "DF", "0.00")
        scale(ws, "G", ws.max_row - 1, 0, 20, reverse=True)
        scale(ws, "J", ws.max_row - 1, 0, 45)

    os.makedirs("outputs", exist_ok=True)
    book.save(TARGET)
    print(f"Брендов: {goods['Бренд'].nunique()}   позиций: {len(goods)}")
    print(f"Дороже прайса знакомому от 12.08 на {median(goods['Дороже, чем знакомому, %'])}% "
          f"по {(goods['Дороже, чем знакомому, %'] != '').sum()} позициям")
    print(f"Дешевле публичного прайса: {median(goods['Дешевле публичной, %'])}% "
          f"по {(goods['Дешевле публичной, %'] != '').sum()} позициям")
    print(f"Сохранено: {TARGET}")


if __name__ == "__main__":
    main()
